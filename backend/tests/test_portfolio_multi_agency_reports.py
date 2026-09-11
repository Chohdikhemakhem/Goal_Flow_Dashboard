from __future__ import annotations

import os
import sys
import unittest
from datetime import date, datetime, timezone
from decimal import Decimal
from io import BytesIO
from pathlib import Path

BACKEND_DIR = Path(__file__).resolve().parents[1]
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

TEST_DB_PATH = BACKEND_DIR / "tmp_portfolio_multi_agency.db"
os.environ["DATABASE_URL"] = f"sqlite:///{TEST_DB_PATH.as_posix()}"
os.environ["JWT_SECRET_KEY"] = "PortfolioMultiAgencyTestsSecretKey!2026"
os.environ["ENVIRONMENT"] = "test"
os.environ["COOKIE_SECURE"] = "false"
os.environ["BONUS_MODULE_ACTIVE"] = "false"
os.environ["BOOTSTRAP_SUPERADMIN_EMAIL"] = ""
os.environ["BOOTSTRAP_SUPERADMIN_PASSWORD"] = ""

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select

from app.core.config import get_settings

get_settings.cache_clear()

from app.main import app
from app.api.v1.reports import _portfolio_report_base_rows
from app.core.security import hash_password
from app.db.base import Base
from app.db.migrations import add_missing_columns
from app.db.session import SessionLocal, engine
from app.models.entities import Agency, Agent, ImportBatch, LoanRaw, User
from app.models.enums import ImportBatchType, UserRole


HISTORICAL_HEADERS = [
    "CLIENT NO",
    "CLIENT NAME",
    "CLIENT FIRST NAME",
    "TOTAL DUE AMT",
    "Échéance principal",
    "TOTAL PRINCIPAL DUE AMT",
    "TOTAL INTEREST DUE AMT",
    "JOURS DE RETARD",
    "ENCOURS",
    "TAUX DE REMBOURSEMENT",
]



class PortfolioMultiAgencyReportTestCase(unittest.TestCase):
    admin_email = "multi.agency.admin@microcred.com.tn"
    admin_password = "MultiAgencyAdmin!2026"

    def setUp(self) -> None:
        Base.metadata.drop_all(bind=engine)
        Base.metadata.create_all(bind=engine)
        add_missing_columns(engine)
        self.client = TestClient(app)
        self._seed()

    def tearDown(self) -> None:
        self.client.close()

    def _seed(self) -> None:
        with SessionLocal() as db:
            db.add(
                User(
                    email=self.admin_email,
                    full_name="Multi Agency Admin",
                    hashed_password=hash_password(self.admin_password),
                    role=UserRole.ADMIN,
                    password_changed_at=datetime.now(timezone.utc),
                )
            )

            self.agency_ids: dict[str, int] = {}
            for agency_name in ["Tunis Centre", "Ariana", "Sousse"]:
                agency = Agency(name=agency_name)
                db.add(agency)
                db.flush()
                self.agency_ids[agency_name] = agency.id

            for agent_name, agency_name in [
                ("Ahmed Ben Ali", "Tunis Centre"),
                ("Fatma Ben Salem", "Ariana"),
                ("Ali Trabelsi", "Sousse"),
            ]:
                db.add(Agent(name=agent_name, agency_id=self.agency_ids[agency_name]))

            batch = ImportBatch(
                batch_type=ImportBatchType.CURRENT_STATE,
                snapshot_date=date(2026, 7, 31),
                file_name="multi_agency_test.xlsx",
            )
            db.add(batch)
            db.flush()

            rows = [
                ("CL-001", "Client A", "Ahmed Ben Ali", "Tunis Centre", Decimal("5000"), 10),
                ("CL-002", "Client B", "Ahmed Ben Ali", "Tunis Centre", Decimal("3200"), 5),
                ("CL-003", "Client C", "Fatma Ben Salem", "Ariana", Decimal("4100"), 8),
                ("CL-004", "Client D", "Ali Trabelsi", "Sousse", Decimal("2900"), 3),
                # days_overdue <= 2: must be excluded from the arrears report.
                ("CL-005", "Client E", "Ali Trabelsi", "Sousse", Decimal("1000"), 0),
            ]
            for index, (client_id, client_name, agent_name, agency_name, outstanding, days_overdue) in enumerate(rows):
                db.add(
                    LoanRaw(
                        contract_no=f"CT-{index:03d}",
                        client_id=client_id,
                        client_name=client_name,
                        client_first_name="Prenom",
                        agency_name=agency_name,
                        agent_name=agent_name,
                        teg_rate=Decimal("0"),
                        disbursement_amount=outstanding * 2,
                        principal_outstanding=outstanding,
                        principal_due=Decimal("100"),
                        interest_due_amt=Decimal("10"),
                        total_scheduled_amount=Decimal("100"),
                        days_overdue=days_overdue,
                        total_due=Decimal("110"),
                        disbursement_date=date(2026, 1, 10),
                        snapshot_date=date(2026, 7, 31),
                        import_batch_id=batch.id,
                    )
                )
            db.commit()

    def _login(self) -> None:
        response = self.client.post(
            "/api/v1/auth/login",
            json={"email": self.admin_email, "password": self.admin_password},
        )
        self.assertEqual(response.status_code, 200, response.text)

    def _download_excel(self, params):
        return self.client.get("/api/v1/reports/portfolio.xlsx", params=params)

    def _sheet_rows(self, content: bytes) -> list[list]:
        workbook = load_workbook(BytesIO(content))
        sheet = workbook.active
        return [[cell.value for cell in row] for row in sheet.iter_rows()]

    # ------------------------------------------------------------------
    # Test 1 / Test 4 (deselection) — one agency keeps the historical template
    # ------------------------------------------------------------------
    def test_single_agency_via_agency_id_keeps_historical_template(self) -> None:
        self._login()
        response = self._download_excel(
            {"report_type": "arrears_list", "agency_id": self.agency_ids["Tunis Centre"]}
        )
        self.assertEqual(response.status_code, 200, response.text)
        self.assertIn("portfolio_arrears_list.xlsx", response.headers.get("content-disposition", ""))

        rows = self._sheet_rows(response.content)
        flat_texts = {str(value) for row in rows for value in row if value is not None}
        self.assertNotIn("AGENCE", flat_texts)
        self.assertTrue(any(value.startswith("GP: ") for value in flat_texts))
        # Grouped historical template: section row 5, header row 6.
        header = [value for value in rows[5] if value is not None]
        self.assertEqual(header, HISTORICAL_HEADERS)

    def test_single_agency_via_agency_ids_keeps_historical_template(self) -> None:
        self._login()
        response = self._download_excel(
            {"report_type": "arrears_list", "agency_ids": str(self.agency_ids["Tunis Centre"])}
        )
        self.assertEqual(response.status_code, 200, response.text)
        self.assertIn("portfolio_arrears_list.xlsx", response.headers.get("content-disposition", ""))


    # ------------------------------------------------------------------
    # Test 2 — two agencies: consolidated multi-agency template
    # ------------------------------------------------------------------
    def test_two_agencies_use_multi_agency_template(self) -> None:
        self._login()
        agency_ids = f"{self.agency_ids['Tunis Centre']},{self.agency_ids['Ariana']}"
        response = self._download_excel({"report_type": "arrears_list", "agency_ids": agency_ids})
        self.assertEqual(response.status_code, 200, response.text)
        self.assertIn(
            "portfolio_arrears_list_multi_agences.xlsx",
            response.headers.get("content-disposition", ""),
        )

        rows = self._sheet_rows(response.content)
        header = [value for value in rows[4] if value is not None]
        self.assertEqual(["AGENCE", "AGENT"] + HISTORICAL_HEADERS, header)

        data_rows = [row for row in rows[5:] if row[0] is not None]
        self.assertEqual(3, len(data_rows))
        self.assertEqual({row[0] for row in data_rows}, {"Tunis Centre", "Ariana"})
        self.assertEqual({row[1] for row in data_rows}, {"Ahmed Ben Ali", "Fatma Ben Salem"})

    # ------------------------------------------------------------------
    # Test 3 — three agencies: all data in one single consolidated table
    # ------------------------------------------------------------------
    def test_three_agencies_consolidated_in_one_table(self) -> None:
        self._login()
        agency_ids = ",".join(str(value) for value in self.agency_ids.values())
        response = self._download_excel({"report_type": "arrears_list", "agency_ids": agency_ids})
        self.assertEqual(response.status_code, 200, response.text)

        rows = self._sheet_rows(response.content)
        data_rows = [row for row in rows[5:] if row[0] is not None]
        self.assertEqual(4, len(data_rows))
        self.assertEqual({row[0] for row in data_rows}, {"Tunis Centre", "Ariana", "Sousse"})

    # ------------------------------------------------------------------
    # Test 7 — totals: multi report equals sum of single-agency reports
    # ------------------------------------------------------------------
    def test_multi_agency_totals_match_single_agency_reports(self) -> None:
        self._login()

        def contracts_and_total(agency_ids: list[int]) -> tuple[set[str], Decimal]:
            with SessionLocal() as db:
                user = db.scalars(select(User).where(User.email == self.admin_email)).one()
                loans = _portfolio_report_base_rows(
                    db=db,
                    user=user,
                    agency_id=agency_ids[0] if len(agency_ids) == 1 else None,
                    agency_ids=agency_ids if len(agency_ids) > 1 else None,
                )
                contracts = {loan.contract_no for loan in loans}
                total = sum(
                    (loan.principal_outstanding or Decimal("0")) + (loan.principal_due or Decimal("0"))
                    for loan in loans
                )
                return contracts, Decimal(total)

        tunis = self.agency_ids["Tunis Centre"]
        ariana = self.agency_ids["Ariana"]

        single_contracts: set[str] = set()
        single_total = Decimal("0")
        for agency_id in (tunis, ariana):
            contracts, total = contracts_and_total([agency_id])
            single_contracts |= contracts
            single_total += total

        multi_contracts, multi_total = contracts_and_total([tunis, ariana])

        self.assertEqual(single_contracts, multi_contracts)
        self.assertEqual(single_total, multi_total)
        self.assertEqual({"CT-000", "CT-001", "CT-002"}, multi_contracts)

    # ------------------------------------------------------------------
    # Test 5 / Test 6 — validation, authorization and PDF naming
    # ------------------------------------------------------------------
    def test_unknown_agency_in_multi_selection_is_rejected(self) -> None:
        self._login()
        agency_ids = f"{self.agency_ids['Tunis Centre']},99999"
        response = self._download_excel({"report_type": "arrears_list", "agency_ids": agency_ids})
        self.assertEqual(response.status_code, 404, response.text)
        self.assertIn("agency_not_found", response.text)

    def test_admin_without_agency_is_rejected(self) -> None:
        self._login()
        response = self._download_excel({"report_type": "arrears_list"})
        self.assertEqual(response.status_code, 400, response.text)
        self.assertIn("agency_required", response.text)

    def test_multi_agency_pdf_uses_dedicated_template_name(self) -> None:
        self._login()
        agency_ids = f"{self.agency_ids['Tunis Centre']},{self.agency_ids['Sousse']}"
        response = self.client.get(
            "/api/v1/reports/portfolio.pdf",
            params={"report_type": "arrears_list", "agency_ids": agency_ids},
        )
        self.assertEqual(response.status_code, 200, response.text)
        self.assertIn(
            "portfolio_arrears_list_multi_agences.pdf",
            response.headers.get("content-disposition", ""),
        )
        self.assertTrue(response.content.startswith(b"%PDF"))


if __name__ == "__main__":
    unittest.main()

