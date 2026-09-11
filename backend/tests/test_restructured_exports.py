from __future__ import annotations

import os
import sys
import unittest
from datetime import date, datetime, timezone
from decimal import Decimal
from io import BytesIO
from pathlib import Path


BACKEND_DIR = Path(__file__).resolve().parents[1]
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

TEST_DB_PATH = BACKEND_DIR / "tmp_restructured_exports.db"
os.environ["DATABASE_URL"] = f"sqlite:///{TEST_DB_PATH.as_posix()}"
os.environ["JWT_SECRET_KEY"] = "RestructuredExportTestsNeedAVeryStrongSecret!2026"
os.environ["ENVIRONMENT"] = "test"
os.environ["COOKIE_SECURE"] = "false"
os.environ["BONUS_MODULE_ACTIVE"] = "false"
os.environ["BOOTSTRAP_SUPERADMIN_EMAIL"] = ""
os.environ["BOOTSTRAP_SUPERADMIN_PASSWORD"] = ""

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.core.config import get_settings

get_settings.cache_clear()

from app.main import app
from app.core.security import hash_password
from app.db.base import Base
from app.db.migrations import add_missing_columns
from app.db.session import SessionLocal, engine
from app.models.entities import Agency, Agent, ImportBatch, LoanRaw, RestructuredAnalysisResult, RestructuredContract, User
from app.models.enums import ImportBatchType, UserRole
from app.services.restructured import CLOSURE_STATUS_ASSUMED_CLOSED, CLOSURE_STATUS_IN_PROGRESS, CLOSURE_STATUS_NA, CONSOLIDATED_FAMILY, RESTRUCTURED_FAMILY


class RestructuredExportsTestCase(unittest.TestCase):
    super_admin_email = "super.admin@microcred.com.tn"
    super_admin_password = "UltraSecureAdmin!2026"

    def setUp(self) -> None:
        Base.metadata.drop_all(bind=engine)
        Base.metadata.create_all(bind=engine)
        add_missing_columns(engine)
        self.client = TestClient(app)

    def tearDown(self) -> None:
        self.client.close()

    def _create_user(self, *, email: str, password: str, role: UserRole, full_name: str) -> None:
        with SessionLocal() as db:
            user = User(
                email=email,
                full_name=full_name,
                hashed_password=hash_password(password),
                role=role,
                is_active=True,
                must_change_password=False,
                password_changed_at=datetime.now(timezone.utc),
            )
            db.add(user)
            db.commit()

    def _seed_super_admin(self) -> None:
        self._create_user(
            email=self.super_admin_email,
            password=self.super_admin_password,
            role=UserRole.SUPER_ADMIN,
            full_name="Security Super Admin",
        )

    def _login(self) -> None:
        response = self.client.post(
            "/api/v1/auth/login",
            json={"email": self.super_admin_email, "password": self.super_admin_password},
        )
        self.assertEqual(response.status_code, 200, response.text)

    def _seed_dashboard_scope(self) -> dict[str, int]:
        with SessionLocal() as db:
            agency_north = Agency(name="Agence Nord")
            agency_south = Agency(name="Agence Sud")
            db.add_all([agency_north, agency_south])
            db.flush()

            agent_north = Agent(name="GP Nord", agency_id=agency_north.id)
            agent_south = Agent(name="GP Sud", agency_id=agency_south.id)
            db.add_all([agent_north, agent_south])
            db.flush()

            current_batch = ImportBatch(
                batch_type=ImportBatchType.CURRENT_STATE,
                period=None,
                snapshot_date=date(2026, 7, 31),
                file_name="current_state.xlsx",
            )
            db.add(current_batch)
            db.flush()

            db.add_all(
                [
                    RestructuredContract(
                        contract_no="RST-001",
                        normalized_contract_no="RST-001",
                        credit_family=RESTRUCTURED_FAMILY,
                        delay_date=date(2026, 4, 1),
                        category_desc="Credit classique",
                        total_due=Decimal("1500.00"),
                        loan_duration=8,
                    ),
                    RestructuredContract(
                        contract_no="RST-002",
                        normalized_contract_no="RST-002",
                        credit_family=RESTRUCTURED_FAMILY,
                        delay_date=date(2026, 4, 15),
                        category_desc="Credit classique",
                        total_due=Decimal("2500.00"),
                        loan_duration=10,
                    ),
                    RestructuredContract(
                        contract_no="RST-003",
                        normalized_contract_no="RST-003",
                        credit_family=RESTRUCTURED_FAMILY,
                        delay_date=date(2026, 5, 10),
                        category_desc="Credit classique",
                        total_due=Decimal("3500.00"),
                        loan_duration=12,
                    ),
                    RestructuredContract(
                        contract_no="CONS-001",
                        normalized_contract_no="CONS-001",
                        credit_family=CONSOLIDATED_FAMILY,
                        delay_date=date(2026, 3, 20),
                        category_desc="Credits Consolides",
                        total_due=Decimal("4100.00"),
                        loan_duration=6,
                    ),
                ]
            )

            db.add_all(
                [
                    LoanRaw(
                        contract_no="RST-001",
                        client_name="Crédite",
                        client_first_name="Élodie",
                        client_id="CLI-001",
                        agency_name=agency_north.name,
                        agent_name=agent_north.name,
                        category_desc="Commerce",
                        disbursement_amount=Decimal("1000.00"),
                        principal_outstanding=Decimal("900.00"),
                        principal_due=Decimal("100.00"),
                        total_scheduled_amount=Decimal("100.00"),
                        days_overdue=15,
                        total_due=Decimal("1500.00"),
                        status="active",
                        disbursement_date=date(2026, 7, 2),
                        snapshot_date=current_batch.snapshot_date,
                        import_batch_id=current_batch.id,
                    ),
                    LoanRaw(
                        contract_no="RST-002",
                        client_name="Client",
                        client_first_name="Nord",
                        client_id="CLI-002",
                        agency_name=agency_north.name,
                        agent_name=agent_north.name,
                        category_desc="Agriculture",
                        disbursement_amount=Decimal("2000.00"),
                        principal_outstanding=Decimal("1500.00"),
                        principal_due=Decimal("200.00"),
                        total_scheduled_amount=Decimal("200.00"),
                        days_overdue=45,
                        total_due=Decimal("2500.00"),
                        status="active",
                        disbursement_date=date(2026, 7, 5),
                        snapshot_date=current_batch.snapshot_date,
                        import_batch_id=current_batch.id,
                    ),
                    LoanRaw(
                        contract_no="RST-003",
                        client_name="Client",
                        client_first_name="Sud",
                        client_id="CLI-003",
                        agency_name=agency_south.name,
                        agent_name=agent_south.name,
                        category_desc="Service",
                        disbursement_amount=Decimal("3000.00"),
                        principal_outstanding=Decimal("2700.00"),
                        principal_due=Decimal("300.00"),
                        total_scheduled_amount=Decimal("300.00"),
                        days_overdue=0,
                        total_due=Decimal("3500.00"),
                        status="active",
                        disbursement_date=date(2026, 7, 8),
                        snapshot_date=current_batch.snapshot_date,
                        import_batch_id=current_batch.id,
                    ),
                    LoanRaw(
                        contract_no="CONS-001",
                        client_name="Client",
                        client_first_name="Consolide",
                        client_id="CLI-004",
                        agency_name=agency_south.name,
                        agent_name=agent_south.name,
                        category_desc="Credits Consolides",
                        disbursement_amount=Decimal("4000.00"),
                        principal_outstanding=Decimal("3300.00"),
                        principal_due=Decimal("400.00"),
                        total_scheduled_amount=Decimal("400.00"),
                        days_overdue=130,
                        total_due=Decimal("4100.00"),
                        status="active",
                        disbursement_date=date(2026, 7, 9),
                        snapshot_date=current_batch.snapshot_date,
                        import_batch_id=current_batch.id,
                    ),
                ]
            )

            db.add_all(
                [
                    RestructuredAnalysisResult(
                        contract_no="RST-001",
                        agency_name=agency_north.name,
                        agent_name=agent_north.name,
                        delay_date=date(2026, 4, 1),
                        total_due=Decimal("1500.00"),
                        loan_duration=8,
                        schedule_status="OK",
                        consecutive_paid_count=4,
                        max_series_installments="4,5,6,7",
                        max_series_dates="01/04/2026,01/05/2026,01/06/2026,01/07/2026",
                        paid_installments_after_delay=4,
                        anomaly_detected=False,
                        anomaly_detail=None,
                        paid_last_four_status="Oui",
                        last_paid_installment_no=8,
                        last_paid_due_date=date(2026, 7, 1),
                        closure_status=CLOSURE_STATUS_ASSUMED_CLOSED,
                    ),
                    RestructuredAnalysisResult(
                        contract_no="RST-002",
                        agency_name=agency_north.name,
                        agent_name=agent_north.name,
                        delay_date=date(2026, 4, 15),
                        total_due=Decimal("2500.00"),
                        loan_duration=10,
                        schedule_status="OK",
                        consecutive_paid_count=2,
                        max_series_installments="4,5",
                        max_series_dates="15/04/2026,15/05/2026",
                        paid_installments_after_delay=2,
                        anomaly_detected=True,
                        anomaly_detail="Trou entre 5 et 7",
                        paid_last_four_status="Non",
                        last_paid_installment_no=7,
                        last_paid_due_date=date(2026, 6, 15),
                        closure_status=CLOSURE_STATUS_IN_PROGRESS,
                    ),
                    RestructuredAnalysisResult(
                        contract_no="RST-003",
                        agency_name=agency_south.name,
                        agent_name=agent_south.name,
                        delay_date=date(2026, 5, 10),
                        total_due=Decimal("3500.00"),
                        loan_duration=12,
                        schedule_status="OK",
                        consecutive_paid_count=0,
                        max_series_installments=None,
                        max_series_dates=None,
                        paid_installments_after_delay=0,
                        anomaly_detected=False,
                        anomaly_detail=None,
                        paid_last_four_status="N/A",
                        last_paid_installment_no=None,
                        last_paid_due_date=None,
                        closure_status=CLOSURE_STATUS_NA,
                    ),
                    RestructuredAnalysisResult(
                        contract_no="CONS-001",
                        agency_name=agency_south.name,
                        agent_name=agent_south.name,
                        delay_date=date(2026, 3, 20),
                        total_due=Decimal("4100.00"),
                        loan_duration=6,
                        schedule_status="OK",
                        consecutive_paid_count=3,
                        max_series_installments="1,2,3",
                        max_series_dates="20/03/2026,20/04/2026,20/05/2026",
                        paid_installments_after_delay=3,
                        anomaly_detected=False,
                        anomaly_detail=None,
                        paid_last_four_status="Oui",
                        last_paid_installment_no=6,
                        last_paid_due_date=date(2026, 7, 20),
                        closure_status=CLOSURE_STATUS_ASSUMED_CLOSED,
                    ),
                ]
            )
            db.commit()

            return {
                "agency_north_id": agency_north.id,
                "agency_south_id": agency_south.id,
                "agent_north_id": agent_north.id,
                "agent_south_id": agent_south.id,
            }

    def _load_workbook(self, response):
        self.assertEqual(response.status_code, 200, response.text)
        self.assertIn(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            response.headers.get("content-type", ""),
        )
        return load_workbook(filename=BytesIO(response.content))

    def _find_header_row(self, sheet) -> int:
        for row_index in range(1, sheet.max_row + 1):
            if sheet.cell(row=row_index, column=1).value == "Numero de contrat":
                return row_index
        raise AssertionError("Header row not found in exported workbook")

    def _extract_data_rows(self, sheet) -> list[list[object]]:
        header_row = self._find_header_row(sheet)
        rows: list[list[object]] = []
        for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
            first_cell = row[0] if row else None
            if not first_cell:
                continue
            rows.append(list(row))
        return rows

    def test_export_without_filter_returns_all_restructured_rows_not_only_current_page(self) -> None:
        self._seed_super_admin()
        self._seed_dashboard_scope()
        self._login()

        list_response = self.client.get("/api/v1/credits-restructures/contrats?family=restructured&limit=1&offset=0")
        self.assertEqual(list_response.status_code, 200, list_response.text)
        payload = list_response.json()
        self.assertEqual(payload["total"], 3)
        self.assertEqual(len(payload["items"]), 1)

        export_response = self.client.get("/api/v1/credits-restructures/contrats/export.xlsx?family=restructured")
        workbook = self._load_workbook(export_response)
        rows = self._extract_data_rows(workbook.active)
        exported_contracts = [row[0] for row in rows]

        self.assertEqual(len(exported_contracts), 3)
        self.assertEqual(set(exported_contracts), {"RST-001", "RST-002", "RST-003"})
        self.assertIn("detail_contrats_restructures_", export_response.headers.get("content-disposition", ""))

    def test_export_respects_agency_filter_and_sort(self) -> None:
        self._seed_super_admin()
        ids = self._seed_dashboard_scope()
        self._login()

        response = self.client.get(
            f"/api/v1/credits-restructures/contrats/export.xlsx?family=restructured&agency_id={ids['agency_north_id']}&sort_key=total_due&sort_direction=desc"
        )
        workbook = self._load_workbook(response)
        rows = self._extract_data_rows(workbook.active)

        self.assertEqual([row[0] for row in rows], ["RST-002", "RST-001"])
        self.assertTrue(all(row[3] == "Agence Nord" for row in rows))

    def test_export_respects_agent_filter(self) -> None:
        self._seed_super_admin()
        ids = self._seed_dashboard_scope()
        self._login()

        response = self.client.get(
            f"/api/v1/credits-restructures/contrats/export.xlsx?family=restructured&agent_id={ids['agent_south_id']}"
        )
        workbook = self._load_workbook(response)
        rows = self._extract_data_rows(workbook.active)

        self.assertEqual([row[0] for row in rows], ["RST-003"])
        self.assertEqual(rows[0][4], "GP Sud")

    def test_export_respects_family_tab_and_search(self) -> None:
        self._seed_super_admin()
        self._seed_dashboard_scope()
        self._login()

        consolidated_response = self.client.get("/api/v1/credits-restructures/contrats/export.xlsx?family=consolidated")
        consolidated_rows = self._extract_data_rows(self._load_workbook(consolidated_response).active)
        self.assertEqual([row[0] for row in consolidated_rows], ["CONS-001"])

        search_response = self.client.get("/api/v1/credits-restructures/contrats/export.xlsx?family=restructured&q=RST-001")
        search_rows = self._extract_data_rows(self._load_workbook(search_response).active)
        self.assertEqual([row[0] for row in search_rows], ["RST-001"])

    def test_export_preserves_french_characters_in_workbook(self) -> None:
        self._seed_super_admin()
        self._seed_dashboard_scope()
        self._login()

        response = self.client.get("/api/v1/credits-restructures/contrats/export.xlsx?family=restructured&q=RST-001")
        workbook = self._load_workbook(response)
        sheet = workbook.active
        flattened = " | ".join(
            "" if value is None else str(value)
            for row in sheet.iter_rows(values_only=True)
            for value in row
        )

        self.assertIn("Crédite Élodie", flattened)
        self.assertNotIn("Crée", flattened)

    def test_export_requires_authentication(self) -> None:
        response = self.client.get("/api/v1/credits-restructures/contrats/export.xlsx?family=restructured")
        self.assertEqual(response.status_code, 401, response.text)

    def test_detail_table_and_export_include_total_scheduled_amount(self) -> None:
        self._seed_super_admin()
        self._seed_dashboard_scope()
        self._login()

        response = self.client.get("/api/v1/credits-restructures/contrats?family=restructured&limit=10&offset=0")
        self.assertEqual(response.status_code, 200, response.text)
        items = response.json()["items"]
        row_by_contract = {item["contract_no"]: item for item in items}
        self.assertEqual(str(row_by_contract["RST-001"]["total_scheduled_amount"]), "100.00")
        self.assertEqual(str(row_by_contract["RST-002"]["total_scheduled_amount"]), "200.00")
        self.assertEqual(str(row_by_contract["RST-003"]["total_scheduled_amount"]), "300.00")

        export_response = self.client.get("/api/v1/credits-restructures/contrats/export.xlsx?family=restructured")
        workbook = self._load_workbook(export_response)
        sheet = workbook.active
        header_row = self._find_header_row(sheet)
        headers = [sheet.cell(row=header_row, column=index).value for index in range(1, sheet.max_column + 1)]
        self.assertIn("Echeance", headers)
        echeance_index = headers.index("Echeance")

        data_rows = self._extract_data_rows(sheet)
        exported = {row[0]: row for row in data_rows}
        self.assertEqual(exported["RST-001"][echeance_index], 100)
        self.assertEqual(exported["RST-002"][echeance_index], 200)
        self.assertEqual(exported["RST-003"][echeance_index], 300)


if __name__ == "__main__":
    unittest.main()
