from __future__ import annotations

import os
import sys
import unittest
import json
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from io import BytesIO
from pathlib import Path


BACKEND_DIR = Path(__file__).resolve().parents[1]
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

TEST_DB_PATH = BACKEND_DIR / "tmp_security_controls.db"
os.environ["DATABASE_URL"] = f"sqlite:///{TEST_DB_PATH.as_posix()}"
os.environ["JWT_SECRET_KEY"] = "SecurityTestsRequireAVeryStrongSecret!2026"
os.environ["ENVIRONMENT"] = "test"
os.environ["COOKIE_SECURE"] = "false"
os.environ["BONUS_MODULE_ACTIVE"] = "false"
os.environ["BOOTSTRAP_SUPERADMIN_EMAIL"] = ""
os.environ["BOOTSTRAP_SUPERADMIN_PASSWORD"] = ""

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select

from app.core.config import get_settings

get_settings.cache_clear()

from app.main import app
from app.core.security import hash_password
from app.db.base import Base
from app.db.migrations import add_missing_columns
from app.db.session import SessionLocal, engine
from app.models.entities import (
    AcmLimitConfig,
    AcmRateVersion,
    AcmRateVersionAudit,
    AcmRateVersionDetail,
    ActivitySector,
    Agency,
    Agent,
    CategorySectorMapping,
    DailyMetric,
    ImportBatch,
    LoanRaw,
    LoginAudit,
    ParReductionTarget,
    SupportActionAudit,
    TaegDailySnapshot,
    Target,
    User,
)
from app.models.enums import ImportBatchType, TargetType, UserRole
from app.services.taeg_history import (
    create_acm_rate_version,
    create_taeg_daily_snapshot_for_batch,
    delete_taeg_daily_snapshots_for_period,
    get_applicable_acm_block,
    resolve_acm_rate_map,
    update_acm_rate_version,
)
from app.services.user_accounts import ensure_agent_accounts
from app.schemas.imports import LoanImportRow
from app.services.imports import _upsert_snapshot_rows


class SecurityControlsTestCase(unittest.TestCase):
    super_admin_email = "super.admin@microcred.com.tn"
    super_admin_password = "UltraSecureAdmin!2026"

    def setUp(self) -> None:
        Base.metadata.drop_all(bind=engine)
        Base.metadata.create_all(bind=engine)
        add_missing_columns(engine)
        self.client = TestClient(app)

    def tearDown(self) -> None:
        self.client.close()

    def _create_user(
        self,
        *,
        email: str,
        password: str,
        role: UserRole,
        full_name: str,
        agency_id: int | None = None,
        agent_id: int | None = None,
        is_active: bool = True,
        must_change_password: bool = False,
        password_changed_at: datetime | None = None,
    ) -> int:
        with SessionLocal() as db:
            user = User(
                email=email,
                full_name=full_name,
                hashed_password=hash_password(password),
                role=role,
                agency_id=agency_id,
                agent_id=agent_id,
                is_active=is_active,
                must_change_password=must_change_password,
                password_changed_at=password_changed_at or datetime.now(timezone.utc),
            )
            db.add(user)
            db.commit()
            db.refresh(user)
            return user.id

    def _seed_super_admin(self) -> None:
        self._create_user(
            email=self.super_admin_email,
            password=self.super_admin_password,
            role=UserRole.SUPER_ADMIN,
            full_name="Security Super Admin",
        )

    def _login(self, email: str, password: str):
        return self.client.post(
            "/api/v1/auth/login",
            json={"email": email, "password": password},
        )

    def _build_current_state_import_file(self, *, contract_no: str = "IMPORT-001", snapshot: date = date(2026, 7, 31)) -> bytes:
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(
            [
                "CONTRACT_NO",
                "CLIENT_NO",
                "BRANCH",
                "DAO_NAME",
                "DISBURSEMENT_AMOUNT",
                "PRINCIPAL_OUTSTANDING",
                "TOTAL_PRINCIPAL_DUE_AMT",
                "TOTAL_CUR_NO_OF_DAYS_OVERDUE",
                "TOTAL_DUE_AMT",
                "DISBURSEMENT_DATE",
                "DateEOD",
                "CATEGORY_DESC",
            ]
        )
        sheet.append(
            [
                contract_no,
                "CLI-001",
                "Agence Import",
                "Agent Import",
                1500,
                1000,
                100,
                0,
                100,
                date(2026, 7, 15),
                snapshot,
                "Micro",
            ]
        )
        buffer = BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    def _build_restructured_contracts_import_file(self, *, contract_no: str = "R-0001") -> bytes:
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(
            [
                "Numero contrat",
                "Colonne1",
                "N° contrat",
                "date de décalage",
                "dans MCR",
                "CATEGORY_DESC",
                "Total_Due",
                "LOAN_DURATION",
            ]
        )
        sheet.append(
            [
                contract_no,
                "Restructure",
                contract_no,
                date(2026, 1, 15),
                "Oui",
                "Credits Restructurés",
                1250,
                12,
            ]
        )
        buffer = BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    def test_login_uses_httponly_cookies_and_passwords_are_not_exposed(self) -> None:
        self._seed_super_admin()

        login_response = self._login(self.super_admin_email, self.super_admin_password)
        self.assertEqual(login_response.status_code, 200, login_response.text)
        payload = login_response.json()
        self.assertNotIn("access_token", payload)
        self.assertNotIn("refresh_token", payload)

        set_cookie_header = login_response.headers.get("set-cookie", "")
        self.assertIn("microcred_access=", set_cookie_header)
        self.assertIn("microcred_refresh=", set_cookie_header)
        self.assertIn("HttpOnly", set_cookie_header)
        self.assertIn("SameSite=Strict", set_cookie_header)

        create_response = self.client.post(
            "/api/v1/users",
            json={
                "email": "new.agent@microcred.com.tn",
                "full_name": "New Agent",
                "role": "portfolio_manager",
                "agency_id": None,
                "agent_id": None,
            },
        )
        self.assertEqual(create_response.status_code, 400, create_response.text)

        with SessionLocal() as db:
            agency = Agency(name="Agence Test")
            db.add(agency)
            db.flush()
            agent = Agent(name="Agent Test", agency_id=agency.id)
            db.add(agent)
            db.commit()
            agency_id = agency.id
            agent_id = agent.id

        create_response = self.client.post(
            "/api/v1/users",
            json={
                "email": "new.agent@microcred.com.tn",
                "full_name": "New Agent",
                "role": "portfolio_manager",
                "agency_id": agency_id,
                "agent_id": agent_id,
            },
        )
        self.assertEqual(create_response.status_code, 200, create_response.text)
        create_body = create_response.json()
        self.assertIn("temporary_password", create_body)
        self.assertTrue(create_body["generated_password"])
        self.assertNotIn("initial_password", create_body)
        self.assertNotIn("hashed_password", create_body)

        temporary_password = create_body["temporary_password"]
        users_response = self.client.get("/api/v1/users?limit=500")
        self.assertEqual(users_response.status_code, 200, users_response.text)
        listed_user = next(item for item in users_response.json()["items"] if item["email"] == "new.agent@microcred.com.tn")
        self.assertEqual(listed_user["temporary_password"], temporary_password)

        new_login = self._login("new.agent@microcred.com.tn", temporary_password)
        self.assertEqual(new_login.status_code, 200, new_login.text)
        self.assertTrue(new_login.json()["user"]["must_change_password"])

        blocked_response = self.client.get("/api/v1/lookups/features")
        self.assertEqual(blocked_response.status_code, 403, blocked_response.text)

        change_response = self.client.post(
            "/api/v1/auth/change-password",
            json={"current_password": temporary_password, "new_password": "AgentActivation!2026"},
        )
        self.assertEqual(change_response.status_code, 200, change_response.text)
        self.assertFalse(change_response.json()["must_change_password"])
        self.client.post("/api/v1/auth/logout")

        self._login(self.super_admin_email, self.super_admin_password)
        users_response = self.client.get("/api/v1/users?limit=500")
        self.assertEqual(users_response.status_code, 200, users_response.text)
        listed_user = next(item for item in users_response.json()["items"] if item["email"] == "new.agent@microcred.com.tn")
        self.assertIsNone(listed_user["temporary_password"])

    def test_logout_invalidates_existing_access_token(self) -> None:
        self._seed_super_admin()

        login_response = self._login(self.super_admin_email, self.super_admin_password)
        self.assertEqual(login_response.status_code, 200, login_response.text)

        access_cookie = self.client.cookies.get("microcred_access")
        self.assertTrue(access_cookie)

        me_response = self.client.get("/api/v1/auth/me")
        self.assertEqual(me_response.status_code, 200, me_response.text)

        logout_response = self.client.post("/api/v1/auth/logout")
        self.assertEqual(logout_response.status_code, 200, logout_response.text)

        reused_token_response = self.client.get(
            "/api/v1/auth/me",
            headers={"Authorization": f"Bearer {access_cookie}"},
        )
        self.assertEqual(reused_token_response.status_code, 401, reused_token_response.text)
        self.assertEqual(reused_token_response.json()["detail"]["code"], "revoked_token")

    def test_required_password_change_blocks_business_routes_until_completed(self) -> None:
        initial_password = "InitialAdmin!2026"
        new_password = "ChangedAdmin!2026"
        self._create_user(
            email="readonly.admin@microcred.com.tn",
            password=initial_password,
            role=UserRole.ADMIN,
            full_name="Read Only Admin",
            must_change_password=True,
        )

        login_response = self._login("readonly.admin@microcred.com.tn", initial_password)
        self.assertEqual(login_response.status_code, 200, login_response.text)
        self.assertTrue(login_response.json()["user"]["must_change_password"])

        me_response = self.client.get("/api/v1/auth/me")
        self.assertEqual(me_response.status_code, 200, me_response.text)

        blocked_response = self.client.get("/api/v1/lookups/features")
        self.assertEqual(blocked_response.status_code, 403, blocked_response.text)
        self.assertEqual(blocked_response.json()["detail"]["code"], "password_change_required")

        wrong_current_password = self.client.post(
            "/api/v1/auth/change-password",
            json={"current_password": "WrongPassword!2026", "new_password": new_password},
        )
        self.assertEqual(wrong_current_password.status_code, 400, wrong_current_password.text)
        self.assertEqual(wrong_current_password.json()["detail"]["code"], "invalid_current_password")

        weak_password = self.client.post(
            "/api/v1/auth/change-password",
            json={"current_password": initial_password, "new_password": "lowercaseonly123!"},
        )
        self.assertEqual(weak_password.status_code, 400, weak_password.text)
        self.assertEqual(weak_password.json()["detail"]["code"], "weak_password")

        change_response = self.client.post(
            "/api/v1/auth/change-password",
            json={"current_password": initial_password, "new_password": new_password},
        )
        self.assertEqual(change_response.status_code, 200, change_response.text)
        self.assertFalse(change_response.json()["must_change_password"])

        allowed_response = self.client.get("/api/v1/lookups/features")
        self.assertEqual(allowed_response.status_code, 200, allowed_response.text)

        self.client.post("/api/v1/auth/logout")
        old_password_login = self._login("readonly.admin@microcred.com.tn", initial_password)
        self.assertEqual(old_password_login.status_code, 401, old_password_login.text)
        new_password_login = self._login("readonly.admin@microcred.com.tn", new_password)
        self.assertEqual(new_password_login.status_code, 200, new_password_login.text)

    def test_super_admin_creates_and_resets_accounts_with_forced_password_change(self) -> None:
        self._seed_super_admin()
        login_response = self._login(self.super_admin_email, self.super_admin_password)
        self.assertEqual(login_response.status_code, 200, login_response.text)

        initial_password = "RegionalAdmin!2026"
        create_response = self.client.post(
            "/api/v1/users",
            json={
                "email": "regional.admin@microcred.com.tn",
                "full_name": "Regional Admin",
                "role": "admin",
                "initial_password": initial_password,
            },
        )
        self.assertEqual(create_response.status_code, 200, create_response.text)
        created_body = create_response.json()
        self.assertEqual(created_body["temporary_password"], initial_password)
        self.assertFalse(created_body["generated_password"])
        self.assertTrue(created_body["user"]["must_change_password"])
        user_id = created_body["user"]["id"]

        self.client.post("/api/v1/auth/logout")
        first_login = self._login("regional.admin@microcred.com.tn", initial_password)
        self.assertEqual(first_login.status_code, 200, first_login.text)
        self.assertTrue(first_login.json()["user"]["must_change_password"])
        blocked_response = self.client.get("/api/v1/lookups/features")
        self.assertEqual(blocked_response.status_code, 403, blocked_response.text)

        change_response = self.client.post(
            "/api/v1/auth/change-password",
            json={"current_password": initial_password, "new_password": "RegionalChanged!2026"},
        )
        self.assertEqual(change_response.status_code, 200, change_response.text)
        self.assertFalse(change_response.json()["must_change_password"])

        self.client.post("/api/v1/auth/logout")
        login_response = self._login(self.super_admin_email, self.super_admin_password)
        self.assertEqual(login_response.status_code, 200, login_response.text)

        reset_response = self.client.post(f"/api/v1/users/{user_id}/reset-password")
        self.assertEqual(reset_response.status_code, 200, reset_response.text)
        reset_body = reset_response.json()
        reset_password = reset_body["temporary_password"]
        self.assertTrue(reset_body["generated_password"])
        self.assertTrue(reset_body["user"]["must_change_password"])

        users_response = self.client.get("/api/v1/users?limit=500")
        self.assertEqual(users_response.status_code, 200, users_response.text)
        listed_user = next(item for item in users_response.json()["items"] if item["id"] == user_id)
        self.assertEqual(listed_user["temporary_password"], reset_password)

        self.client.post("/api/v1/auth/logout")
        old_password_login = self._login("regional.admin@microcred.com.tn", "RegionalChanged!2026")
        self.assertEqual(old_password_login.status_code, 401, old_password_login.text)
        reset_login = self._login("regional.admin@microcred.com.tn", reset_password)
        self.assertEqual(reset_login.status_code, 200, reset_login.text)
        self.assertTrue(reset_login.json()["user"]["must_change_password"])

    def test_successful_login_creates_audit_entry(self) -> None:
        self._seed_super_admin()
        login_response = self._login(self.super_admin_email, self.super_admin_password)
        self.assertEqual(login_response.status_code, 200, login_response.text)

        with SessionLocal() as db:
            audits = db.scalars(select(LoginAudit).order_by(LoginAudit.id.desc())).all()
            self.assertEqual(len(audits), 1)
            self.assertEqual(audits[0].email, self.super_admin_email)

        audit_response = self.client.get("/api/v1/users/login-audit")
        self.assertEqual(audit_response.status_code, 200, audit_response.text)
        self.assertEqual(audit_response.json()["total_connections"], 1)

    def test_support_role_has_limited_scope_and_forced_password_change(self) -> None:
        self._seed_super_admin()
        self.assertEqual(self._login(self.super_admin_email, self.super_admin_password).status_code, 200)

        create_response = self.client.post(
            "/api/v1/users",
            json={
                "email": "support.agent@microcred.com.tn",
                "full_name": "Support Agent",
                "role": "support",
                "initial_password": "SupportInitial!2026",
            },
        )
        self.assertEqual(create_response.status_code, 200, create_response.text)
        body = create_response.json()
        self.assertEqual(body["user"]["role"], "support")
        self.assertTrue(body["user"]["must_change_password"])

        self.client.post("/api/v1/auth/logout")
        first_login = self._login("support.agent@microcred.com.tn", "SupportInitial!2026")
        self.assertEqual(first_login.status_code, 200, first_login.text)
        self.assertTrue(first_login.json()["user"]["must_change_password"])

        blocked_before_change = self.client.get("/api/v1/imports/batches")
        self.assertEqual(blocked_before_change.status_code, 403, blocked_before_change.text)
        self.assertEqual(blocked_before_change.json()["detail"]["code"], "password_change_required")

        change_response = self.client.post(
            "/api/v1/auth/change-password",
            json={"current_password": "SupportInitial!2026", "new_password": "SupportChanged!2026"},
        )
        self.assertEqual(change_response.status_code, 200, change_response.text)
        self.assertFalse(change_response.json()["must_change_password"])

        allowed_imports = self.client.get("/api/v1/imports/batches")
        self.assertEqual(allowed_imports.status_code, 200, allowed_imports.text)
        allowed_users = self.client.get("/api/v1/users?limit=500")
        self.assertEqual(allowed_users.status_code, 200, allowed_users.text)

        denied_metrics = self.client.get("/api/v1/metrics/summary")
        self.assertEqual(denied_metrics.status_code, 403, denied_metrics.text)
        denied_taeg = self.client.get("/api/v1/taeg/periods")
        self.assertEqual(denied_taeg.status_code, 403, denied_taeg.text)
        denied_lookups = self.client.get("/api/v1/lookups/features")
        self.assertEqual(denied_lookups.status_code, 403, denied_lookups.text)

    def test_committee_member_closed_months_uses_imported_snapshot_months_before_current_month(self) -> None:
        password = "CommitteeViewer!2026"
        self._create_user(
            email="committee.member@microcred.com.tn",
            password=password,
            role=UserRole.COMMITTEE_MEMBER,
            full_name="Committee Member",
        )
        login_response = self._login("committee.member@microcred.com.tn", password)
        self.assertEqual(login_response.status_code, 200, login_response.text)

        current_month_start = date.today().replace(day=1)
        previous_month_end = current_month_start - timedelta(days=1)
        previous_month_start = previous_month_end.replace(day=1)
        two_months_back_end = previous_month_start - timedelta(days=1)
        two_months_back_start = two_months_back_end.replace(day=1)
        current_month_snapshot = min(date.today(), current_month_start + timedelta(days=10))

        with SessionLocal() as db:
            previous_historical = ImportBatch(
                batch_type=ImportBatchType.HISTORICAL_MONTH,
                period=f"{previous_month_end.year:04d}-{previous_month_end.month:02d}",
                snapshot_date=previous_month_end,
                file_name="historical-previous.xlsx",
            )
            previous_snapshot = ImportBatch(
                batch_type=ImportBatchType.SNAPSHOT,
                period=None,
                snapshot_date=previous_month_end,
                file_name="snapshot-previous.xlsx",
            )
            two_months_historical = ImportBatch(
                batch_type=ImportBatchType.HISTORICAL_MONTH,
                period=f"{two_months_back_end.year:04d}-{two_months_back_end.month:02d}",
                snapshot_date=two_months_back_end,
                file_name="historical-two-months-back.xlsx",
            )
            current_state_batch = ImportBatch(
                batch_type=ImportBatchType.CURRENT_STATE,
                period=None,
                snapshot_date=current_month_snapshot,
                file_name="current-state-current-month.xlsx",
            )
            db.add_all(
                [
                    previous_historical,
                    previous_snapshot,
                    two_months_historical,
                    current_state_batch,
                ]
            )
            db.commit()

        response = self.client.get("/api/v1/metrics/closed-months")
        self.assertEqual(response.status_code, 200, response.text)
        payload = response.json()

        expected_previous_key = f"{previous_month_start.year:04d}-{previous_month_start.month:02d}"
        expected_two_months_key = f"{two_months_back_start.year:04d}-{two_months_back_start.month:02d}"
        current_month_key = f"{current_month_start.year:04d}-{current_month_start.month:02d}"

        self.assertEqual([item["key"] for item in payload], [expected_previous_key, expected_two_months_key])
        self.assertEqual(payload[0]["batch_id"], 1)
        self.assertEqual(payload[0]["period"], expected_previous_key)
        self.assertEqual(payload[1]["period"], expected_two_months_key)
        self.assertNotIn(current_month_key, [item["key"] for item in payload])

    def test_committee_member_can_access_dashboard_lookups_but_not_current_state(self) -> None:
        password = "CommitteeLookup!2026"
        self._create_user(
            email="committee.lookup@microcred.com.tn",
            password=password,
            role=UserRole.COMMITTEE_MEMBER,
            full_name="Committee Lookup",
        )
        with SessionLocal() as db:
            agency = Agency(name="Agence Comite")
            db.add(agency)
            db.flush()
            db.add(Agent(name="Agent Comite", agency_id=agency.id))
            db.commit()

        login_response = self._login("committee.lookup@microcred.com.tn", password)
        self.assertEqual(login_response.status_code, 200, login_response.text)

        agencies_response = self.client.get("/api/v1/lookups/agencies?limit=500")
        self.assertEqual(agencies_response.status_code, 200, agencies_response.text)
        self.assertEqual(agencies_response.json()["items"][0]["name"], "Agence Comite")

        agents_response = self.client.get("/api/v1/lookups/agents?limit=500")
        self.assertEqual(agents_response.status_code, 200, agents_response.text)
        self.assertEqual(agents_response.json()["items"][0]["name"], "Agent Comite")

        features_response = self.client.get("/api/v1/lookups/features")
        self.assertEqual(features_response.status_code, 200, features_response.text)

        active_snapshot_response = self.client.get("/api/v1/metrics/active-snapshot")
        self.assertEqual(active_snapshot_response.status_code, 403, active_snapshot_response.text)

    def test_support_can_manage_users_without_plaintext_password_access(self) -> None:
        self._seed_super_admin()
        support_id = self._create_user(
            email="support.operator@microcred.com.tn",
            password="SupportChanged!2026",
            role=UserRole.SUPPORT,
            full_name="Support Operator",
        )
        admin_id = self._create_user(
            email="regional.admin@microcred.com.tn",
            password="RegionalAdmin!2026",
            role=UserRole.ADMIN,
            full_name="Regional Admin",
        )

        login_response = self._login("support.operator@microcred.com.tn", "SupportChanged!2026")
        self.assertEqual(login_response.status_code, 200, login_response.text)

        users_response = self.client.get("/api/v1/users?limit=500")
        self.assertEqual(users_response.status_code, 200, users_response.text)
        listed_support = next(item for item in users_response.json()["items"] if item["id"] == support_id)
        self.assertIsNone(listed_support["temporary_password"])

        create_response = self.client.post(
            "/api/v1/users",
            json={
                "email": "new.manager@microcred.com.tn",
                "full_name": "New Manager",
                "role": "admin",
                "initial_password": "NewManager!2026",
            },
        )
        self.assertEqual(create_response.status_code, 200, create_response.text)
        created_user = create_response.json()["user"]
        self.assertIsNone(create_response.json()["temporary_password"])

        update_response = self.client.put(
            f"/api/v1/users/{created_user['id']}",
            json={
                "full_name": "Updated Manager",
                "role": "support",
            },
        )
        self.assertEqual(update_response.status_code, 200, update_response.text)
        self.assertEqual(update_response.json()["role"], "support")
        self.assertEqual(update_response.json()["full_name"], "Updated Manager")

        status_response = self.client.patch(
            f"/api/v1/users/{admin_id}/status",
            json={"is_active": False},
        )
        self.assertEqual(status_response.status_code, 200, status_response.text)
        self.assertFalse(status_response.json()["is_active"])

        reset_response = self.client.post(f"/api/v1/users/{admin_id}/reset-password", json={})
        self.assertEqual(reset_response.status_code, 200, reset_response.text)
        self.assertIsNone(reset_response.json()["temporary_password"])
        self.assertTrue(reset_response.json()["user"]["must_change_password"])

        delete_response = self.client.delete(f"/api/v1/users/{created_user['id']}")
        self.assertEqual(delete_response.status_code, 200, delete_response.text)

        protected_response = self.client.post("/api/v1/users/1/reset-password", json={})
        self.assertEqual(protected_response.status_code, 403, protected_response.text)

        super_admin_role_forbidden = self.client.post(
            "/api/v1/users",
            json={
                "email": "forbidden@microcred.com.tn",
                "full_name": "Forbidden",
                "role": "super_admin",
                "initial_password": "ForbiddenAdmin!2026",
            },
        )
        self.assertEqual(super_admin_role_forbidden.status_code, 403, super_admin_role_forbidden.text)

        self_role_change_response = self.client.put(
            f"/api/v1/users/{support_id}",
            json={"role": "admin"},
        )
        self.assertEqual(self_role_change_response.status_code, 400, self_role_change_response.text)

        self_delete_response = self.client.delete(f"/api/v1/users/{support_id}")
        self.assertEqual(self_delete_response.status_code, 400, self_delete_response.text)

        with SessionLocal() as db:
            audit_rows = db.scalars(
                select(SupportActionAudit).where(SupportActionAudit.actor_user_id == support_id)
            ).all()
            actions = {row.action for row in audit_rows}
            self.assertIn("user_create", actions)
            self.assertIn("user_update", actions)
            self.assertIn("user_status_update", actions)
            self.assertIn("password_reset", actions)
            self.assertIn("user_delete", actions)

    def test_support_can_access_users_audits_but_cannot_manage_super_admin_accounts(self) -> None:
        self._seed_super_admin()
        self._create_user(
            email="support.audit@microcred.com.tn",
            password="SupportChanged!2026",
            role=UserRole.SUPPORT,
            full_name="Support Audit",
        )

        login_response = self._login("support.audit@microcred.com.tn", "SupportChanged!2026")
        self.assertEqual(login_response.status_code, 200, login_response.text)

        gp_audit_response = self.client.get("/api/v1/users/gp-mcr-audit")
        self.assertEqual(gp_audit_response.status_code, 200, gp_audit_response.text)

        login_audit_response = self.client.get("/api/v1/users/login-audit")
        self.assertEqual(login_audit_response.status_code, 200, login_audit_response.text)
        self.assertGreaterEqual(login_audit_response.json()["total_connections"], 1)

        excel_export = self.client.get("/api/v1/users/login-audit.xlsx")
        self.assertEqual(excel_export.status_code, 200, excel_export.text)

        pdf_export = self.client.get("/api/v1/users/login-audit.pdf")
        self.assertEqual(pdf_export.status_code, 200, pdf_export.text)

        protected_update = self.client.put(
            "/api/v1/users/1",
            json={"full_name": "Blocked Super Admin"},
        )
        self.assertEqual(protected_update.status_code, 403, protected_update.text)

        protected_status = self.client.patch(
            "/api/v1/users/1/status",
            json={"is_active": False},
        )
        self.assertEqual(protected_status.status_code, 403, protected_status.text)

        protected_delete = self.client.delete("/api/v1/users/1")
        self.assertEqual(protected_delete.status_code, 403, protected_delete.text)

    def test_support_can_preview_and_import_restructured_contracts_without_response_shape_regression(self) -> None:
        self._seed_super_admin()
        self._create_user(
            email="support.import@microcred.com.tn",
            password="SupportChanged!2026",
            role=UserRole.SUPPORT,
            full_name="Support Import",
        )

        login_response = self._login("support.import@microcred.com.tn", "SupportChanged!2026")
        self.assertEqual(login_response.status_code, 200, login_response.text)

        content = self._build_restructured_contracts_import_file(contract_no="RST-IMPORT-001")
        files = {
            "file": (
                "liste_restructures.xlsx",
                content,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        }

        preview_response = self.client.post("/api/v1/imports/restructured-contracts/preview", files=files)
        self.assertEqual(preview_response.status_code, 200, preview_response.text)
        preview_payload = preview_response.json()
        self.assertEqual(preview_payload["file_name"], "liste_restructures.xlsx")
        self.assertEqual(preview_payload["rows_seen"], 1)
        self.assertIn("existing_contracts_count", preview_payload)
        self.assertIn("imported_list_count", preview_payload)
        self.assertIn("detected_in_mcr_count", preview_payload)
        self.assertIn("kept_count", preview_payload)
        self.assertIn("add_count", preview_payload)
        self.assertIn("update_count", preview_payload)
        self.assertIn("delete_count", preview_payload)
        self.assertIn("pending_cleanup_count", preview_payload)
        self.assertIsInstance(preview_payload.get("deletions"), list)
        self.assertIn("message", preview_payload)

        import_response = self.client.post("/api/v1/imports/restructured-contracts", files=files)
        self.assertEqual(import_response.status_code, 200, import_response.text)
        import_payload = import_response.json()
        self.assertEqual(import_payload["rows_seen"], 1)
        self.assertEqual(import_payload["inserted"], 1)
        self.assertEqual(import_payload["updated"], 0)
        self.assertIn("deleted", import_payload)
        self.assertIn("kept_count", import_payload)
        self.assertIn("pending_cleanup_count", import_payload)
        self.assertIn("recalculated_contracts", import_payload)
        self.assertIn("import_log_id", import_payload)
        self.assertIn("message", import_payload)

    def test_current_state_import_endpoint_returns_real_import_result_fields_and_is_safe_on_retry(self) -> None:
        self._seed_super_admin()
        login_response = self._login(self.super_admin_email, self.super_admin_password)
        self.assertEqual(login_response.status_code, 200, login_response.text)

        content = self._build_current_state_import_file(contract_no="IMPORT-CURRENT-001", snapshot=date(2026, 7, 31))
        files = {
            "file": (
                "current-state.xlsx",
                content,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        }

        first_response = self.client.post("/api/v1/imports/current-state", files=files)
        self.assertEqual(first_response.status_code, 200, first_response.text)
        first_payload = first_response.json()
        self.assertEqual(first_payload["rows_seen"], 1)
        self.assertEqual(first_payload["snapshot_date"], "2026-07-31")
        self.assertEqual(first_payload["inserted"], 1)
        self.assertIn("import_batch", first_payload)

        second_response = self.client.post("/api/v1/imports/current-state", files=files)
        self.assertEqual(second_response.status_code, 200, second_response.text)
        second_payload = second_response.json()
        self.assertEqual(second_payload["rows_seen"], 1)
        self.assertEqual(second_payload["snapshot_date"], "2026-07-31")
        self.assertEqual(second_payload["inserted"], 1)

        with SessionLocal() as db:
            current_batches = db.scalars(
                select(ImportBatch).where(ImportBatch.batch_type == ImportBatchType.CURRENT_STATE)
            ).all()
            all_rows = db.scalars(select(LoanRaw).where(LoanRaw.contract_no == "IMPORT-CURRENT-001")).all()
            self.assertEqual(len(current_batches), 1)
            self.assertEqual(len(all_rows), 1)
            self.assertEqual(current_batches[0].snapshot_date, date(2026, 7, 31))

    def test_super_admin_can_configure_acm_limits(self) -> None:
        self._seed_super_admin()
        login_response = self._login(self.super_admin_email, self.super_admin_password)
        self.assertEqual(login_response.status_code, 200, login_response.text)

        put_response = self.client.put(
            "/api/v1/users/acm-limits",
            json={
                "par_0_limit": 5,
                "par_30_limit": 4.25,
                "par_120_limit": 1.1,
                "cohort_1_30_limit": 2.0,
                "cohort_31_60_limit": 1.5,
                "cohort_61_90_limit": 0.9,
                "cohort_91_120_limit": 0.5,
            },
        )
        self.assertEqual(put_response.status_code, 200, put_response.text)
        self.assertTrue(put_response.json()["configured"])

        get_response = self.client.get("/api/v1/users/acm-limits")
        self.assertEqual(get_response.status_code, 200, get_response.text)
        self.assertEqual(str(get_response.json()["par_0_limit"]), "5.00")

    def test_metrics_chart_exposes_segmented_disbursement_categories(self) -> None:
        self._seed_super_admin()
        with SessionLocal() as db:
            agency = Agency(name="Agence Segments")
            db.add(agency)
            db.flush()
            agent = Agent(name="Agent Segments", agency_id=agency.id)
            db.add(agent)
            db.flush()
            batch = ImportBatch(
                batch_type=ImportBatchType.CURRENT_STATE,
                period=None,
                snapshot_date=date(2026, 6, 30),
                file_name="segments.xlsx",
            )
            db.add(batch)
            db.flush()
            common = dict(
                agency_name=agency.name,
                agent_name=agent.name,
                client_name="Client",
                client_first_name="Test",
                client_rating="Classe 0",
                principal_outstanding=1000,
                principal_due=0,
                total_scheduled_amount=0,
                days_overdue=0,
                total_due=0,
                status="active",
                snapshot_date=date(2026, 6, 30),
                import_batch_id=batch.id,
                maturity_date=None,
                next_schedule_date=None,
            )
            db.add_all([
                LoanRaw(contract_no="A1", client_id="1", category_desc="Afari services", disbursement_amount=15000, disbursement_date=date(2026, 6, 2), **common),
                LoanRaw(contract_no="A2", client_id="2", category_desc="TPME", disbursement_amount=25000, disbursement_date=date(2026, 6, 3), **common),
                LoanRaw(contract_no="A3", client_id="3", category_desc="Micro", disbursement_amount=5000, disbursement_date=date(2026, 6, 4), **common),
            ])
            db.commit()

        login_response = self._login(self.super_admin_email, self.super_admin_password)
        self.assertEqual(login_response.status_code, 200, login_response.text)
        chart_response = self.client.get("/api/v1/metrics/charts")
        self.assertEqual(chart_response.status_code, 200, chart_response.text)
        point = chart_response.json()["points"][0]
        self.assertEqual(float(point["afari_disbursement_volume"]), 15000.0)
        self.assertEqual(point["afari_disbursement_count"], 1)
        self.assertEqual(float(point["tpme_disbursement_volume"]), 25000.0)
        self.assertEqual(point["tpme_disbursement_count"], 1)
        self.assertEqual(float(point["micro_disbursement_volume"]), 5000.0)
        self.assertEqual(point["micro_disbursement_count"], 1)

    def test_metrics_chart_exposes_segmented_quality_cohorts(self) -> None:
        self._seed_super_admin()
        with SessionLocal() as db:
            agency = Agency(name="Agence Cohortes")
            db.add(agency)
            db.flush()
            agent = Agent(name="Agent Cohortes", agency_id=agency.id)
            db.add(agent)
            db.flush()
            batch = ImportBatch(
                batch_type=ImportBatchType.CURRENT_STATE,
                period=None,
                snapshot_date=date(2026, 6, 30),
                file_name="quality-cohorts.xlsx",
            )
            db.add(batch)
            db.flush()
            common = dict(
                agency_name=agency.name,
                agent_name=agent.name,
                client_name="Client",
                client_first_name="Test",
                client_rating="Classe 0",
                principal_outstanding=0,
                principal_due=0,
                total_scheduled_amount=0,
                days_overdue=0,
                total_due=0,
                status="active",
                snapshot_date=date(2026, 6, 30),
                import_batch_id=batch.id,
                maturity_date=None,
                next_schedule_date=None,
            )
            db.add_all([
                LoanRaw(contract_no="Q1", client_id="1", days_overdue=5, disbursement_amount=5000, disbursement_date=date(2026, 6, 2), principal_outstanding=500, principal_due=0, **common),
                LoanRaw(contract_no="Q2", client_id="2", days_overdue=20, disbursement_amount=6000, disbursement_date=date(2026, 6, 3), principal_outstanding=750, principal_due=0, **common),
                LoanRaw(contract_no="Q3", client_id="3", days_overdue=45, disbursement_amount=7000, disbursement_date=date(2026, 6, 4), principal_outstanding=350, principal_due=0, **common),
            ])
            db.commit()

        login_response = self._login(self.super_admin_email, self.super_admin_password)
        self.assertEqual(login_response.status_code, 200, login_response.text)

        chart_response = self.client.get("/api/v1/metrics/charts")
        self.assertEqual(chart_response.status_code, 200, chart_response.text)
        point = next(item for item in chart_response.json()["points"] if item["label"] == "Agence Cohortes")
        self.assertEqual(float(point["par_1_15"]), 500.0)
        self.assertEqual(float(point["par_16_30"]), 750.0)
        self.assertEqual(float(point["par_31_60"]), 350.0)

    def test_admin_can_create_and_update_par_reduction_target_to_reach(self) -> None:
        self._seed_super_admin()
        with SessionLocal() as db:
            agency = Agency(name="Agence Objectif PAR")
            db.add(agency)
            db.flush()
            agent = Agent(name="Agent Objectif PAR", agency_id=agency.id)
            db.add(agent)
            db.commit()
            agency_id = agency.id
            agent_id = agent.id

        login_response = self._login(self.super_admin_email, self.super_admin_password)
        self.assertEqual(login_response.status_code, 200, login_response.text)
        payload = {
            "agency_id": agency_id,
            "agent_id": agent_id,
            "month": 8,
            "year": 2026,
            "target_par30": 30000,
            "target_cohort_1_15": 2000,
            "target_cohort_16_30": 10000,
        }

        create_response = self.client.post("/api/v1/par-reduction-targets", json=payload)
        self.assertEqual(create_response.status_code, 200, create_response.text)
        created = create_response.json()
        self.assertEqual(float(created["target_par30"]), 30000.0)
        target_id = created["id"]

        with SessionLocal() as db:
            target = db.get(ParReductionTarget, target_id)
            self.assertIsNotNone(target)
            self.assertEqual(target.target_par30, Decimal("30000.00"))
            self.assertIsNone(getattr(target, "legacy_target_par30_reduction", None))

        update_payload = {**payload, "target_par30": 25000}
        update_response = self.client.put(f"/api/v1/par-reduction-targets/{target_id}", json=update_payload)
        self.assertEqual(update_response.status_code, 200, update_response.text)
        self.assertEqual(float(update_response.json()["target_par30"]), 25000.0)

    def test_super_admin_can_delete_user_account_and_related_references(self) -> None:
        self._seed_super_admin()
        login_response = self._login(self.super_admin_email, self.super_admin_password)
        self.assertEqual(login_response.status_code, 200, login_response.text)

        create_response = self.client.post(
            "/api/v1/users",
            json={
                "email": "delete.me@microcred.com.tn",
                "full_name": "Delete Me",
                "role": "admin",
                "initial_password": "DeleteMeAdmin!2026",
            },
        )
        self.assertEqual(create_response.status_code, 200, create_response.text)
        user_id = create_response.json()["user"]["id"]

        self_delete_response = self.client.delete(f"/api/v1/users/{self.client.get('/api/v1/auth/me').json()['id']}")
        self.assertEqual(self_delete_response.status_code, 400, self_delete_response.text)
        self.assertEqual(self_delete_response.json()["detail"]["code"], "self_delete_forbidden")

        delete_response = self.client.delete(f"/api/v1/users/{user_id}")
        self.assertEqual(delete_response.status_code, 200, delete_response.text)

        users_response = self.client.get("/api/v1/users?limit=500")
        self.assertEqual(users_response.status_code, 200, users_response.text)
        self.assertFalse(any(item["id"] == user_id for item in users_response.json()["items"]))

        deleted_login = self._login("delete.me@microcred.com.tn", "DeleteMeAdmin!2026")
        self.assertEqual(deleted_login.status_code, 401, deleted_login.text)

    def test_deleting_portfolio_manager_deletes_unused_associated_agent(self) -> None:
        self._seed_super_admin()
        login_response = self._login(self.super_admin_email, self.super_admin_password)
        self.assertEqual(login_response.status_code, 200, login_response.text)

        with SessionLocal() as db:
            agency = Agency(name="Agence Delete GP")
            db.add(agency)
            db.flush()
            agent = Agent(name="Delete GP", agency_id=agency.id)
            db.add(agent)
            db.flush()
            agency_id = agency.id
            agent_id = agent.id

            db.add(
                DailyMetric(
                    agency_id=agency_id,
                    agent_id=agent_id,
                    date=date(2026, 6, 1),
                    disbursement_count=1,
                    disbursement_volume=1000,
                    nb_clients=1,
                    outstanding=900,
                    healthy_outstanding=900,
                    par_0=0,
                    par_1_30=0,
                    par_31_60=0,
                    par_30=0,
                )
            )
            db.add(
                Target(
                    target_type=TargetType.AGENT,
                    agency_id=agency_id,
                    agent_id=agent_id,
                    month=6,
                    year=2026,
                    target_disbursement_count=1,
                )
            )
            db.commit()

        create_response = self.client.post(
            "/api/v1/users",
            json={
                "email": "delete.gp@microcred.com.tn",
                "full_name": "Delete GP",
                "role": "portfolio_manager",
                "agency_id": agency_id,
                "agent_id": agent_id,
            },
        )
        self.assertEqual(create_response.status_code, 200, create_response.text)
        user_id = create_response.json()["user"]["id"]

        delete_response = self.client.delete(f"/api/v1/users/{user_id}")
        self.assertEqual(delete_response.status_code, 200, delete_response.text)

        with SessionLocal() as db:
            self.assertIsNone(db.get(User, user_id))
            self.assertIsNone(db.get(Agent, agent_id))
            self.assertEqual(db.query(DailyMetric).filter(DailyMetric.agent_id == agent_id).count(), 0)
            self.assertEqual(db.query(Target).filter(Target.agent_id == agent_id).count(), 0)

    def test_login_rate_limiting_locks_account_after_repeated_failures(self) -> None:
        self._seed_super_admin()

        for _ in range(6):
            bad_login = self._login(self.super_admin_email, "WrongPassword!2026")
            self.assertEqual(bad_login.status_code, 401, bad_login.text)

        locked_response = self._login(self.super_admin_email, self.super_admin_password)
        self.assertEqual(locked_response.status_code, 429, locked_response.text)
        self.assertEqual(locked_response.json()["detail"]["code"], "account_locked")

        with SessionLocal() as db:
            user = db.query(User).filter(User.email == self.super_admin_email).one()
            self.assertGreaterEqual(user.failed_login_attempts, 6)
            self.assertIsNotNone(user.locked_until)

    def test_password_expires_after_40_days_and_requires_change_on_login(self) -> None:
        password = "ExpiredAdmin!2026"
        self._create_user(
            email="expired.admin@microcred.com.tn",
            password=password,
            role=UserRole.ADMIN,
            full_name="Expired Password Admin",
            password_changed_at=datetime.now(timezone.utc) - timedelta(days=41),
        )

        login_response = self._login("expired.admin@microcred.com.tn", password)
        self.assertEqual(login_response.status_code, 200, login_response.text)
        self.assertTrue(login_response.json()["user"]["must_change_password"])

        blocked_response = self.client.get("/api/v1/lookups/features")
        self.assertEqual(blocked_response.status_code, 403, blocked_response.text)
        self.assertEqual(blocked_response.json()["detail"]["code"], "password_change_required")

    def test_automatically_created_portfolio_manager_account_is_active(self) -> None:
        with SessionLocal() as db:
            agency = Agency(name="Agence Auto Account")
            db.add(agency)
            db.flush()
            agent = Agent(name="Auto Account GP", agency_id=agency.id)
            db.add(agent)
            db.commit()

            created_credentials = ensure_agent_accounts(db)
            db.commit()
            user = db.scalar(select(User).where(User.agent_id == agent.id))
            self.assertEqual(created_credentials, [])
            self.assertIsNotNone(user)
            self.assertTrue(user.is_active)
            self.assertTrue(user.must_change_password)
            self.assertIsNone(user.password_changed_at)

            provisioned = ensure_agent_accounts(db, issue_credentials=True)
            db.commit()
            db.refresh(user)
            self.assertEqual(len(provisioned), 1)
            self.assertEqual(provisioned[0]["email"], user.email)
            self.assertTrue(provisioned[0]["temporary_password"])
            self.assertTrue(user.is_active)
            self.assertTrue(user.must_change_password)
            self.assertIsNotNone(user.password_changed_at)
            self.assertEqual(user.temporary_password, provisioned[0]["temporary_password"])

            provisioned_again = ensure_agent_accounts(db, issue_credentials=True)
            db.commit()
            self.assertEqual(len(provisioned_again), 1)
            self.assertEqual(provisioned_again[0]["temporary_password"], user.temporary_password)

            user.is_active = False
            db.commit()
            created_again = ensure_agent_accounts(db)
            db.commit()
            db.refresh(user)
            self.assertEqual(created_again, [])
            self.assertFalse(user.is_active)

    def test_existing_portfolio_manager_waiting_for_first_login_is_reprovisioned(self) -> None:
        with SessionLocal() as db:
            agency = Agency(name="Agence Reprovision")
            db.add(agency)
            db.flush()
            agent = Agent(name="Reprovision GP", agency_id=agency.id)
            db.add(agent)
            db.flush()
            user = User(
                email="reprovision.gp@microcred.com.tn",
                full_name="Reprovision GP",
                hashed_password=hash_password("OldTemporary!2026"),
                role=UserRole.PORTFOLIO_MANAGER,
                agency_id=agency.id,
                agent_id=agent.id,
                is_active=True,
                must_change_password=True,
                password_changed_at=datetime.now(timezone.utc),
                temporary_password=None,
            )
            db.add(user)
            db.commit()

            provisioned = ensure_agent_accounts(db, issue_credentials=True)
            db.commit()
            db.refresh(user)

            self.assertEqual(len(provisioned), 1)
            self.assertEqual(provisioned[0]["email"], user.email)
            self.assertTrue(provisioned[0]["temporary_password"])
            self.assertEqual(user.temporary_password, provisioned[0]["temporary_password"])
            self.assertTrue(user.must_change_password)

    def test_auto_create_agents_keeps_one_account_for_gp_across_agencies(self) -> None:
        with SessionLocal() as db:
            agency_one = Agency(name="Agence Multi One")
            agency_two = Agency(name="Agence Multi Two")
            db.add_all([agency_one, agency_two])
            db.flush()
            db.add_all(
                [
                    Agent(name="WAJIH.TRABELSI", agency_id=agency_one.id),
                    Agent(name="WAJIH TRABELSI", agency_id=agency_two.id),
                ]
            )
            db.commit()

            provisioned = ensure_agent_accounts(db, issue_credentials=True)
            db.commit()

            users = db.scalars(
                select(User).where(User.role == UserRole.PORTFOLIO_MANAGER)
            ).all()
            self.assertEqual(len(provisioned), 1)
            self.assertEqual(len(users), 1)
            self.assertEqual(users[0].temporary_password, provisioned[0]["temporary_password"])
            self.assertTrue(users[0].must_change_password)

            provisioned_again = ensure_agent_accounts(db, issue_credentials=True)
            db.commit()
            users_again = db.scalars(
                select(User).where(User.role == UserRole.PORTFOLIO_MANAGER)
            ).all()
            self.assertEqual(len(provisioned_again), 1)
            self.assertEqual(len(users_again), 1)

    def test_auto_create_agents_merges_existing_duplicate_gp_accounts(self) -> None:
        with SessionLocal() as db:
            agency_one = Agency(name="Agence Duplicate One")
            agency_two = Agency(name="Agence Duplicate Two")
            db.add_all([agency_one, agency_two])
            db.flush()
            agent_one = Agent(name="Duplicate GP", agency_id=agency_one.id)
            agent_two = Agent(name="Duplicate.GP", agency_id=agency_two.id)
            db.add_all([agent_one, agent_two])
            db.flush()
            db.add_all(
                [
                    User(
                        email="duplicate.gp.one@microcred.com.tn",
                        full_name="Duplicate GP",
                        hashed_password=hash_password("DuplicateOne!2026"),
                        role=UserRole.PORTFOLIO_MANAGER,
                        agency_id=agency_one.id,
                        agent_id=agent_one.id,
                        is_active=True,
                        must_change_password=True,
                        temporary_password="DuplicateOne!2026",
                        password_changed_at=datetime.now(timezone.utc),
                    ),
                    User(
                        email="duplicate.gp.two@microcred.com.tn",
                        full_name="Duplicate GP",
                        hashed_password=hash_password("DuplicateTwo!2026"),
                        role=UserRole.PORTFOLIO_MANAGER,
                        agency_id=agency_two.id,
                        agent_id=agent_two.id,
                        is_active=True,
                        must_change_password=True,
                        temporary_password="DuplicateTwo!2026",
                        password_changed_at=datetime.now(timezone.utc),
                    ),
                ]
            )
            db.commit()

            provisioned = ensure_agent_accounts(db, issue_credentials=True)
            db.commit()

            users = db.scalars(
                select(User).where(User.role == UserRole.PORTFOLIO_MANAGER)
            ).all()
            self.assertEqual(len(users), 1)
            self.assertEqual(len(provisioned), 1)
            self.assertEqual(users[0].agent_id, agent_one.id)

    def test_portfolio_manager_scope_combines_same_gp_across_agencies(self) -> None:
        with SessionLocal() as db:
            agency_one = Agency(name="Agence Nord")
            agency_two = Agency(name="Agence Sud")
            db.add_all([agency_one, agency_two])
            db.flush()

            agent_one = Agent(name="Shared Agent", agency_id=agency_one.id)
            agent_two = Agent(name="Shared Agent", agency_id=agency_two.id)
            db.add_all([agent_one, agent_two])
            db.flush()

            pm_user = User(
                email="shared.agent.nord@microcred.com.tn",
                full_name="Shared Agent North",
                hashed_password=hash_password("SharedAgent!2026"),
                role=UserRole.PORTFOLIO_MANAGER,
                agency_id=agency_one.id,
                agent_id=agent_one.id,
                is_active=True,
                password_changed_at=datetime.now(timezone.utc),
            )
            db.add(pm_user)
            db.flush()

            db.add_all(
                [
                    LoanRaw(
                        contract_no="NORD-001",
                        client_name="Client Nord",
                        client_first_name="Alpha",
                        client_id="CLI-N-1",
                        client_ncni="NCNI-N-1",
                        agency_name=agency_one.name,
                        agent_name=agent_one.name,
                        client_rating="Classe 0",
                        category_desc="Test",
                        disbursement_amount=1000,
                        principal_outstanding=900,
                        principal_due=100,
                        total_scheduled_amount=200,
                        days_overdue=0,
                        total_due=80,
                        status="active",
                        disbursement_date=date(2026, 5, 10),
                        maturity_date=date(2027, 5, 10),
                        next_schedule_date=date(2026, 6, 5),
                        snapshot_date=date(2026, 5, 20),
                    ),
                    LoanRaw(
                        contract_no="SUD-001",
                        client_name="Client Sud",
                        client_first_name="Beta",
                        client_id="CLI-S-1",
                        client_ncni="NCNI-S-1",
                        agency_name=agency_two.name,
                        agent_name=agent_two.name,
                        client_rating="Classe 1",
                        category_desc="Test",
                        disbursement_amount=7000,
                        principal_outstanding=4000,
                        principal_due=500,
                        total_scheduled_amount=500,
                        days_overdue=45,
                        total_due=300,
                        status="active",
                        disbursement_date=date(2026, 5, 12),
                        maturity_date=date(2027, 5, 12),
                        next_schedule_date=date(2026, 6, 7),
                        snapshot_date=date(2026, 5, 20),
                    ),
                ]
            )
            db.commit()

        login_response = self._login("shared.agent.nord@microcred.com.tn", "SharedAgent!2026")
        self.assertEqual(login_response.status_code, 200, login_response.text)

        summary_response = self.client.get("/api/v1/metrics/summary")
        self.assertEqual(summary_response.status_code, 200, summary_response.text)
        summary = summary_response.json()
        self.assertEqual(summary["credits_count"], 2)
        self.assertEqual(summary["disbursement_count"], 2)
        self.assertEqual(summary["nb_clients"], 2)
        self.assertEqual(summary["outstanding"], "5500.00")
        self.assertEqual(summary["par_30"], "4500.00")

        metrics_response = self.client.get("/api/v1/metrics/daily?limit=50")
        self.assertEqual(metrics_response.status_code, 200, metrics_response.text)
        metrics_body = metrics_response.json()
        self.assertEqual(metrics_body["total"], 2)
        self.assertEqual(
            {item["agency"]["name"] for item in metrics_body["items"]},
            {"Agence Nord", "Agence Sud"},
        )

    def test_portfolio_manager_scope_matches_agent_name_punctuation_variants(self) -> None:
        with SessionLocal() as db:
            agency = Agency(name="Agence Centre")
            db.add(agency)
            db.flush()

            agent = Agent(name="ABIR.BEN.HOUIDI", agency_id=agency.id)
            db.add(agent)
            db.flush()

            pm_user = User(
                email="abir.ben@microcred.com.tn",
                full_name="ABIR.BEN.HOUIDI",
                hashed_password=hash_password("AbirSecure!2026"),
                role=UserRole.PORTFOLIO_MANAGER,
                agency_id=agency.id,
                agent_id=agent.id,
                is_active=True,
                password_changed_at=datetime.now(timezone.utc),
            )
            db.add(pm_user)
            db.flush()

            db.add(
                LoanRaw(
                    contract_no="CENTRE-001",
                    client_name="Client Centre",
                    client_first_name="Gamma",
                    client_id="CLI-C-1",
                    client_ncni="NCNI-C-1",
                    agency_name=agency.name,
                    agent_name="ABIR,BEN,HOUIDI",
                    client_rating="Classe 0",
                    category_desc="Test",
                    disbursement_amount=2500,
                    principal_outstanding=1200,
                    principal_due=100,
                    total_scheduled_amount=300,
                    days_overdue=0,
                    total_due=90,
                    status="active",
                    disbursement_date=date(2026, 5, 14),
                    maturity_date=date(2027, 5, 14),
                    next_schedule_date=date(2026, 6, 9),
                    snapshot_date=date(2026, 5, 20),
                )
            )
            db.commit()

        login_response = self._login("abir.ben@microcred.com.tn", "AbirSecure!2026")
        self.assertEqual(login_response.status_code, 200, login_response.text)

        summary_response = self.client.get("/api/v1/metrics/summary")
        self.assertEqual(summary_response.status_code, 200, summary_response.text)
        summary = summary_response.json()
        self.assertEqual(summary["credits_count"], 1)
        self.assertEqual(summary["disbursement_count"], 1)
        self.assertEqual(summary["nb_clients"], 1)
        self.assertEqual(summary["outstanding"], "1300.00")

    def test_security_headers_are_applied(self) -> None:
        response = self.client.get("/health")
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.headers.get("X-Frame-Options"), "DENY")
        self.assertEqual(response.headers.get("X-Content-Type-Options"), "nosniff")
        self.assertEqual(response.headers.get("Referrer-Policy"), "strict-origin-when-cross-origin")
        self.assertIn("default-src 'self'", response.headers.get("Content-Security-Policy", ""))

    def test_super_admin_can_delete_snapshots_but_not_current_state(self) -> None:
        self._seed_super_admin()
        snapshot_date = date(2026, 5, 10)
        current_date = date(2026, 5, 11)

        with SessionLocal() as db:
            agency = Agency(name="Agence Snapshot")
            db.add(agency)
            db.flush()
            agent = Agent(name="Snapshot Agent", agency_id=agency.id)
            db.add(agent)
            db.flush()
            snapshot_batch = ImportBatch(
                batch_type=ImportBatchType.SNAPSHOT,
                period=None,
                snapshot_date=snapshot_date,
                file_name="snapshot.xlsx",
            )
            current_batch = ImportBatch(
                batch_type=ImportBatchType.CURRENT_STATE,
                period=None,
                snapshot_date=current_date,
                file_name="current.xlsx",
            )
            db.add_all([snapshot_batch, current_batch])
            db.flush()

            db.add(
                LoanRaw(
                    contract_no="SNAPSHOT-001",
                    client_name="Snapshot",
                    client_first_name="Client",
                    client_id="CLI-SNAPSHOT",
                    agency_name=agency.name,
                    agent_name=agent.name,
                    disbursement_amount=1000,
                    principal_outstanding=800,
                    principal_due=100,
                    total_scheduled_amount=100,
                    days_overdue=0,
                    total_due=50,
                    status="active",
                    disbursement_date=date(2026, 5, 1),
                    snapshot_date=snapshot_date,
                    import_batch_id=snapshot_batch.id,
                )
            )
            db.add(
                DailyMetric(
                    agency_id=agency.id,
                    agent_id=agent.id,
                    date=snapshot_date,
                    disbursement_count=1,
                    disbursement_volume=1000,
                    nb_clients=1,
                    outstanding=900,
                    healthy_outstanding=900,
                    par_0=0,
                    par_1_30=0,
                    par_31_60=0,
                    par_30=0,
                )
            )
            db.commit()
            snapshot_batch_id = snapshot_batch.id
            current_batch_id = current_batch.id

        login_response = self._login(self.super_admin_email, self.super_admin_password)
        self.assertEqual(login_response.status_code, 200, login_response.text)

        delete_response = self.client.post(
            "/api/v1/imports/snapshots/delete",
            json={"batch_ids": [snapshot_batch_id]},
        )
        self.assertEqual(delete_response.status_code, 200, delete_response.text)
        self.assertEqual(delete_response.json()["deleted_batch_ids"], [snapshot_batch_id])
        self.assertEqual(delete_response.json()["deleted_rows"], 1)

        with SessionLocal() as db:
            self.assertIsNone(db.get(ImportBatch, snapshot_batch_id))
            self.assertIsNotNone(db.get(ImportBatch, current_batch_id))
            self.assertEqual(db.query(LoanRaw).filter(LoanRaw.import_batch_id == snapshot_batch_id).count(), 0)
            self.assertEqual(db.query(DailyMetric).filter(DailyMetric.date == snapshot_date).count(), 0)

        protected_response = self.client.post(
            "/api/v1/imports/snapshots/delete",
            json={"batch_ids": [current_batch_id]},
        )
        self.assertEqual(protected_response.status_code, 400, protected_response.text)
        self.assertEqual(protected_response.json()["detail"]["code"], "snapshot_delete_error")

    def test_date_to_only_uses_mcr_for_that_month(self) -> None:
        self._seed_super_admin()

        with SessionLocal() as db:
            agency = Agency(name="Agence Historique")
            db.add(agency)
            db.flush()
            agent = Agent(name="Historique Agent", agency_id=agency.id)
            db.add(agent)
            db.flush()

            historical_batch = ImportBatch(
                batch_type=ImportBatchType.HISTORICAL_MONTH,
                period="2025-12",
                snapshot_date=date(2025, 12, 31),
                file_name="mcr-2025-12.xlsx",
            )
            current_batch = ImportBatch(
                batch_type=ImportBatchType.CURRENT_STATE,
                period=None,
                snapshot_date=date(2026, 5, 16),
                file_name="mcr-current.xlsx",
            )
            db.add_all([historical_batch, current_batch])
            db.flush()

            db.add_all(
                [
                    LoanRaw(
                        contract_no="HIST-BEFORE-MONTH",
                        client_name="Client",
                        client_first_name="Avant Mois",
                        client_id="CLI-BEFORE-MONTH",
                        agency_name=agency.name,
                        agent_name=agent.name,
                        disbursement_amount=300,
                        principal_outstanding=200,
                        principal_due=50,
                        total_scheduled_amount=100,
                        days_overdue=0,
                        total_due=0,
                        status="active",
                        disbursement_date=date(2025, 11, 30),
                        snapshot_date=date(2025, 12, 31),
                        import_batch_id=historical_batch.id,
                    ),
                    LoanRaw(
                        contract_no="HIST-DEC-001",
                        client_name="Client",
                        client_first_name="Decembre",
                        client_id="CLI-DEC-001",
                        agency_name=agency.name,
                        agent_name=agent.name,
                        disbursement_amount=1200,
                        principal_outstanding=900,
                        principal_due=100,
                        total_scheduled_amount=100,
                        days_overdue=0,
                        total_due=0,
                        status="active",
                        disbursement_date=date(2025, 12, 10),
                        snapshot_date=date(2025, 12, 31),
                        import_batch_id=historical_batch.id,
                    ),
                    LoanRaw(
                        contract_no="HIST-DEC-LATE",
                        client_name="Client",
                        client_first_name="Apres Filtre",
                        client_id="CLI-DEC-002",
                        agency_name=agency.name,
                        agent_name=agent.name,
                        disbursement_amount=5000,
                        principal_outstanding=4000,
                        principal_due=200,
                        total_scheduled_amount=100,
                        days_overdue=0,
                        total_due=0,
                        status="active",
                        disbursement_date=date(2025, 12, 20),
                        snapshot_date=date(2025, 12, 31),
                        import_batch_id=historical_batch.id,
                    ),
                    LoanRaw(
                        contract_no="CURRENT-001",
                        client_name="Client",
                        client_first_name="Courant",
                        client_id="CLI-CURRENT",
                        agency_name=agency.name,
                        agent_name=agent.name,
                        disbursement_amount=9999,
                        principal_outstanding=8000,
                        principal_due=500,
                        total_scheduled_amount=100,
                        days_overdue=45,
                        total_due=0,
                        status="active",
                        disbursement_date=date(2026, 5, 10),
                        snapshot_date=date(2026, 5, 16),
                        import_batch_id=current_batch.id,
                    ),
                ]
            )
            db.commit()

        login_response = self._login(self.super_admin_email, self.super_admin_password)
        self.assertEqual(login_response.status_code, 200, login_response.text)

        response = self.client.get("/api/v1/metrics/summary?date_to=2025-12-16")
        self.assertEqual(response.status_code, 200, response.text)
        summary = response.json()

        self.assertEqual(summary["disbursement_count"], 2)
        self.assertEqual(summary["disbursement_volume"], "1500.00")
        self.assertEqual(summary["outstanding"], "1250.00")
        self.assertEqual(summary["par_30"], "0.00")

        credits_response = self.client.get("/api/v1/metrics/current-credits?date_to=2025-12-16")
        self.assertEqual(credits_response.status_code, 200, credits_response.text)
        credits_payload = credits_response.json()
        self.assertEqual(credits_payload["total"], 2)
        self.assertEqual(
            {item["contract_no"] for item in credits_payload["items"]},
            {"HIST-BEFORE-MONTH", "HIST-DEC-001"},
        )

        charts_response = self.client.get("/api/v1/metrics/charts?date_to=2025-12-16")
        self.assertEqual(charts_response.status_code, 200, charts_response.text)
        charts_payload = charts_response.json()
        self.assertEqual(charts_payload["mode"], "GLOBAL_BY_AGENCY")
        self.assertEqual(len(charts_payload["points"]), 1)
        self.assertEqual(charts_payload["points"][0]["label"], "Agence Historique")
        self.assertEqual(charts_payload["points"][0]["disbursement_volume"], "1500.00")
        self.assertEqual(charts_payload["points"][0]["outstanding"], "1250.00")

    def test_super_admin_gp_mcr_audit_detects_new_and_missing_gp_accounts(self) -> None:
        self._seed_super_admin()

        with SessionLocal() as db:
            agency = Agency(name="Agence Audit")
            missing_agency = Agency(name="Agence Ancienne")
            db.add_all([agency, missing_agency])
            db.flush()

            new_agent = Agent(name="New GP", agency_id=agency.id)
            pending_agent = Agent(name="Pending GP", agency_id=agency.id)
            absent_agent = Agent(name="Absent GP", agency_id=missing_agency.id)
            db.add_all([new_agent, pending_agent, absent_agent])
            db.flush()

            pending_user = User(
                email="pending.gp@microcred.com.tn",
                full_name="Pending GP",
                hashed_password=hash_password("PendingSecure!2026"),
                role=UserRole.PORTFOLIO_MANAGER,
                agency_id=agency.id,
                agent_id=pending_agent.id,
                is_active=True,
                must_change_password=True,
                password_changed_at=None,
                temporary_password=None,
            )
            absent_user = User(
                email="absent.gp@microcred.com.tn",
                full_name="Absent GP",
                hashed_password=hash_password("AbsentSecure!2026"),
                role=UserRole.PORTFOLIO_MANAGER,
                agency_id=missing_agency.id,
                agent_id=absent_agent.id,
                is_active=True,
                must_change_password=False,
                password_changed_at=datetime.now(timezone.utc),
            )
            db.add_all([pending_user, absent_user])
            db.flush()

            batch = ImportBatch(
                batch_type=ImportBatchType.CURRENT_STATE,
                period=None,
                snapshot_date=date(2026, 6, 7),
                file_name="mcr-audit.xlsx",
            )
            db.add(batch)
            db.flush()

            db.add_all(
                [
                    LoanRaw(
                        contract_no="AUDIT-NEW-001",
                        client_name="Client",
                        client_first_name="Nouveau",
                        client_id="CLI-AUDIT-NEW",
                        agency_name=agency.name,
                        agent_name=new_agent.name,
                        disbursement_amount=1000,
                        principal_outstanding=800,
                        principal_due=100,
                        total_scheduled_amount=100,
                        days_overdue=0,
                        total_due=0,
                        status="active",
                        disbursement_date=date(2026, 6, 1),
                        snapshot_date=date(2026, 6, 7),
                        import_batch_id=batch.id,
                    ),
                    LoanRaw(
                        contract_no="AUDIT-PENDING-001",
                        client_name="Client",
                        client_first_name="Pending",
                        client_id="CLI-AUDIT-PENDING",
                        agency_name=agency.name,
                        agent_name=pending_agent.name,
                        disbursement_amount=2000,
                        principal_outstanding=1500,
                        principal_due=100,
                        total_scheduled_amount=100,
                        days_overdue=0,
                        total_due=0,
                        status="active",
                        disbursement_date=date(2026, 6, 2),
                        snapshot_date=date(2026, 6, 7),
                        import_batch_id=batch.id,
                    ),
                ]
            )
            db.commit()

        login_response = self._login(self.super_admin_email, self.super_admin_password)
        self.assertEqual(login_response.status_code, 200, login_response.text)

        response = self.client.get("/api/v1/users/gp-mcr-audit")
        self.assertEqual(response.status_code, 200, response.text)
        audit = response.json()

        self.assertEqual(audit["import_batch"]["file_name"], "mcr-audit.xlsx")
        action_rows = {row["agent_name"]: row for row in audit["new_gp_without_accounts"]}
        self.assertEqual(action_rows["New GP"]["action"], "CREATE_ACCOUNT")
        self.assertFalse(action_rows["New GP"]["has_account"])
        self.assertEqual(action_rows["Pending GP"]["action"], "PROVISION_PASSWORD")
        self.assertTrue(action_rows["Pending GP"]["has_account"])
        self.assertEqual(action_rows["Pending GP"]["loan_count"], 1)

        missing_rows = {row["agent_name"]: row for row in audit["existing_accounts_not_in_mcr"]}
        self.assertIn("Absent GP", missing_rows)
        self.assertEqual(missing_rows["Absent GP"]["email"], "absent.gp@microcred.com.tn")

    def test_current_credits_can_filter_multiple_risk_categories(self) -> None:
        self._seed_super_admin()

        with SessionLocal() as db:
            agency = Agency(name="Agence Risque")
            db.add(agency)
            db.flush()
            agent = Agent(name="Risk GP", agency_id=agency.id)
            db.add(agent)
            db.flush()
            batch = ImportBatch(
                batch_type=ImportBatchType.CURRENT_STATE,
                period=None,
                snapshot_date=date(2026, 6, 18),
                file_name="mcr-risk.xlsx",
            )
            db.add(batch)
            db.flush()

            def loan(contract_no: str, days_overdue: int) -> LoanRaw:
                return LoanRaw(
                    contract_no=contract_no,
                    client_name="Client",
                    client_first_name=contract_no,
                    client_id=f"CLI-{contract_no}",
                    agency_name=agency.name,
                    agent_name=agent.name,
                    disbursement_amount=1000,
                    principal_outstanding=900,
                    principal_due=100,
                    total_scheduled_amount=75,
                    days_overdue=days_overdue,
                    total_due=0,
                    status="active",
                    disbursement_date=date(2026, 6, 5),
                    snapshot_date=date(2026, 6, 18),
                    import_batch_id=batch.id,
                )

            db.add_all(
                [
                    loan("RISK-000", 0),
                    loan("RISK-010", 10),
                    loan("RISK-045", 45),
                    loan("RISK-075", 75),
                    loan("RISK-130", 130),
                ]
            )
            db.commit()

        login_response = self._login(self.super_admin_email, self.super_admin_password)
        self.assertEqual(login_response.status_code, 200, login_response.text)

        response = self.client.get(
            "/api/v1/metrics/current-credits?risk_categories=cohort_1_30,par_120&limit=20"
        )
        self.assertEqual(response.status_code, 200, response.text)
        payload = response.json()
        self.assertEqual(payload["total"], 2)
        self.assertEqual(
            {item["contract_no"] for item in payload["items"]},
            {"RISK-010", "RISK-130"},
        )

        par30_response = self.client.get("/api/v1/metrics/current-credits?risk_categories=par_30&limit=20")
        self.assertEqual(par30_response.status_code, 200, par30_response.text)
        par30_payload = par30_response.json()
        self.assertEqual(par30_payload["total"], 3)
        self.assertEqual(
            {item["contract_no"] for item in par30_payload["items"]},
            {"RISK-045", "RISK-075", "RISK-130"},
        )

        invalid_response = self.client.get("/api/v1/metrics/current-credits?risk_categories=unknown")
        self.assertEqual(invalid_response.status_code, 400, invalid_response.text)

    def test_taeg_module_permissions_are_limited_to_admin_scopes(self) -> None:
        self._seed_super_admin()
        admin_password = "RegionalAdmin!2026"
        manager_password = "AgencyManager!2026"
        self._create_user(
            email="regional.admin@microcred.com.tn",
            password=admin_password,
            role=UserRole.ADMIN,
            full_name="Regional Admin",
        )
        self._create_user(
            email="agency.manager@microcred.com.tn",
            password=manager_password,
            role=UserRole.AGENCY_MANAGER,
            full_name="Agency Manager",
        )

        login_response = self._login("regional.admin@microcred.com.tn", admin_password)
        self.assertEqual(login_response.status_code, 200, login_response.text)
        admin_periods_response = self.client.get("/api/v1/taeg/periods")
        self.assertEqual(admin_periods_response.status_code, 200, admin_periods_response.text)

        self.client.post("/api/v1/auth/logout")
        manager_login = self._login("agency.manager@microcred.com.tn", manager_password)
        self.assertEqual(manager_login.status_code, 200, manager_login.text)
        manager_periods_response = self.client.get("/api/v1/taeg/periods")
        self.assertEqual(manager_periods_response.status_code, 403, manager_periods_response.text)

    def test_taeg_dashboard_uses_only_credits_disbursed_in_selected_snapshot_month(self) -> None:
        self._seed_super_admin()

        with SessionLocal() as db:
            agency = Agency(name="Agence TAEG")
            db.add(agency)
            db.flush()
            agent = Agent(name="Agent TAEG", agency_id=agency.id)
            db.add(agent)
            db.flush()

            sector = ActivitySector(name="Commerce", acm_rate=Decimal("0.12"))
            db.add(sector)
            db.flush()
            db.add(CategorySectorMapping(category_desc="Commerce local", activity_sector_id=sector.id))

            current_batch = ImportBatch(
                batch_type=ImportBatchType.CURRENT_STATE,
                period=None,
                snapshot_date=date(2026, 6, 20),
                file_name="mcr-taeg-current.xlsx",
            )
            historical_batch = ImportBatch(
                batch_type=ImportBatchType.HISTORICAL_MONTH,
                period="2026-05",
                snapshot_date=date(2026, 5, 31),
                file_name="mcr-taeg-2026-05.xlsx",
            )
            db.add_all([current_batch, historical_batch])
            db.flush()

            common = dict(
                agency_name=agency.name,
                agent_name=agent.name,
                client_name="Client",
                client_first_name="TAEG",
                client_rating="Classe 0",
                principal_outstanding=1000,
                principal_due=100,
                total_scheduled_amount=100,
                days_overdue=0,
                total_due=0,
                status="active",
                snapshot_date=date(2026, 6, 20),
                import_batch_id=current_batch.id,
                maturity_date=None,
                next_schedule_date=None,
                category_desc="Commerce local",
            )
            db.add_all(
                [
                    LoanRaw(
                        contract_no="TAEG-JUNE-001",
                        client_id="TAEG-001",
                        disbursement_amount=1000,
                        disbursement_date=date(2026, 6, 2),
                        teg_rate=Decimal("0.20"),
                        **common,
                    ),
                    LoanRaw(
                        contract_no="TAEG-JUNE-002",
                        client_id="TAEG-002",
                        disbursement_amount=3000,
                        disbursement_date=date(2026, 6, 18),
                        teg_rate=Decimal("0.10"),
                        **common,
                    ),
                    LoanRaw(
                        contract_no="TAEG-MAY-IGNORED",
                        client_id="TAEG-003",
                        disbursement_amount=5000,
                        disbursement_date=date(2026, 5, 28),
                        teg_rate=Decimal("0.25"),
                        **common,
                    ),
                    LoanRaw(
                        contract_no="TAEG-JULY-IGNORED",
                        client_id="TAEG-004",
                        disbursement_amount=6000,
                        disbursement_date=date(2026, 7, 3),
                        teg_rate=Decimal("0.08"),
                        **common,
                    ),
                ]
            )
            create_acm_rate_version(
                db,
                created_by_user_id=None,
                effective_start_date=date(2026, 1, 1),
                is_open_ended=True,
                details=[{"sector_id": sector.id, "acm_rate": Decimal("0.1200")}],
                comment="Bloc courant",
                created_at=datetime(2026, 1, 2, 8, 0),
            )
            db.commit()

        login_response = self._login(self.super_admin_email, self.super_admin_password)
        self.assertEqual(login_response.status_code, 200, login_response.text)

        dashboard_response = self.client.get("/api/v1/taeg/dashboard?period_key=current")
        self.assertEqual(dashboard_response.status_code, 200, dashboard_response.text)
        payload = dashboard_response.json()
        self.assertEqual(payload["period"]["snapshot_date"], "2026-06-20")
        self.assertEqual(len(payload["rows"]), 1)
        row = payload["rows"][0]
        self.assertEqual(row["sector_name"], "Commerce")
        self.assertEqual(row["credits_count"], 2)
        self.assertEqual(Decimal(row["disbursement_amount"]), Decimal("4000.00"))
        self.assertEqual(Decimal(row["taeg_calculated"]), Decimal("500.0000"))
        self.assertEqual(Decimal(row["taeg_weighted_rate"]), Decimal("0.1250"))
        self.assertEqual(row["status"], "non_conforme")

        details_response = self.client.get("/api/v1/taeg/details?period_key=current&limit=20")
        self.assertEqual(details_response.status_code, 200, details_response.text)
        details_payload = details_response.json()
        self.assertEqual(details_payload["total"], 2)
        self.assertEqual(
            {item["contract_no"] for item in details_payload["items"]},
            {"TAEG-JUNE-001", "TAEG-JUNE-002"},
        )

    def test_upsert_snapshot_rows_updates_teg_rate_on_existing_snapshot_row(self) -> None:
        snapshot = date(2026, 6, 30)

        with SessionLocal() as db:
            batch_a = ImportBatch(
                batch_type=ImportBatchType.HISTORICAL_MONTH,
                period="2026-06",
                snapshot_date=snapshot,
                file_name="historical-taeg.xlsx",
            )
            batch_b = ImportBatch(
                batch_type=ImportBatchType.CURRENT_STATE,
                period=None,
                snapshot_date=snapshot,
                file_name="current-taeg.xlsx",
            )
            db.add_all([batch_a, batch_b])
            db.flush()

            db.add(
                LoanRaw(
                    contract_no="UPSERT-001",
                    client_name="Client",
                    client_first_name="Legacy",
                    client_id="CLI-UPSERT-001",
                    agency_name="Agence Legacy",
                    agent_name="Agent Legacy",
                    client_rating="Classe 0",
                    category_desc="Commerce local",
                    teg_rate=Decimal("0.0000"),
                    disbursement_amount=Decimal("1000.00"),
                    principal_outstanding=Decimal("900.00"),
                    principal_due=Decimal("100.00"),
                    total_scheduled_amount=Decimal("100.00"),
                    days_overdue=0,
                    total_due=Decimal("0.00"),
                    status="active",
                    disbursement_date=snapshot,
                    maturity_date=None,
                    next_schedule_date=None,
                    snapshot_date=snapshot,
                    import_batch_id=batch_a.id,
                )
            )
            db.flush()

            rows = [
                LoanImportRow(
                    contract_no="UPSERT-001",
                    client_name="Client",
                    client_first_name="Updated",
                    client_id="CLI-UPSERT-001",
                    client_ncni=None,
                    agency_name="Agence Legacy",
                    agent_name="Agent Legacy",
                    client_rating="Classe 0",
                    category_desc="Commerce local",
                    teg_rate=Decimal("0.2350"),
                    disbursement_amount=Decimal("2500.00"),
                    principal_outstanding=Decimal("2000.00"),
                    principal_due=Decimal("250.00"),
                    total_scheduled_amount=Decimal("250.00"),
                    days_overdue=0,
                    total_due=Decimal("0.00"),
                    disbursement_date=snapshot,
                    maturity_date=None,
                    next_schedule_date=None,
                    snapshot_date=snapshot,
                    status="active",
                )
            ]

            inserted, updated, duplicates = _upsert_snapshot_rows(
                db=db,
                rows=rows,
                snapshot_date=snapshot,
                import_batch_id=batch_b.id,
            )
            db.commit()

            self.assertEqual(inserted, 0)
            self.assertEqual(updated, 1)
            self.assertEqual(duplicates, 0)

            updated_row = db.scalar(
                select(LoanRaw).where(
                    LoanRaw.contract_no == "UPSERT-001",
                    LoanRaw.snapshot_date == snapshot,
                )
            )
            self.assertIsNotNone(updated_row)
            self.assertEqual(updated_row.import_batch_id, batch_b.id)
            self.assertEqual(updated_row.teg_rate, Decimal("0.2350"))
            self.assertEqual(updated_row.disbursement_amount, Decimal("2500.00"))

    def test_taeg_semesters_and_custom_periods_use_only_available_selected_months(self) -> None:
        self._seed_super_admin()

        with SessionLocal() as db:
            agency = Agency(name="Agence TAEG S")
            db.add(agency)
            db.flush()
            agent = Agent(name="Agent TAEG S", agency_id=agency.id)
            db.add(agent)
            db.flush()

            sector = ActivitySector(name="Services", acm_rate=Decimal("0.18"))
            db.add(sector)
            db.flush()
            db.add(CategorySectorMapping(category_desc="Services locaux", activity_sector_id=sector.id))

            jan_batch = ImportBatch(
                batch_type=ImportBatchType.HISTORICAL_MONTH,
                period="2026-01",
                snapshot_date=date(2026, 1, 31),
                file_name="mcr-2026-01.xlsx",
            )
            feb_batch = ImportBatch(
                batch_type=ImportBatchType.HISTORICAL_MONTH,
                period="2026-02",
                snapshot_date=date(2026, 2, 28),
                file_name="mcr-2026-02.xlsx",
            )
            aug_batch = ImportBatch(
                batch_type=ImportBatchType.HISTORICAL_MONTH,
                period="2026-08",
                snapshot_date=date(2026, 8, 31),
                file_name="mcr-2026-08.xlsx",
            )
            db.add_all([jan_batch, feb_batch, aug_batch])
            db.flush()

            def add_loan(batch: ImportBatch, contract_no: str, amount: str, teg_rate: str, disb_date: date):
                db.add(
                    LoanRaw(
                        contract_no=contract_no,
                        client_name="Client",
                        client_first_name=contract_no,
                        client_id=f"CLI-{contract_no}",
                        agency_name=agency.name,
                        agent_name=agent.name,
                        client_rating="Classe 0",
                        category_desc="Services locaux",
                        teg_rate=Decimal(teg_rate),
                        disbursement_amount=Decimal(amount),
                        principal_outstanding=Decimal("1000.00"),
                        principal_due=Decimal("100.00"),
                        total_scheduled_amount=Decimal("100.00"),
                        days_overdue=0,
                        total_due=Decimal("0.00"),
                        status="active",
                        disbursement_date=disb_date,
                        maturity_date=None,
                        next_schedule_date=None,
                        snapshot_date=batch.snapshot_date,
                        import_batch_id=batch.id,
                    )
                )

            add_loan(jan_batch, "TAEG-S1-JAN", "1000.00", "0.1000", date(2026, 1, 10))
            add_loan(jan_batch, "TAEG-S1-OFF", "9000.00", "0.5000", date(2026, 3, 1))
            add_loan(feb_batch, "TAEG-S1-FEB", "2000.00", "0.2000", date(2026, 2, 11))
            add_loan(aug_batch, "TAEG-S2-AUG", "4000.00", "0.3000", date(2026, 8, 6))

            create_acm_rate_version(
                db,
                created_by_user_id=None,
                effective_start_date=date(2026, 1, 1),
                effective_end_date=date(2026, 6, 30),
                details=[{"sector_id": sector.id, "acm_rate": Decimal("0.1800")}],
                comment="Bloc S1",
                created_at=datetime(2026, 1, 2, 8, 0),
            )
            create_acm_rate_version(
                db,
                created_by_user_id=None,
                effective_start_date=date(2026, 7, 1),
                is_open_ended=True,
                details=[{"sector_id": sector.id, "acm_rate": Decimal("0.3000")}],
                comment="Bloc S2",
                created_at=datetime(2026, 7, 2, 8, 0),
            )
            db.commit()

            jan_key = f"historical:{jan_batch.id}"
            aug_key = f"historical:{aug_batch.id}"

        login_response = self._login(self.super_admin_email, self.super_admin_password)
        self.assertEqual(login_response.status_code, 200, login_response.text)

        s1_response = self.client.get("/api/v1/taeg/dashboard?period_key=semester:s1")
        self.assertEqual(s1_response.status_code, 200, s1_response.text)
        s1_payload = s1_response.json()
        self.assertEqual(s1_payload["period"]["label"], "S1")
        self.assertEqual(len(s1_payload["used_periods"]), 2)
        self.assertEqual(len(s1_payload["acm_segments"]), 1)
        s1_row = s1_payload["rows"][0]
        self.assertEqual(s1_row["credits_count"], 2)
        self.assertEqual(Decimal(s1_row["disbursement_amount"]), Decimal("3000.00"))
        self.assertEqual(Decimal(s1_row["taeg_calculated"]), Decimal("500.0000"))
        self.assertAlmostEqual(float(s1_row["taeg_weighted_rate"]), 500.0 / 3000.0, places=6)
        self.assertEqual(Decimal(s1_row["acm_rate"]), Decimal("0.1800"))

        custom_response = self.client.get(
            f"/api/v1/taeg/dashboard?period_key=custom&period_keys={jan_key},{aug_key}"
        )
        self.assertEqual(custom_response.status_code, 200, custom_response.text)
        custom_payload = custom_response.json()
        self.assertEqual(custom_payload["period"]["label"], "Personnalise")
        self.assertEqual(len(custom_payload["used_periods"]), 2)
        self.assertEqual(len(custom_payload["acm_segments"]), 2)
        custom_row = custom_payload["rows"][0]
        self.assertEqual(custom_row["credits_count"], 1)
        self.assertEqual(Decimal(custom_row["disbursement_amount"]), Decimal("1000.00"))
        self.assertAlmostEqual(float(custom_row["taeg_weighted_rate"]), 0.1, places=6)
        self.assertEqual(Decimal(custom_row["acm_rate"]), Decimal("0.1800"))

        invalid_custom = self.client.get(f"/api/v1/taeg/dashboard?period_key=custom&period_keys={jan_key}")
        self.assertEqual(invalid_custom.status_code, 400, invalid_custom.text)

    def test_taeg_excel_export_matches_dashboard_without_credit_details(self) -> None:
        self._seed_super_admin()

        with SessionLocal() as db:
            agency = Agency(name="Agence Export TAEG")
            db.add(agency)
            db.flush()

            agent = Agent(name="Agent Export", agency_id=agency.id)
            db.add(agent)
            db.flush()

            sector = ActivitySector(name="Commerce", acm_rate=Decimal("18.50"))
            db.add(sector)
            db.flush()
            db.add(CategorySectorMapping(category_desc="Commerce local", activity_sector_id=sector.id))

            feb_batch = ImportBatch(
                batch_type=ImportBatchType.HISTORICAL_MONTH,
                period="2026-02",
                snapshot_date=date(2026, 2, 28),
                file_name="taeg-feb.xlsx",
            )
            apr_batch = ImportBatch(
                batch_type=ImportBatchType.HISTORICAL_MONTH,
                period="2026-04",
                snapshot_date=date(2026, 4, 30),
                file_name="taeg-apr.xlsx",
            )
            db.add_all([feb_batch, apr_batch])
            db.flush()

            db.add_all(
                [
                    LoanRaw(
                        contract_no="TAEG-EXP-001",
                        client_name="Client",
                        client_first_name="Un",
                        client_id="CLI-TAEG-1",
                        agency_name=agency.name,
                        agent_name=agent.name,
                        category_desc="Commerce local",
                        disbursement_amount=Decimal("2000"),
                        teg_rate=Decimal("12.50"),
                        principal_outstanding=Decimal("1500"),
                        principal_due=Decimal("100"),
                        total_scheduled_amount=Decimal("100"),
                        days_overdue=0,
                        total_due=0,
                        status="active",
                        disbursement_date=date(2026, 2, 8),
                        snapshot_date=date(2026, 2, 28),
                        import_batch_id=feb_batch.id,
                    ),
                    LoanRaw(
                        contract_no="TAEG-EXP-002",
                        client_name="Client",
                        client_first_name="Deux",
                        client_id="CLI-TAEG-2",
                        agency_name=agency.name,
                        agent_name=agent.name,
                        category_desc="Commerce local",
                        disbursement_amount=Decimal("3000"),
                        teg_rate=Decimal("14.00"),
                        principal_outstanding=Decimal("2200"),
                        principal_due=Decimal("120"),
                        total_scheduled_amount=Decimal("120"),
                        days_overdue=0,
                        total_due=0,
                        status="active",
                        disbursement_date=date(2026, 4, 12),
                        snapshot_date=date(2026, 4, 30),
                        import_batch_id=apr_batch.id,
                    ),
                ]
            )
            db.commit()
            feb_key = f"historical:{feb_batch.id}"
            apr_key = f"historical:{apr_batch.id}"

        login_response = self._login(self.super_admin_email, self.super_admin_password)
        self.assertEqual(login_response.status_code, 200, login_response.text)

        response = self.client.get(
            f"/api/v1/taeg/export.xlsx?period_key=custom&period_keys={feb_key},{apr_key}"
        )
        self.assertEqual(response.status_code, 200, response.text)
        self.assertIn(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            response.headers.get("content-type", ""),
        )
        self.assertIn("TAEG_Personnalise_", response.headers.get("content-disposition", ""))

        workbook = load_workbook(filename=BytesIO(response.content))
        sheet = workbook.active
        rendered_values = [
            "" if cell is None else str(cell)
            for row in sheet.iter_rows(values_only=True)
            for cell in row
        ]
        flattened = " | ".join(rendered_values)

        self.assertIn("MicroCred - Rapport TAEG", flattened)
        self.assertIn("Personnalise", flattened)
        self.assertIn("Bloc ACM", flattened)
        self.assertIn("Periode analysee", flattened)
        self.assertIn("Fevrier • Avril", flattened)
        self.assertIn("Commerce", flattened)
        self.assertIn("Nombre total de secteurs", flattened)
        self.assertNotIn("CONTRACT_NO", flattened)
        self.assertNotIn("Liste detaillee des credits", flattened)
        self.assertEqual(len(sheet._charts), 1)
        self.assertIn("Comparaison TAEG pondere / Taux ACM", str(sheet._charts[0].title.tx.rich.p[0].r[0].t))

    def test_taeg_dashboard_splits_month_acm_blocks_by_credit_disbursement_date(self) -> None:
        self._seed_super_admin()

        with SessionLocal() as db:
            agency = Agency(name="Agence TAEG Split")
            db.add(agency)
            db.flush()

            agent = Agent(name="Agent TAEG Split", agency_id=agency.id)
            db.add(agent)
            db.flush()

            sector = ActivitySector(name="Commerce", acm_rate=Decimal("0.12"))
            db.add(sector)
            db.flush()
            db.add(CategorySectorMapping(category_desc="Commerce local", activity_sector_id=sector.id))

            current_batch = ImportBatch(
                batch_type=ImportBatchType.CURRENT_STATE,
                period=None,
                snapshot_date=date(2026, 7, 31),
                file_name="taeg-split-july.xlsx",
            )
            db.add(current_batch)
            db.flush()

            db.add_all(
                [
                    LoanRaw(
                        contract_no="TAEG-SPLIT-001",
                        client_name="Client",
                        client_first_name="Un",
                        client_id="CLI-SPLIT-1",
                        agency_name=agency.name,
                        agent_name=agent.name,
                        category_desc="Commerce local",
                        disbursement_amount=Decimal("1000"),
                        teg_rate=Decimal("0.11"),
                        principal_outstanding=Decimal("900"),
                        principal_due=Decimal("50"),
                        total_scheduled_amount=Decimal("50"),
                        days_overdue=0,
                        total_due=0,
                        status="active",
                        disbursement_date=date(2026, 7, 12),
                        snapshot_date=current_batch.snapshot_date,
                        import_batch_id=current_batch.id,
                    ),
                    LoanRaw(
                        contract_no="TAEG-SPLIT-002",
                        client_name="Client",
                        client_first_name="Deux",
                        client_id="CLI-SPLIT-2",
                        agency_name=agency.name,
                        agent_name=agent.name,
                        category_desc="Commerce local",
                        disbursement_amount=Decimal("2000"),
                        teg_rate=Decimal("0.25"),
                        principal_outstanding=Decimal("1800"),
                        principal_due=Decimal("80"),
                        total_scheduled_amount=Decimal("80"),
                        days_overdue=0,
                        total_due=0,
                        status="active",
                        disbursement_date=date(2026, 7, 13),
                        snapshot_date=current_batch.snapshot_date,
                        import_batch_id=current_batch.id,
                    ),
                ]
            )

            create_acm_rate_version(
                db,
                created_by_user_id=None,
                effective_start_date=date(2026, 1, 1),
                effective_end_date=date(2026, 7, 12),
                details=[{"sector_id": sector.id, "acm_rate": Decimal("0.1200")}],
                comment="Bloc A",
                created_at=datetime(2026, 1, 5, 9, 0),
            )
            create_acm_rate_version(
                db,
                created_by_user_id=None,
                effective_start_date=date(2026, 7, 13),
                is_open_ended=True,
                details=[{"sector_id": sector.id, "acm_rate": Decimal("0.2000")}],
                comment="Bloc B",
                created_at=datetime(2026, 7, 13, 9, 0),
            )
            db.commit()

        login_response = self._login(self.super_admin_email, self.super_admin_password)
        self.assertEqual(login_response.status_code, 200, login_response.text)

        dashboard_response = self.client.get("/api/v1/taeg/dashboard?period_key=current")
        self.assertEqual(dashboard_response.status_code, 200, dashboard_response.text)
        payload = dashboard_response.json()

        self.assertEqual(len(payload["acm_segments"]), 2)
        self.assertEqual(payload["active_acm_segment"]["range_start"], "2026-07-01")
        self.assertEqual(payload["active_acm_segment"]["range_end"], "2026-07-12")
        self.assertEqual(payload["acm_segments"][0]["range_start"], "2026-07-01")
        self.assertEqual(payload["acm_segments"][0]["range_end"], "2026-07-12")
        self.assertEqual(payload["acm_segments"][1]["range_start"], "2026-07-13")
        self.assertEqual(payload["acm_segments"][1]["range_end"], "2026-07-31")
        self.assertIsNone(payload["acm_coverage_message"])
        self.assertEqual(payload["acm_segments"][0]["credits_count"], 1)
        self.assertEqual(payload["acm_segments"][1]["credits_count"], 1)

        row = payload["rows"][0]
        self.assertEqual(row["credits_count"], 1)
        self.assertEqual(row["status"], "conforme")
        self.assertEqual(row["compliant_credits_count"], 1)
        self.assertEqual(row["non_compliant_credits_count"], 0)
        self.assertEqual(row["uncovered_credits_count"], 0)
        self.assertAlmostEqual(float(row["acm_rate"]), 0.12, places=6)

        details_response = self.client.get("/api/v1/taeg/details?period_key=current&limit=20&offset=0")
        self.assertEqual(details_response.status_code, 200, details_response.text)
        detail_items = {
            item["contract_no"]: item
            for item in details_response.json()["items"]
        }
        self.assertEqual(set(detail_items.keys()), {"TAEG-SPLIT-001"})
        self.assertEqual(Decimal(detail_items["TAEG-SPLIT-001"]["acm_rate"]), Decimal("0.1200"))
        self.assertEqual(detail_items["TAEG-SPLIT-001"]["status"], "conforme")
        self.assertEqual(detail_items["TAEG-SPLIT-001"]["acm_block_label"], "01/01/2026 -> 12/07/2026")
        second_segment_response = self.client.get(
            "/api/v1/taeg/details",
            params={
                "period_key": "current",
                "limit": 20,
                "offset": 0,
                "acm_block_id": payload["acm_segments"][1]["version_id"],
                "period_start": payload["acm_segments"][1]["range_start"],
                "period_end": payload["acm_segments"][1]["range_end"],
            },
        )
        self.assertEqual(second_segment_response.status_code, 200, second_segment_response.text)
        second_segment_items = {item["contract_no"]: item for item in second_segment_response.json()["items"]}
        self.assertEqual(set(second_segment_items.keys()), {"TAEG-SPLIT-002"})
        self.assertEqual(Decimal(second_segment_items["TAEG-SPLIT-002"]["acm_rate"]), Decimal("0.2000"))
        self.assertEqual(second_segment_items["TAEG-SPLIT-002"]["status"], "non_conforme")
        self.assertEqual(second_segment_items["TAEG-SPLIT-002"]["acm_block_label"], "13/07/2026 -> Jusqu'a modification")

        fallback_response = self.client.get(
            "/api/v1/taeg/dashboard",
            params={
                "period_key": "current",
                "acm_block_id": 999999,
                "period_start": "2026-07-01",
                "period_end": "2026-07-31",
            },
        )
        self.assertEqual(fallback_response.status_code, 200, fallback_response.text)
        fallback_payload = fallback_response.json()
        self.assertEqual(fallback_payload["active_acm_segment"]["range_start"], "2026-07-01")
        self.assertEqual(fallback_payload["active_acm_segment"]["range_end"], "2026-07-12")

    def test_taeg_summary_status_uses_displayed_weighted_values(self) -> None:
        self._seed_super_admin()

        with SessionLocal() as db:
            agency = Agency(name="Agence TAEG Summary Status")
            db.add(agency)
            db.flush()

            agent = Agent(name="Agent TAEG Summary Status", agency_id=agency.id)
            db.add(agent)
            db.flush()

            sector = ActivitySector(name="Agriculture", acm_rate=Decimal("30.09"))
            db.add(sector)
            db.flush()
            db.add(CategorySectorMapping(category_desc="Agriculture locale", activity_sector_id=sector.id))

            current_batch = ImportBatch(
                batch_type=ImportBatchType.CURRENT_STATE,
                period=None,
                snapshot_date=date(2026, 6, 30),
                file_name="taeg-summary-status.xlsx",
            )
            db.add(current_batch)
            db.flush()

            common = {
                "client_name": "Client",
                "agency_name": agency.name,
                "agent_name": agent.name,
                "category_desc": "Agriculture locale",
                "principal_outstanding": Decimal("1000"),
                "principal_due": Decimal("100"),
                "total_scheduled_amount": Decimal("100"),
                "days_overdue": 0,
                "total_due": 0,
                "status": "active",
                "snapshot_date": current_batch.snapshot_date,
                "import_batch_id": current_batch.id,
            }
            db.add_all(
                [
                    LoanRaw(
                        contract_no="TAEG-SUMMARY-001",
                        client_first_name="A",
                        client_id="CLI-SUMMARY-1",
                        disbursement_amount=Decimal("1000"),
                        teg_rate=Decimal("40.00"),
                        disbursement_date=date(2026, 6, 10),
                        **common,
                    ),
                    LoanRaw(
                        contract_no="TAEG-SUMMARY-002",
                        client_first_name="B",
                        client_id="CLI-SUMMARY-2",
                        disbursement_amount=Decimal("9000"),
                        teg_rate=Decimal("27.00"),
                        disbursement_date=date(2026, 6, 11),
                        **common,
                    ),
                ]
            )
            create_acm_rate_version(
                db,
                created_by_user_id=None,
                effective_start_date=date(2026, 1, 1),
                is_open_ended=True,
                details=[{"sector_id": sector.id, "acm_rate": Decimal("30.09")}],
                comment="Bloc agriculture",
                created_at=datetime(2026, 1, 2, 8, 0),
            )
            db.commit()

        login_response = self._login(self.super_admin_email, self.super_admin_password)
        self.assertEqual(login_response.status_code, 200, login_response.text)

        dashboard_response = self.client.get("/api/v1/taeg/dashboard?period_key=current")
        self.assertEqual(dashboard_response.status_code, 200, dashboard_response.text)
        row = dashboard_response.json()["rows"][0]

        self.assertEqual(row["sector_name"], "Agriculture")
        self.assertEqual(Decimal(row["taeg_weighted_rate"]), Decimal("28.3000"))
        self.assertEqual(Decimal(row["acm_rate"]), Decimal("30.0900"))
        self.assertEqual(row["non_compliant_credits_count"], 1)
        self.assertEqual(row["status"], "conforme")

    def test_taeg_semester_s1_uses_applicable_acm_block_credit_by_credit(self) -> None:
        self._seed_super_admin()

        with SessionLocal() as db:
            agency = Agency(name="Agence TAEG S1 ACM")
            db.add(agency)
            db.flush()

            agent = Agent(name="Agent TAEG S1 ACM", agency_id=agency.id)
            db.add(agent)
            db.flush()

            sector = ActivitySector(name="Commerce S1", acm_rate=Decimal("0.12"))
            db.add(sector)
            db.flush()
            db.add(CategorySectorMapping(category_desc="Commerce S1", activity_sector_id=sector.id))

            march_batch = ImportBatch(
                batch_type=ImportBatchType.HISTORICAL_MONTH,
                period="2026-03",
                snapshot_date=date(2026, 3, 31),
                file_name="taeg-s1-march.xlsx",
            )
            april_batch = ImportBatch(
                batch_type=ImportBatchType.HISTORICAL_MONTH,
                period="2026-04",
                snapshot_date=date(2026, 4, 30),
                file_name="taeg-s1-april.xlsx",
            )
            db.add_all([march_batch, april_batch])
            db.flush()

            common = {
                "client_name": "Client",
                "agency_name": agency.name,
                "agent_name": agent.name,
                "category_desc": "Commerce S1",
                "principal_outstanding": Decimal("1000"),
                "principal_due": Decimal("100"),
                "total_scheduled_amount": Decimal("100"),
                "days_overdue": 0,
                "total_due": 0,
                "status": "active",
            }
            db.add_all(
                [
                    LoanRaw(
                        contract_no="TAEG-S1-BLOCK-A",
                        client_first_name="Bloc A",
                        client_id="CLI-S1-A",
                        disbursement_amount=Decimal("1000"),
                        teg_rate=Decimal("0.11"),
                        disbursement_date=date(2026, 3, 20),
                        snapshot_date=march_batch.snapshot_date,
                        import_batch_id=march_batch.id,
                        **common,
                    ),
                    LoanRaw(
                        contract_no="TAEG-S1-BLOCK-B",
                        client_first_name="Bloc B",
                        client_id="CLI-S1-B",
                        disbursement_amount=Decimal("2000"),
                        teg_rate=Decimal("0.25"),
                        disbursement_date=date(2026, 4, 20),
                        snapshot_date=april_batch.snapshot_date,
                        import_batch_id=april_batch.id,
                        **common,
                    ),
                ]
            )

            create_acm_rate_version(
                db,
                created_by_user_id=None,
                effective_start_date=date(2026, 1, 1),
                effective_end_date=date(2026, 3, 31),
                details=[{"sector_id": sector.id, "acm_rate": Decimal("0.1200")}],
                comment="Bloc S1 A",
                created_at=datetime(2026, 1, 5, 8, 0),
            )
            create_acm_rate_version(
                db,
                created_by_user_id=None,
                effective_start_date=date(2026, 4, 1),
                is_open_ended=True,
                details=[{"sector_id": sector.id, "acm_rate": Decimal("0.2000")}],
                comment="Bloc S1 B",
                created_at=datetime(2026, 4, 1, 8, 0),
            )
            db.commit()

        login_response = self._login(self.super_admin_email, self.super_admin_password)
        self.assertEqual(login_response.status_code, 200, login_response.text)

        dashboard_response = self.client.get("/api/v1/taeg/dashboard?period_key=semester:s1")
        self.assertEqual(dashboard_response.status_code, 200, dashboard_response.text)
        payload = dashboard_response.json()
        self.assertEqual(len(payload["acm_segments"]), 2)
        self.assertEqual(payload["active_acm_segment"]["block_label"], "Bloc ACM du 01/01/2026")
        row = payload["rows"][0]

        self.assertEqual(row["credits_count"], 1)
        self.assertEqual(row["status"], "conforme")
        self.assertEqual(row["compliant_credits_count"], 1)
        self.assertEqual(row["non_compliant_credits_count"], 0)
        self.assertEqual(row["uncovered_credits_count"], 0)
        self.assertAlmostEqual(float(row["acm_rate"]), 0.12, places=6)
        self.assertEqual({segment["block_label"] for segment in payload["acm_segments"]}, {"Bloc ACM du 01/01/2026", "Bloc ACM du 01/04/2026"})

        details_response = self.client.get("/api/v1/taeg/details?period_key=semester:s1&limit=20&offset=0")
        self.assertEqual(details_response.status_code, 200, details_response.text)
        detail_items = {item["contract_no"]: item for item in details_response.json()["items"]}
        self.assertEqual(set(detail_items.keys()), {"TAEG-S1-BLOCK-A"})
        self.assertEqual(detail_items["TAEG-S1-BLOCK-A"]["acm_block_label"], "01/01/2026 -> 31/03/2026")
        self.assertEqual(detail_items["TAEG-S1-BLOCK-A"]["status"], "conforme")
        second_segment = payload["acm_segments"][1]
        second_details_response = self.client.get(
            "/api/v1/taeg/details",
            params={
                "period_key": "semester:s1",
                "limit": 20,
                "offset": 0,
                "acm_block_id": second_segment["version_id"],
                "period_start": second_segment["range_start"],
                "period_end": second_segment["range_end"],
            },
        )
        self.assertEqual(second_details_response.status_code, 200, second_details_response.text)
        second_detail_items = {item["contract_no"]: item for item in second_details_response.json()["items"]}
        self.assertEqual(set(second_detail_items.keys()), {"TAEG-S1-BLOCK-B"})
        self.assertEqual(second_detail_items["TAEG-S1-BLOCK-B"]["acm_block_label"], "01/04/2026 -> Jusqu'a modification")
        self.assertEqual(second_detail_items["TAEG-S1-BLOCK-B"]["status"], "non_conforme")

    def test_taeg_dashboard_marks_uncovered_credit_when_no_acm_block_applies(self) -> None:
        self._seed_super_admin()

        with SessionLocal() as db:
            agency = Agency(name="Agence TAEG Uncovered")
            db.add(agency)
            db.flush()

            agent = Agent(name="Agent TAEG Uncovered", agency_id=agency.id)
            db.add(agent)
            db.flush()

            sector = ActivitySector(name="Services", acm_rate=Decimal("0.22"))
            db.add(sector)
            db.flush()
            db.add(CategorySectorMapping(category_desc="Services locaux", activity_sector_id=sector.id))

            current_batch = ImportBatch(
                batch_type=ImportBatchType.CURRENT_STATE,
                period=None,
                snapshot_date=date(2026, 7, 31),
                file_name="taeg-uncovered-july.xlsx",
            )
            db.add(current_batch)
            db.flush()

            db.add(
                LoanRaw(
                    contract_no="TAEG-UNCOVERED-001",
                    client_name="Client",
                    client_first_name="Unique",
                    client_id="CLI-UNCOV-1",
                    agency_name=agency.name,
                    agent_name=agent.name,
                    category_desc="Services locaux",
                    disbursement_amount=Decimal("1500"),
                    teg_rate=Decimal("0.18"),
                    principal_outstanding=Decimal("1200"),
                    principal_due=Decimal("60"),
                    total_scheduled_amount=Decimal("60"),
                    days_overdue=0,
                    total_due=0,
                    status="active",
                    disbursement_date=date(2026, 7, 12),
                    snapshot_date=current_batch.snapshot_date,
                    import_batch_id=current_batch.id,
                )
            )

            create_acm_rate_version(
                db,
                created_by_user_id=None,
                effective_start_date=date(2026, 7, 13),
                is_open_ended=True,
                details=[{"sector_id": sector.id, "acm_rate": Decimal("0.2200")}],
                comment="Bloc tardif",
                created_at=datetime(2026, 7, 13, 10, 0),
            )
            db.commit()

        login_response = self._login(self.super_admin_email, self.super_admin_password)
        self.assertEqual(login_response.status_code, 200, login_response.text)

        dashboard_response = self.client.get("/api/v1/taeg/dashboard?period_key=current")
        self.assertEqual(dashboard_response.status_code, 200, dashboard_response.text)
        payload = dashboard_response.json()
        self.assertIsNotNone(payload["acm_coverage_message"])
        self.assertTrue(any(segment["covered"] is False for segment in payload["acm_segments"]))

        row = payload["rows"][0]
        self.assertIsNone(row["acm_rate"])
        self.assertEqual(row["status"], "non_couvert")
        self.assertEqual(row["uncovered_credits_count"], 1)

        details_response = self.client.get("/api/v1/taeg/details?period_key=current&limit=20&offset=0")
        self.assertEqual(details_response.status_code, 200, details_response.text)
        detail = details_response.json()["items"][0]
        self.assertIsNone(detail["acm_rate"])
        self.assertEqual(detail["status"], "non_couvert")
        self.assertIsNone(detail["acm_block_label"])

    def test_taeg_rejects_overlapping_acm_blocks_for_selected_period(self) -> None:
        self._seed_super_admin()

        with SessionLocal() as db:
            agency = Agency(name="Agence TAEG Overlap")
            db.add(agency)
            db.flush()

            agent = Agent(name="Agent TAEG Overlap", agency_id=agency.id)
            db.add(agent)
            db.flush()

            sector = ActivitySector(name="Secteur Overlap", acm_rate=Decimal("0.10"))
            db.add(sector)
            db.flush()
            db.add(CategorySectorMapping(category_desc="Secteur Overlap", activity_sector_id=sector.id))

            march_batch = ImportBatch(
                batch_type=ImportBatchType.HISTORICAL_MONTH,
                period="2026-03",
                snapshot_date=date(2026, 3, 31),
                file_name="taeg-overlap-march.xlsx",
            )
            db.add(march_batch)
            db.flush()

            db.add(
                LoanRaw(
                    contract_no="TAEG-OVERLAP-001",
                    client_name="Client",
                    client_first_name="Overlap",
                    client_id="CLI-OVERLAP-1",
                    agency_name=agency.name,
                    agent_name=agent.name,
                    category_desc="Secteur Overlap",
                    disbursement_amount=Decimal("1200"),
                    teg_rate=Decimal("0.14"),
                    principal_outstanding=Decimal("1000"),
                    principal_due=Decimal("100"),
                    total_scheduled_amount=Decimal("100"),
                    days_overdue=0,
                    total_due=0,
                    status="active",
                    disbursement_date=date(2026, 3, 20),
                    snapshot_date=march_batch.snapshot_date,
                    import_batch_id=march_batch.id,
                )
            )

            version_a = AcmRateVersion(
                created_by_user_id=None,
                effective_start_date=date(2026, 1, 1),
                effective_end_date=date(2026, 3, 31),
                is_open_ended=False,
                status="historique",
                comment="Bloc overlap A",
                created_at=datetime(2026, 1, 2, 8, 0),
            )
            version_b = AcmRateVersion(
                created_by_user_id=None,
                effective_start_date=date(2026, 3, 15),
                effective_end_date=date(2026, 6, 30),
                is_open_ended=False,
                status="historique",
                comment="Bloc overlap B",
                created_at=datetime(2026, 3, 15, 8, 0),
            )
            db.add_all([version_a, version_b])
            db.flush()
            db.add_all(
                [
                    AcmRateVersionDetail(
                        version_id=version_a.id,
                        activity_sector_id=sector.id,
                        sector_name=sector.name,
                        acm_rate=Decimal("0.1200"),
                    ),
                    AcmRateVersionDetail(
                        version_id=version_b.id,
                        activity_sector_id=sector.id,
                        sector_name=sector.name,
                        acm_rate=Decimal("0.1500"),
                    ),
                ]
            )
            db.commit()

            with self.assertRaisesRegex(ValueError, "plusieurs blocs couvrent la date 20/03/2026"):
                get_applicable_acm_block(db, date(2026, 3, 20))

        login_response = self._login(self.super_admin_email, self.super_admin_password)
        self.assertEqual(login_response.status_code, 200, login_response.text)

        dashboard_response = self.client.get("/api/v1/taeg/dashboard?period_key=semester:s1")
        self.assertEqual(dashboard_response.status_code, 400, dashboard_response.text)
        self.assertIn("plusieurs blocs couvrent la sous-periode", dashboard_response.text)

        details_response = self.client.get("/api/v1/taeg/details?period_key=semester:s1&limit=20&offset=0")
        self.assertEqual(details_response.status_code, 400, details_response.text)
        self.assertIn("plusieurs blocs couvrent la sous-periode", details_response.text)

    def test_create_acm_rate_version_rejects_gap_between_blocks(self) -> None:
        with SessionLocal() as db:
            sector = ActivitySector(name="Agriculture", acm_rate=Decimal("0.10"))
            db.add(sector)
            db.flush()

            create_acm_rate_version(
                db,
                created_by_user_id=None,
                effective_start_date=date(2026, 1, 1),
                effective_end_date=date(2026, 1, 31),
                details=[{"sector_id": sector.id, "acm_rate": Decimal("0.1000")}],
                comment="Janvier",
                created_at=datetime(2026, 1, 1, 8, 0),
            )

            with self.assertRaisesRegex(ValueError, "intervalle non couvert"):
                create_acm_rate_version(
                    db,
                    created_by_user_id=None,
                    effective_start_date=date(2026, 3, 1),
                    is_open_ended=True,
                    details=[{"sector_id": sector.id, "acm_rate": Decimal("0.1200")}],
                    comment="Mars",
                    created_at=datetime(2026, 3, 1, 8, 0),
                )

    def test_taeg_details_filters_and_export_follow_active_table_view(self) -> None:
        self._seed_super_admin()

        with SessionLocal() as db:
            agency = Agency(name="Agence TAEG Filters")
            db.add(agency)
            db.flush()

            agent = Agent(name="Agent TAEG Filters", agency_id=agency.id)
            db.add(agent)
            db.flush()

            commerce = ActivitySector(name="Commerce", acm_rate=Decimal("10.00"))
            services = ActivitySector(name="Services", acm_rate=Decimal("20.00"))
            db.add_all([commerce, services])
            db.flush()

            db.add_all([
                CategorySectorMapping(category_desc="Commerce local", activity_sector_id=commerce.id),
                CategorySectorMapping(category_desc="Services locaux", activity_sector_id=services.id),
            ])

            current_batch = ImportBatch(
                batch_type=ImportBatchType.CURRENT_STATE,
                period=None,
                snapshot_date=date(2026, 7, 15),
                file_name="taeg-current-filters.xlsx",
            )
            db.add(current_batch)
            db.flush()

            common = {
                "agency_name": agency.name,
                "agent_name": agent.name,
                "client_rating": "Classe 0",
                "principal_outstanding": Decimal("1000.00"),
                "principal_due": Decimal("100.00"),
                "total_scheduled_amount": Decimal("100.00"),
                "days_overdue": 0,
                "total_due": Decimal("0.00"),
                "status": "active",
                "snapshot_date": current_batch.snapshot_date,
                "import_batch_id": current_batch.id,
            }
            db.add_all([
                LoanRaw(
                    contract_no="TAEG-FILTER-001",
                    client_name="Client",
                    client_first_name="Commerce 1",
                    client_id="CLI-TAEG-F-1",
                    category_desc="Commerce local",
                    disbursement_amount=Decimal("2000.00"),
                    disbursement_date=date(2026, 7, 2),
                    teg_rate=Decimal("12.00"),
                    **common,
                ),
                LoanRaw(
                    contract_no="TAEG-FILTER-002",
                    client_name="Client",
                    client_first_name="Commerce 2",
                    client_id="CLI-TAEG-F-2",
                    category_desc="Commerce local",
                    disbursement_amount=Decimal("3000.00"),
                    disbursement_date=date(2026, 7, 7),
                    teg_rate=Decimal("11.00"),
                    **common,
                ),
                LoanRaw(
                    contract_no="TAEG-FILTER-003",
                    client_name="Client",
                    client_first_name="Service 1",
                    client_id="CLI-TAEG-F-3",
                    category_desc="Services locaux",
                    disbursement_amount=Decimal("1500.00"),
                    disbursement_date=date(2026, 7, 9),
                    teg_rate=Decimal("18.00"),
                    **common,
                ),
            ])
            create_acm_rate_version(
                db,
                created_by_user_id=None,
                effective_start_date=date(2026, 1, 1),
                is_open_ended=True,
                details=[
                    {"sector_id": commerce.id, "acm_rate": Decimal("10.00")},
                    {"sector_id": services.id, "acm_rate": Decimal("20.00")},
                ],
                comment="Bloc filtres",
                created_at=datetime(2026, 1, 2, 8, 0),
            )
            db.commit()

        login_response = self._login(self.super_admin_email, self.super_admin_password)
        self.assertEqual(login_response.status_code, 200, login_response.text)

        details_filters = json.dumps({
            "sector_name": {"mode": "contains", "value": "commerce"},
            "disbursement_amount": {"mode": "gt", "value": 2500},
        })
        details_response = self.client.get(
            "/api/v1/taeg/details",
            params={
                "period_key": "current",
                "limit": 20,
                "offset": 0,
                "table_filters": details_filters,
                "sort_key": "disbursement_amount",
                "sort_direction": "desc",
            },
        )
        self.assertEqual(details_response.status_code, 200, details_response.text)
        details_payload = details_response.json()
        self.assertEqual(details_payload["total"], 1)
        self.assertEqual(len(details_payload["items"]), 1)
        self.assertEqual(details_payload["items"][0]["contract_no"], "TAEG-FILTER-002")

        export_filters = json.dumps({
            "status": {"mode": "in", "selected": ["non_conforme"]},
        })
        export_response = self.client.get(
            "/api/v1/taeg/export.xlsx",
            params={
                "period_key": "current",
                "table_filters": export_filters,
                "sort_key": "disbursement_amount",
                "sort_direction": "desc",
            },
        )
        self.assertEqual(export_response.status_code, 200, export_response.text)

        workbook = load_workbook(filename=BytesIO(export_response.content))
        sheet = workbook.active
        rendered_values = [
            "" if cell is None else str(cell)
            for row in sheet.iter_rows(values_only=True)
            for cell in row
        ]
        flattened = " | ".join(rendered_values)

        self.assertIn("Commerce", flattened)
        self.assertNotIn("Services", flattened)

    def test_taeg_monthly_history_endpoints_return_daily_points_grouped_by_sector(self) -> None:
        self._seed_super_admin()

        with SessionLocal() as db:
            agency = Agency(name="Agence Historique TAEG")
            db.add(agency)
            db.flush()
            agent = Agent(name="Agent Historique TAEG", agency_id=agency.id)
            db.add(agent)
            db.flush()

            commerce = ActivitySector(name="Commerce", acm_rate=Decimal("0.12"))
            services = ActivitySector(name="Services", acm_rate=Decimal("0.15"))
            db.add_all([commerce, services])
            db.flush()

            db.add_all(
                [
                    TaegDailySnapshot(
                        snapshot_date=date(2026, 3, 5),
                        imported_at=datetime(2026, 3, 5, 8, 0),
                        period_month=3,
                        period_year=2026,
                        agency_id=agency.id,
                        agent_id=agent.id,
                        sector_id=commerce.id,
                        sector_name="Commerce",
                        credits_count=2,
                        disbursement_amount=Decimal("4000.00"),
                        taeg_calculated=Decimal("520.0000"),
                        taeg_weighted_rate=Decimal("0.1300"),
                        acm_rate=Decimal("0.1200"),
                        status="non_conforme",
                    ),
                    TaegDailySnapshot(
                        snapshot_date=date(2026, 3, 5),
                        imported_at=datetime(2026, 3, 5, 8, 0),
                        period_month=3,
                        period_year=2026,
                        agency_id=agency.id,
                        agent_id=agent.id,
                        sector_id=services.id,
                        sector_name="Services",
                        credits_count=1,
                        disbursement_amount=Decimal("2500.00"),
                        taeg_calculated=Decimal("300.0000"),
                        taeg_weighted_rate=Decimal("0.1200"),
                        acm_rate=Decimal("0.1500"),
                        status="conforme",
                    ),
                    TaegDailySnapshot(
                        snapshot_date=date(2026, 3, 12),
                        imported_at=datetime(2026, 3, 12, 8, 0),
                        period_month=3,
                        period_year=2026,
                        agency_id=agency.id,
                        agent_id=agent.id,
                        sector_id=commerce.id,
                        sector_name="Commerce",
                        credits_count=3,
                        disbursement_amount=Decimal("5500.00"),
                        taeg_calculated=Decimal("660.0000"),
                        taeg_weighted_rate=Decimal("0.1200"),
                        acm_rate=Decimal("0.1200"),
                        status="conforme",
                    ),
                ]
            )
            db.commit()

        login_response = self._login(self.super_admin_email, self.super_admin_password)
        self.assertEqual(login_response.status_code, 200, login_response.text)

        periods_response = self.client.get("/api/v1/taeg/monthly-history/periods")
        self.assertEqual(periods_response.status_code, 200, periods_response.text)
        periods_payload = periods_response.json()
        self.assertEqual(periods_payload[0]["key"], "2026-03")
        self.assertEqual(periods_payload[0]["days_count"], 2)

        history_response = self.client.get("/api/v1/taeg/monthly-history?period=2026-03")
        self.assertEqual(history_response.status_code, 200, history_response.text)
        history_payload = history_response.json()
        self.assertEqual(history_payload["period"]["key"], "2026-03")
        self.assertEqual(history_payload["days_count"], 2)
        self.assertEqual({item["sector_name"] for item in history_payload["sectors"]}, {"Commerce", "Services"})
        self.assertEqual(len(history_payload["points"]), 3)

    def test_delete_taeg_daily_snapshots_for_period_removes_only_selected_month(self) -> None:
        with SessionLocal() as db:
            db.add_all(
                [
                    TaegDailySnapshot(
                        snapshot_date=date(2026, 3, 2),
                        imported_at=datetime(2026, 3, 2, 7, 0),
                        period_month=3,
                        period_year=2026,
                        sector_name="Commerce",
                        credits_count=1,
                        disbursement_amount=Decimal("1000.00"),
                        taeg_calculated=Decimal("100.0000"),
                        taeg_weighted_rate=Decimal("0.1000"),
                        acm_rate=Decimal("0.1200"),
                        status="conforme",
                    ),
                    TaegDailySnapshot(
                        snapshot_date=date(2026, 3, 18),
                        imported_at=datetime(2026, 3, 18, 7, 0),
                        period_month=3,
                        period_year=2026,
                        sector_name="Services",
                        credits_count=1,
                        disbursement_amount=Decimal("1500.00"),
                        taeg_calculated=Decimal("165.0000"),
                        taeg_weighted_rate=Decimal("0.1100"),
                        acm_rate=Decimal("0.1500"),
                        status="conforme",
                    ),
                    TaegDailySnapshot(
                        snapshot_date=date(2026, 4, 1),
                        imported_at=datetime(2026, 4, 1, 7, 0),
                        period_month=4,
                        period_year=2026,
                        sector_name="Commerce",
                        credits_count=1,
                        disbursement_amount=Decimal("1700.00"),
                        taeg_calculated=Decimal("204.0000"),
                        taeg_weighted_rate=Decimal("0.1200"),
                        acm_rate=Decimal("0.1200"),
                        status="conforme",
                    ),
                ]
            )
            db.commit()

            deleted = delete_taeg_daily_snapshots_for_period(db, "2026-03")
            db.commit()

            self.assertEqual(deleted, 2)
            remaining = db.scalars(select(TaegDailySnapshot).order_by(TaegDailySnapshot.snapshot_date.asc())).all()
            self.assertEqual(len(remaining), 1)
            self.assertEqual(remaining[0].snapshot_date, date(2026, 4, 1))

    def test_taeg_publish_acm_rate_versions_with_actor_history(self) -> None:
        self._seed_super_admin()
        login_response = self._login(self.super_admin_email, self.super_admin_password)
        self.assertEqual(login_response.status_code, 200, login_response.text)

        create_response = self.client.post(
            "/api/v1/taeg/sectors",
            json={"name": "Agriculture", "acm_rate": 0},
        )
        self.assertEqual(create_response.status_code, 200, create_response.text)
        sector_id = create_response.json()["id"]

        publish_response = self.client.post(
            "/api/v1/taeg/acm-rate-versions",
            json={
                "effective_start_date": "2026-01-01",
                "effective_end_date": "2026-06-30",
                "is_open_ended": False,
                "comment": "Version S1",
                "details": [
                    {"sector_id": sector_id, "acm_rate": 12.5},
                ],
            },
        )
        self.assertEqual(publish_response.status_code, 200, publish_response.text)

        publish_response_2 = self.client.post(
            "/api/v1/taeg/acm-rate-versions",
            json={
                "effective_start_date": "2026-07-01",
                "is_open_ended": True,
                "comment": "Version S2",
                "details": [
                    {"sector_id": sector_id, "acm_rate": 13.75},
                ],
            },
        )
        self.assertEqual(publish_response_2.status_code, 200, publish_response_2.text)

        versions_response = self.client.get("/api/v1/taeg/acm-rate-versions")
        self.assertEqual(versions_response.status_code, 200, versions_response.text)
        versions_payload = versions_response.json()
        self.assertGreaterEqual(len(versions_payload), 2)
        self.assertEqual(versions_payload[0]["created_by_name"], "Security Super Admin")
        self.assertEqual(versions_payload[0]["effective_start_date"], "2026-07-01")
        self.assertEqual(versions_payload[0]["effective_end_date"], None)
        self.assertTrue(versions_payload[0]["is_open_ended"])
        self.assertEqual(versions_payload[0]["sector_count"], 1)
        self.assertEqual(versions_payload[1]["effective_end_date"], "2026-06-30")

        version_detail_response = self.client.get(f"/api/v1/taeg/acm-rate-versions/{versions_payload[0]['id']}")
        self.assertEqual(version_detail_response.status_code, 200, version_detail_response.text)
        detail_payload = version_detail_response.json()
        self.assertEqual(detail_payload["created_by_name"], "Security Super Admin")
        self.assertEqual(detail_payload["details"][0]["sector_name"], "Agriculture")
        self.assertEqual(Decimal(detail_payload["details"][0]["acm_rate"]), Decimal("13.7500"))
        self.assertEqual(detail_payload["sector_count"], 1)

        with SessionLocal() as db:
            self.assertEqual(db.scalar(select(func.count(AcmRateVersion.id))), len(versions_payload))

    def test_taeg_update_acm_rate_version_tracks_audit_and_recalculates_snapshots(self) -> None:
        self._seed_super_admin()
        login_response = self._login(self.super_admin_email, self.super_admin_password)
        self.assertEqual(login_response.status_code, 200, login_response.text)

        with SessionLocal() as db:
            agency = Agency(name="Agence ACM")
            db.add(agency)
            db.flush()
            agent = Agent(name="Agent ACM", agency_id=agency.id)
            db.add(agent)
            db.flush()
            sector = ActivitySector(name="Agriculture", acm_rate=Decimal("0.1200"))
            db.add(sector)
            db.flush()
            db.add(CategorySectorMapping(category_desc="Agriculture locale", activity_sector_id=sector.id))

            batch = ImportBatch(
                batch_type=ImportBatchType.CURRENT_STATE,
                period=None,
                snapshot_date=date(2026, 7, 15),
                file_name="taeg-update-audit.xlsx",
            )
            db.add(batch)
            db.flush()

            db.add(
                LoanRaw(
                    contract_no="TAEG-UPDATE-001",
                    client_name="Client",
                    client_first_name="Audit",
                    client_id="CLI-TAEG-UPD-1",
                    agency_name="Agence ACM",
                    agent_name="Agent ACM",
                    category_desc="Agriculture locale",
                    disbursement_amount=Decimal("2000.00"),
                    disbursement_date=date(2026, 7, 10),
                    teg_rate=Decimal("15.00"),
                    principal_outstanding=Decimal("1000.00"),
                    principal_due=Decimal("100.00"),
                    total_scheduled_amount=Decimal("100.00"),
                    days_overdue=0,
                    total_due=Decimal("0.00"),
                    status="active",
                    snapshot_date=batch.snapshot_date,
                    import_batch_id=batch.id,
                )
            )
            version = create_acm_rate_version(
                db,
                created_by_user_id=None,
                effective_start_date=date(2026, 1, 1),
                is_open_ended=True,
                details=[{"sector_id": sector.id, "acm_rate": Decimal("0.1200")}],
                comment="Bloc initial",
                created_at=datetime(2026, 1, 3, 8, 0),
            )
            create_taeg_daily_snapshot_for_batch(db, batch.id)
            db.commit()
            version_id = version.id
            self.assertIsNotNone(
                db.scalar(select(TaegDailySnapshot).where(TaegDailySnapshot.snapshot_date == date(2026, 7, 15)))
            )

        update_response = self.client.put(
            f"/api/v1/taeg/acm-rate-versions/{version_id}",
            json={
                "effective_start_date": "2026-01-01",
                "is_open_ended": True,
                "comment": "Bloc corrige",
                "modification_comment": "Correction du taux Agriculture",
                "confirm_impact": True,
                "details": [
                    {"sector_id": 1, "acm_rate": 16.0},
                ],
            },
        )
        self.assertEqual(update_response.status_code, 200, update_response.text)
        payload = update_response.json()
        self.assertEqual(payload["comment"], "Bloc corrige")
        self.assertEqual(Decimal(payload["details"][0]["acm_rate"]), Decimal("16.0000"))
        self.assertEqual(payload["affected_snapshot_count"], 1)
        self.assertTrue(payload["may_affect_existing_results"])
        self.assertEqual(len(payload["audit_entries"]), 1)
        self.assertEqual(payload["audit_entries"][0]["comment"], "Correction du taux Agriculture")

        with SessionLocal() as db:
            self.assertEqual(db.scalar(select(func.count(AcmRateVersionAudit.id))), 1)
            snapshot = db.scalar(select(TaegDailySnapshot).where(TaegDailySnapshot.snapshot_date == date(2026, 7, 15)))
            self.assertIsNotNone(snapshot)
            self.assertEqual(Decimal(snapshot.acm_rate), Decimal("16.0000"))

    def test_taeg_update_acm_rate_version_requires_confirmation_when_results_exist(self) -> None:
        self._seed_super_admin()
        login_response = self._login(self.super_admin_email, self.super_admin_password)
        self.assertEqual(login_response.status_code, 200, login_response.text)

        with SessionLocal() as db:
            sector = ActivitySector(name="Commerce", acm_rate=Decimal("0.1200"))
            db.add(sector)
            db.flush()
            db.add(CategorySectorMapping(category_desc="Commerce local", activity_sector_id=sector.id))
            batch = ImportBatch(
                batch_type=ImportBatchType.CURRENT_STATE,
                period=None,
                snapshot_date=date(2026, 7, 15),
                file_name="taeg-confirmation.xlsx",
            )
            db.add(batch)
            db.flush()
            db.add(
                LoanRaw(
                    contract_no="TAEG-CONFIRM-001",
                    client_name="Client",
                    client_first_name="Confirm",
                    client_id="CLI-TAEG-CONF-1",
                    agency_name="Agence ACM",
                    agent_name="Agent ACM",
                    category_desc="Commerce local",
                    disbursement_amount=Decimal("1500.00"),
                    disbursement_date=date(2026, 7, 12),
                    teg_rate=Decimal("11.00"),
                    principal_outstanding=Decimal("900.00"),
                    principal_due=Decimal("90.00"),
                    total_scheduled_amount=Decimal("90.00"),
                    days_overdue=0,
                    total_due=Decimal("0.00"),
                    status="active",
                    snapshot_date=batch.snapshot_date,
                    import_batch_id=batch.id,
                )
            )
            version = create_acm_rate_version(
                db,
                created_by_user_id=None,
                effective_start_date=date(2026, 1, 1),
                is_open_ended=True,
                details=[{"sector_id": sector.id, "acm_rate": Decimal("0.1200")}],
                comment="Bloc initial",
                created_at=datetime(2026, 1, 3, 8, 0),
            )
            create_taeg_daily_snapshot_for_batch(db, batch.id)
            db.commit()
            version_id = version.id

        response = self.client.put(
            f"/api/v1/taeg/acm-rate-versions/{version_id}",
            json={
                "effective_start_date": "2026-01-01",
                "is_open_ended": True,
                "comment": "Bloc sans confirmation",
                "modification_comment": "Essai sans confirmation",
                "confirm_impact": False,
                "details": [
                    {"sector_id": 1, "acm_rate": 13.5},
                ],
            },
        )
        self.assertEqual(response.status_code, 400, response.text)
        self.assertIn("Cette modification peut changer les resultats TAEG", response.text)

    def test_taeg_update_acm_rate_version_rejects_overlap(self) -> None:
        with SessionLocal() as db:
            sector = ActivitySector(name="Services", acm_rate=Decimal("0.1200"))
            db.add(sector)
            db.flush()

            first = create_acm_rate_version(
                db,
                created_by_user_id=None,
                effective_start_date=date(2026, 1, 1),
                effective_end_date=date(2026, 6, 30),
                details=[{"sector_id": sector.id, "acm_rate": Decimal("0.1200")}],
                comment="S1",
                created_at=datetime(2026, 1, 3, 8, 0),
            )
            create_acm_rate_version(
                db,
                created_by_user_id=None,
                effective_start_date=date(2026, 7, 1),
                is_open_ended=True,
                details=[{"sector_id": sector.id, "acm_rate": Decimal("0.1500")}],
                comment="S2",
                created_at=datetime(2026, 7, 3, 8, 0),
            )

            with self.assertRaisesRegex(ValueError, "chevauche"):
                update_acm_rate_version(
                    db,
                    version_id=first.id,
                    modified_by_user_id=None,
                    effective_start_date=date(2026, 1, 1),
                    effective_end_date=date(2026, 7, 15),
                    is_open_ended=False,
                    comment="Bloc overlap",
                    modification_comment="Chevauchement volontaire",
                    confirm_impact=True,
                    details=[{"sector_id": sector.id, "acm_rate": Decimal("0.1400")}],
                )

    def test_taeg_update_acm_rate_version_requires_super_admin(self) -> None:
        self._create_user(
            email="admin.acm@microcred.com.tn",
            password="AdminAcm!2026",
            role=UserRole.ADMIN,
            full_name="Admin ACM",
        )
        login_response = self._login("admin.acm@microcred.com.tn", "AdminAcm!2026")
        self.assertEqual(login_response.status_code, 200, login_response.text)

        with SessionLocal() as db:
            sector = ActivitySector(name="Habitat", acm_rate=Decimal("0.1200"))
            db.add(sector)
            db.flush()
            version = create_acm_rate_version(
                db,
                created_by_user_id=None,
                effective_start_date=date(2026, 1, 1),
                is_open_ended=True,
                details=[{"sector_id": sector.id, "acm_rate": Decimal("0.1200")}],
                comment="Bloc Habitat",
                created_at=datetime(2026, 1, 3, 8, 0),
            )
            db.commit()
            version_id = version.id

        response = self.client.put(
            f"/api/v1/taeg/acm-rate-versions/{version_id}",
            json={
                "effective_start_date": "2026-01-01",
                "is_open_ended": True,
                "comment": "Bloc non autorise",
                "modification_comment": "Tentative admin",
                "confirm_impact": True,
                "details": [
                    {"sector_id": 1, "acm_rate": 12.0},
                ],
            },
        )
        self.assertEqual(response.status_code, 403, response.text)

    def test_taeg_publish_acm_rate_version_requires_complete_block(self) -> None:
        self._seed_super_admin()
        login_response = self._login(self.super_admin_email, self.super_admin_password)
        self.assertEqual(login_response.status_code, 200, login_response.text)

        first_sector = self.client.post(
            "/api/v1/taeg/sectors",
            json={"name": "Agriculture", "acm_rate": 0},
        )
        second_sector = self.client.post(
            "/api/v1/taeg/sectors",
            json={"name": "Commerce", "acm_rate": 0},
        )
        self.assertEqual(first_sector.status_code, 200, first_sector.text)
        self.assertEqual(second_sector.status_code, 200, second_sector.text)

        response = self.client.post(
            "/api/v1/taeg/acm-rate-versions",
            json={
                "effective_start_date": "2026-01-01",
                "effective_end_date": "2026-06-30",
                "details": [
                    {"sector_id": first_sector.json()["id"], "acm_rate": 12.5},
                ],
            },
        )
        self.assertEqual(response.status_code, 400, response.text)
        self.assertIn("Tous les secteurs actifs doivent etre renseignes", response.text)

    def test_resolve_acm_rate_map_uses_effective_version_for_requested_date(self) -> None:
        actor_id = self._create_user(
            email="actor.acm@microcred.com.tn",
            password="ActorAcmPassword!2026",
            role=UserRole.SUPER_ADMIN,
            full_name="Actor ACM",
        )

        with SessionLocal() as db:
            sector = ActivitySector(name="Commerce", acm_rate=Decimal("0.1200"))
            db.add(sector)
            db.flush()

            create_acm_rate_version(
                db,
                created_by_user_id=actor_id,
                effective_start_date=date(2026, 1, 1),
                effective_end_date=date(2026, 6, 30),
                details=[{"sector_id": sector.id, "acm_rate": Decimal("0.1200")}],
                comment="S1",
                created_at=datetime(2026, 3, 5, 9, 0),
            )
            create_acm_rate_version(
                db,
                created_by_user_id=actor_id,
                effective_start_date=date(2026, 7, 1),
                is_open_ended=True,
                details=[{"sector_id": sector.id, "acm_rate": Decimal("0.1850")}],
                comment="S2",
                created_at=datetime(2026, 7, 7, 10, 0),
            )
            db.commit()
            sector_id = sector.id

            march_map = resolve_acm_rate_map(db, date(2026, 3, 31))
            july_map = resolve_acm_rate_map(db, date(2026, 7, 31))

            self.assertEqual(Decimal(march_map[sector_id]["acm_rate"]), Decimal("0.1200"))
            self.assertEqual(Decimal(july_map[sector_id]["acm_rate"]), Decimal("0.1850"))

    def test_arrears_report_includes_principal_schedule_column(self) -> None:
        from app.api.v1.reports import _portfolio_report_definition

        _title, _subtitle, columns, _predicate = _portfolio_report_definition("arrears_list", None)
        labels = [label for _key, label, _kind in columns]

        self.assertIn("Échéance principal", labels)
        self.assertLess(labels.index("TOTAL DUE AMT"), labels.index("Échéance principal"))
        self.assertLess(labels.index("Échéance principal"), labels.index("JOURS DE RETARD"))


if __name__ == "__main__":
    unittest.main()
