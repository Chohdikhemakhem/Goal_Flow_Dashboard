from __future__ import annotations

import json
import re
import unicodedata
from datetime import date, datetime
from io import BytesIO

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import delete, func, select
from sqlalchemy.orm import Session

from app.api.deps import get_current_user, pagination, require_roles
from app.api.v1.reports import (
    ALT_ROW_FILL,
    HEADER_BG,
    HEADER_FONT,
    THIN_BORDER,
    TITLE_FILL,
    TITLE_FONT,
    _fit_column_widths,
)
from app.db.session import get_db
from app.models.entities import ActivitySector, Agency, Agent, CategorySectorMapping, User
from app.models.enums import UserRole
from app.schemas.common import Message, Page
from app.schemas.domain import (
    AcmRateVersionCreate,
    AcmRateVersionRead,
    AcmRateVersionSummaryRead,
    AcmRateVersionUpdate,
    ActivitySectorCreate,
    ActivitySectorRead,
    ActivitySectorUpdate,
    TaegMonthlyHistoryPeriodRead,
    TaegMonthlyHistoryRead,
    CategorySectorMappingRead,
    CategorySectorMappingUpdate,
    TaegCreditDetailRead,
    TaegDashboardRead,
    TaegPeriodRead,
)
from app.services.taeg import (
    _filter_taeg_summary_rows,
    _sort_taeg_summary_rows,
    format_taeg_period_name,
    format_taeg_used_months_label,
    list_category_sector_mappings,
    list_taeg_periods,
    summarize_taeg_rows,
    taeg_credit_details,
    taeg_dashboard,
)
from app.services.taeg_history import (
    create_acm_rate_version,
    get_acm_rate_version_payload,
    get_taeg_monthly_history,
    list_acm_rate_versions,
    list_taeg_monthly_history_periods,
    update_acm_rate_version,
)
from app.services.committee_access import is_committee_member
from app.services.selection import normalize_agency_ids

router = APIRouter(
    prefix="/taeg",
    tags=["taeg"],
    dependencies=[Depends(require_roles([UserRole.SUPER_ADMIN, UserRole.ADMIN, UserRole.COMMITTEE_MEMBER, UserRole.AGENCY_MANAGER, UserRole.PORTFOLIO_MANAGER, UserRole.REGIONAL_MANAGER_NORD, UserRole.REGIONAL_MANAGER_SUD]))],
)

TAEG_EXPORT_COLUMNS = [
    ("sector_name", "Secteur d'activite", "text"),
    ("credits_count", "Nombre de credits", "int"),
    ("disbursement_amount", "Montant decaisse", "money"),
    ("taeg_calculated", "TAEG Calcule", "money"),
    ("taeg_weighted_rate", "TAEG Pondere (%)", "rate"),
    ("acm_rate", "Taux ACM (%)", "rate"),
    ("status_label", "Statut", "text"),
]


def _safe_export_fragment(value: str) -> str:
    normalized = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode("ascii")
    cleaned = re.sub(r"[^A-Za-z0-9]+", "", normalized)
    return cleaned or "TAEG"


def _taeg_export_filename(period_name: str) -> str:
    timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M")
    return f"TAEG_{_safe_export_fragment(period_name)}_{timestamp}.xlsx"


def _parse_table_filters(raw_value: str | None) -> dict:
    if not raw_value:
        return {}
    try:
        payload = json.loads(raw_value)
    except json.JSONDecodeError as exc:
        raise ValueError("Filtres de colonnes TAEG invalides.") from exc
    if not isinstance(payload, dict):
        raise ValueError("Filtres de colonnes TAEG invalides.")
    return payload


def _apply_taeg_excel_format(cell, value_type: str) -> None:
    if value_type == "money":
        cell.number_format = "#,##0.00"
        cell.alignment = Alignment(horizontal="right", vertical="center")
    elif value_type == "int":
        cell.number_format = "0"
        cell.alignment = Alignment(horizontal="right", vertical="center")
    elif value_type == "rate":
        cell.number_format = '0.00"%"'
        cell.alignment = Alignment(horizontal="right", vertical="center")
    else:
        cell.alignment = Alignment(horizontal="left", vertical="center")


def _build_taeg_excel_buffer(
    *,
    title: str,
    period_name: str,
    segment_label: str,
    segment_block_label: str,
    generated_at: datetime,
    generated_by: str,
    agency_name: str,
    agent_name: str,
    used_months: str,
    rows: list[dict],
    summary: dict,
) -> BytesIO:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "TAEG"

    max_col = len(TAEG_EXPORT_COLUMNS)
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max_col)
    sheet["A1"] = title
    sheet["A1"].fill = TITLE_FILL
    sheet["A1"].font = TITLE_FONT
    sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 28

    metadata = [
        ("Onglet", period_name),
        ("Bloc ACM", segment_block_label),
        ("Periode analysee", segment_label),
        ("Generation", generated_at.strftime("%d/%m/%Y %H:%M")),
        ("Genere par", generated_by),
        ("Agence", agency_name),
        ("Agent", agent_name),
        ("Mois utilises", used_months),
    ]

    meta_label_fill = PatternFill("solid", fgColor="E2E8F0")
    meta_label_font = Font(color="0F172A", bold=True)
    meta_value_fill = PatternFill("solid", fgColor="F8FAFC")

    meta_row = 3
    for label, value in metadata:
        label_cell = sheet.cell(row=meta_row, column=1, value=label)
        label_cell.fill = meta_label_fill
        label_cell.font = meta_label_font
        label_cell.border = THIN_BORDER
        label_cell.alignment = Alignment(horizontal="left", vertical="center")
        sheet.merge_cells(start_row=meta_row, start_column=2, end_row=meta_row, end_column=max_col)
        value_cell = sheet.cell(row=meta_row, column=2, value=value)
        value_cell.fill = meta_value_fill
        value_cell.border = THIN_BORDER
        value_cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        meta_row += 1

    header_row = meta_row + 1
    for col_index, (_key, label, _value_type) in enumerate(TAEG_EXPORT_COLUMNS, start=1):
        cell = sheet.cell(row=header_row, column=col_index, value=label)
        cell.fill = HEADER_BG
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER
    sheet.row_dimensions[header_row].height = 24

    data_row = header_row + 1
    chart_data_start_row = data_row
    for row_index, payload in enumerate(rows, start=data_row):
        for col_index, (key, _label, value_type) in enumerate(TAEG_EXPORT_COLUMNS, start=1):
            value = payload.get(key)
            cell = sheet.cell(row=row_index, column=col_index, value=value)
            cell.border = THIN_BORDER
            if row_index % 2 == 0:
                cell.fill = ALT_ROW_FILL
            if key == "status_label":
                if payload.get("status") == "conforme":
                    cell.font = Font(color="166534", bold=True)
                elif payload.get("status") == "non_couvert":
                    cell.font = Font(color="9A3412", bold=True)
                else:
                    cell.font = Font(color="991B1B", bold=True)
            _apply_taeg_excel_format(cell, value_type)

    if not rows:
        empty_row = data_row
        sheet.merge_cells(start_row=empty_row, start_column=1, end_row=empty_row, end_column=max_col)
        empty_cell = sheet.cell(row=empty_row, column=1, value="Aucune donnee disponible pour cette selection.")
        empty_cell.border = THIN_BORDER
        empty_cell.alignment = Alignment(horizontal="left", vertical="center")
        empty_cell.font = Font(color="64748B", italic=True)

    chart_section_title_row = max(data_row + len(rows), data_row + 1) + 2
    chart_anchor_row = chart_section_title_row + 1
    summary_start = chart_anchor_row + 18

    sheet.merge_cells(start_row=chart_section_title_row, start_column=1, end_row=chart_section_title_row, end_column=max_col)
    chart_title = sheet.cell(
        row=chart_section_title_row,
        column=1,
        value="Diagramme - Comparaison TAEG pondere / Taux ACM",
    )
    chart_title.fill = TITLE_FILL
    chart_title.font = Font(color="FFFFFF", bold=True, size=12)
    chart_title.alignment = Alignment(horizontal="left", vertical="center")

    if rows:
        chart = BarChart()
        chart.type = "col"
        chart.style = 10
        chart.grouping = "clustered"
        chart.overlap = 0
        chart.title = "Comparaison TAEG pondere / Taux ACM"
        chart.x_axis.title = "Secteurs d'activite"
        chart.y_axis.title = "Pourcentage"
        chart.height = 9.5
        chart.width = 18
        chart.legend.position = "r"
        chart.varyColors = False

        weighted_data = Reference(
            sheet,
            min_col=5,
            min_row=header_row,
            max_row=header_row + len(rows),
        )
        acm_data = Reference(
            sheet,
            min_col=6,
            min_row=header_row,
            max_row=header_row + len(rows),
        )
        categories = Reference(
            sheet,
            min_col=1,
            min_row=chart_data_start_row,
            max_row=chart_data_start_row + len(rows) - 1,
        )

        chart.add_data(weighted_data, titles_from_data=True)
        chart.add_data(acm_data, titles_from_data=True)
        chart.set_categories(categories)
        chart.dLbls = DataLabelList()
        chart.dLbls.showVal = True
        chart.dLbls.position = "outEnd"
        if len(chart.series) >= 1:
            chart.series[0].graphicalProperties.solidFill = "2563EB"
        if len(chart.series) >= 2:
            chart.series[1].graphicalProperties.solidFill = "0F766E"
        sheet.add_chart(chart, f"A{chart_anchor_row}")
    else:
        sheet.merge_cells(start_row=chart_anchor_row, start_column=1, end_row=chart_anchor_row + 2, end_column=max_col)
        chart_empty = sheet.cell(
            row=chart_anchor_row,
            column=1,
            value="Aucun diagramme disponible pour cette selection.",
        )
        chart_empty.border = THIN_BORDER
        chart_empty.alignment = Alignment(horizontal="left", vertical="center")
        chart_empty.font = Font(color="64748B", italic=True)

    sheet.merge_cells(start_row=summary_start, start_column=1, end_row=summary_start, end_column=max_col)
    summary_title = sheet.cell(row=summary_start, column=1, value="Synthese")
    summary_title.fill = TITLE_FILL
    summary_title.font = Font(color="FFFFFF", bold=True, size=12)
    summary_title.alignment = Alignment(horizontal="left", vertical="center")

    summary_rows = [
        ("Nombre total de secteurs", summary["total_sectors"], "int"),
        ("Nombre de secteurs conformes", summary["compliant_sectors"], "int"),
        ("Nombre de secteurs non conformes", summary["non_compliant_sectors"], "int"),
        ("Montant total decaisse", summary["total_disbursement"], "money"),
        ("Nombre total de credits", summary["total_credits"], "int"),
        ("Moyenne ponderee globale du TAEG", summary["global_weighted_rate"], "rate"),
    ]
    summary_row_index = summary_start + 1
    for label, value, value_type in summary_rows:
        label_cell = sheet.cell(row=summary_row_index, column=1, value=label)
        label_cell.fill = meta_label_fill
        label_cell.font = meta_label_font
        label_cell.border = THIN_BORDER
        label_cell.alignment = Alignment(horizontal="left", vertical="center")
        value_cell = sheet.cell(row=summary_row_index, column=2, value=value)
        value_cell.fill = meta_value_fill
        value_cell.border = THIN_BORDER
        value_cell.font = Font(bold=True)
        _apply_taeg_excel_format(value_cell, value_type)
        sheet.merge_cells(start_row=summary_row_index, start_column=3, end_row=summary_row_index, end_column=max_col)
        filler_cell = sheet.cell(row=summary_row_index, column=3, value="")
        filler_cell.fill = meta_value_fill
        filler_cell.border = THIN_BORDER
        summary_row_index += 1

    sheet.freeze_panes = f"A{header_row + 1}"
    _fit_column_widths(sheet, header_row=header_row)
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer


@router.get(
    "/periods",
    response_model=list[TaegPeriodRead],
    dependencies=[Depends(require_roles([UserRole.SUPER_ADMIN, UserRole.ADMIN, UserRole.COMMITTEE_MEMBER, UserRole.AGENCY_MANAGER, UserRole.PORTFOLIO_MANAGER, UserRole.REGIONAL_MANAGER_NORD, UserRole.REGIONAL_MANAGER_SUD]))],
)
def get_taeg_periods(
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    items = list_taeg_periods(db)
    if is_committee_member(user):
        return [item for item in items if not item.get("is_current")]
    return items


@router.get(
    "/dashboard",
    response_model=TaegDashboardRead,
    dependencies=[Depends(require_roles([UserRole.SUPER_ADMIN, UserRole.ADMIN, UserRole.COMMITTEE_MEMBER, UserRole.AGENCY_MANAGER, UserRole.PORTFOLIO_MANAGER, UserRole.REGIONAL_MANAGER_NORD, UserRole.REGIONAL_MANAGER_SUD]))],
)
def get_taeg_dashboard(
    period_key: str | None = None,
    period_keys: str | None = None,
    agency_id: str | None = None,
    agent_id: int | None = None,
    acm_block_id: int | None = None,
    period_start: date | None = None,
    period_end: date | None = None,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    if is_committee_member(user) and (period_key or "current") == "current":
        raise HTTPException(status_code=403, detail={"code": "current_state_forbidden", "message": "Le profil Membre comité n'a pas accès à l'état actuel TAEG."})
    try:
        return taeg_dashboard(
            db,
            user=user,
            period_key=period_key,
            period_keys=period_keys,
            agency_id=agency_id,
            agent_id=agent_id,
            acm_block_id=acm_block_id,
            period_start=period_start,
            period_end=period_end,
        )
    except ValueError as exc:
        raise HTTPException(status_code=400, detail={"code": "taeg_dashboard_error", "message": str(exc)}) from exc


@router.get(
    "/export.xlsx",
    dependencies=[Depends(require_roles([UserRole.SUPER_ADMIN, UserRole.ADMIN, UserRole.COMMITTEE_MEMBER, UserRole.AGENCY_MANAGER, UserRole.PORTFOLIO_MANAGER, UserRole.REGIONAL_MANAGER_NORD, UserRole.REGIONAL_MANAGER_SUD]))],
)
def export_taeg_excel(
    period_key: str | None = None,
    period_keys: str | None = None,
    agency_id: str | None = None,
    agent_id: int | None = None,
    acm_block_id: int | None = None,
    period_start: date | None = None,
    period_end: date | None = None,
    table_filters: str | None = None,
    sort_key: str | None = None,
    sort_direction: str | None = None,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    if is_committee_member(user) and (period_key or "current") == "current":
        raise HTTPException(status_code=403, detail={"code": "current_state_forbidden", "message": "Le profil Membre comité n'a pas accès à l'état actuel TAEG."})
    try:
        dashboard = taeg_dashboard(
            db,
            user=user,
            period_key=period_key,
            period_keys=period_keys,
            agency_id=agency_id,
            agent_id=agent_id,
            acm_block_id=acm_block_id,
            period_start=period_start,
            period_end=period_end,
        )
        parsed_table_filters = _parse_table_filters(table_filters)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail={"code": "taeg_export_error", "message": str(exc)}) from exc

    period_name = format_taeg_period_name(dashboard["period"])
    active_segment = dashboard.get("active_acm_segment") or {}
    segment_label = (
        f"{active_segment['range_start'].strftime('%d/%m/%Y')} -> {active_segment['range_end'].strftime('%d/%m/%Y')}"
        if active_segment.get("range_start") and active_segment.get("range_end")
        else "Selection complete"
    )
    segment_block_label = active_segment.get("block_label") or "Aucun bloc ACM"
    used_months = format_taeg_used_months_label(dashboard.get("used_periods", []))
    selected_agency_ids = normalize_agency_ids(agency_id=agency_id)
    selected_agencies = (
        db.scalars(select(Agency).where(Agency.id.in_(selected_agency_ids)).order_by(Agency.name)).all()
        if selected_agency_ids
        else []
    )
    selected_agent = db.get(Agent, agent_id) if agent_id is not None else None
    if selected_agencies:
        agency_name = " + ".join(agency.name for agency in selected_agencies)
    else:
        agency_name = "Toutes les agences" if not selected_agency_ids else "Agence inconnue"
    agent_name = selected_agent.name if selected_agent else ("Agent inconnu" if agent_id is not None else "Tous les agents")

    filtered_rows = _filter_taeg_summary_rows(dashboard["rows"], parsed_table_filters)
    sorted_rows = _sort_taeg_summary_rows(filtered_rows, sort_key, sort_direction)

    export_rows = []
    for row in sorted_rows:
        status_label = (
            "Conforme"
            if row["status"] == "conforme"
            else "Non Conforme"
            if row["status"] == "non_conforme"
            else "Non couvert ACM"
        )
        export_rows.append(
            {
                **row,
                "status_label": status_label,
            }
        )
    summary = summarize_taeg_rows(export_rows)
    generated_at = datetime.now()
    generated_by = f"{user.full_name} ({user.email})" if user.full_name else user.email
    buffer = _build_taeg_excel_buffer(
        title="MicroCred - Rapport TAEG",
        period_name=period_name,
        segment_label=segment_label,
        segment_block_label=segment_block_label,
        generated_at=generated_at,
        generated_by=generated_by,
        agency_name=agency_name,
        agent_name=agent_name,
        used_months=used_months,
        rows=export_rows,
        summary=summary,
    )
    filename = _taeg_export_filename(period_name)
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get(
    "/details",
    response_model=Page[TaegCreditDetailRead],
    dependencies=[Depends(require_roles([UserRole.SUPER_ADMIN, UserRole.ADMIN, UserRole.COMMITTEE_MEMBER, UserRole.AGENCY_MANAGER, UserRole.PORTFOLIO_MANAGER, UserRole.REGIONAL_MANAGER_NORD, UserRole.REGIONAL_MANAGER_SUD]))],
)
def get_taeg_details(
    period_key: str | None = None,
    period_keys: str | None = None,
    agency_id: str | None = None,
    agent_id: int | None = None,
    sector_id: int | None = None,
    acm_block_id: int | None = None,
    period_start: date | None = None,
    period_end: date | None = None,
    table_filters: str | None = None,
    sort_key: str | None = None,
    sort_direction: str | None = None,
    page: tuple[int, int] = Depends(pagination),
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    limit, offset = page
    if is_committee_member(user) and (period_key or "current") == "current":
        raise HTTPException(status_code=403, detail={"code": "current_state_forbidden", "message": "Le profil Membre comité n'a pas accès à l'état actuel TAEG."})
    try:
        parsed_table_filters = _parse_table_filters(table_filters)
        _period, rows, total = taeg_credit_details(
            db,
            user=user,
            period_key=period_key,
            period_keys=period_keys,
            agency_id=agency_id,
            agent_id=agent_id,
            sector_id=sector_id,
            acm_block_id=acm_block_id,
            period_start=period_start,
            period_end=period_end,
            column_filters=parsed_table_filters,
            sort_key=sort_key,
            sort_direction=sort_direction,
            limit=limit,
            offset=offset,
        )
    except ValueError as exc:
        raise HTTPException(status_code=400, detail={"code": "taeg_details_error", "message": str(exc)}) from exc
    return Page(items=rows, total=total, limit=limit, offset=offset)


@router.get(
    "/monthly-history/periods",
    response_model=list[TaegMonthlyHistoryPeriodRead],
    dependencies=[Depends(require_roles([UserRole.SUPER_ADMIN, UserRole.ADMIN, UserRole.COMMITTEE_MEMBER, UserRole.AGENCY_MANAGER, UserRole.PORTFOLIO_MANAGER, UserRole.REGIONAL_MANAGER_NORD, UserRole.REGIONAL_MANAGER_SUD]))],
)
def get_taeg_monthly_history_periods(
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    if is_committee_member(user):
        raise HTTPException(status_code=403, detail={"code": "daily_history_forbidden", "message": "Le profil Membre comité ne peut pas consulter l'historique journalier TAEG."})
    return list_taeg_monthly_history_periods(db)


@router.get(
    "/monthly-history",
    response_model=TaegMonthlyHistoryRead,
    dependencies=[Depends(require_roles([UserRole.SUPER_ADMIN, UserRole.ADMIN, UserRole.COMMITTEE_MEMBER, UserRole.AGENCY_MANAGER, UserRole.PORTFOLIO_MANAGER, UserRole.REGIONAL_MANAGER_NORD, UserRole.REGIONAL_MANAGER_SUD]))],
)
def get_taeg_monthly_history_endpoint(
    period: str,
    agency_id: str | None = None,
    agent_id: int | None = None,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    if is_committee_member(user):
        raise HTTPException(status_code=403, detail={"code": "daily_history_forbidden", "message": "Le profil Membre comité ne peut pas consulter l'historique journalier TAEG."})
    try:
        return get_taeg_monthly_history(
            db,
            period=period,
            agency_id=agency_id,
            agent_id=agent_id,
            user=user,
        )
    except ValueError as exc:
        raise HTTPException(status_code=400, detail={"code": "taeg_monthly_history_error", "message": str(exc)}) from exc


@router.get(
    "/acm-rate-versions",
    response_model=list[AcmRateVersionSummaryRead],
    dependencies=[Depends(require_roles([UserRole.SUPER_ADMIN,UserRole.SUPPORT]))],
)
def get_acm_rate_versions(db: Session = Depends(get_db)):
    return list_acm_rate_versions(db)


@router.post(
    "/acm-rate-versions",
    response_model=AcmRateVersionRead,
    dependencies=[Depends(require_roles([UserRole.SUPER_ADMIN,UserRole.SUPPORT]))],
)
def publish_acm_rate_version(
    payload: AcmRateVersionCreate,
    db: Session = Depends(get_db),
    actor: User = Depends(get_current_user),
):
    try:
        version = create_acm_rate_version(
            db,
            created_by_user_id=actor.id if actor else None,
            effective_start_date=payload.effective_start_date,
            effective_end_date=payload.effective_end_date,
            is_open_ended=payload.is_open_ended,
            comment=payload.comment,
            details=[detail.model_dump() for detail in payload.details],
        )
        db.commit()
    except ValueError as exc:
        db.rollback()
        raise HTTPException(status_code=400, detail={"code": "acm_rate_version_invalid", "message": str(exc)}) from exc
    payload_data = get_acm_rate_version_payload(db, version.id)
    if not payload_data:
        raise HTTPException(status_code=500, detail={"code": "acm_rate_version_missing", "message": "Version ACM introuvable apres creation."})
    return payload_data


@router.put(
    "/acm-rate-versions/{version_id}",
    response_model=AcmRateVersionRead,
    dependencies=[Depends(require_roles([UserRole.SUPER_ADMIN,UserRole.SUPPORT]))],
)
def update_acm_rate_version_endpoint(
    version_id: int,
    payload: AcmRateVersionUpdate,
    db: Session = Depends(get_db),
    actor: User = Depends(get_current_user),
):
    try:
        version, _recalculated_batch_ids = update_acm_rate_version(
            db,
            version_id=version_id,
            modified_by_user_id=actor.id if actor else None,
            effective_start_date=payload.effective_start_date,
            effective_end_date=payload.effective_end_date,
            is_open_ended=payload.is_open_ended,
            comment=payload.comment,
            modification_comment=payload.modification_comment,
            confirm_impact=payload.confirm_impact,
            details=[detail.model_dump() for detail in payload.details],
        )
        db.commit()
    except ValueError as exc:
        db.rollback()
        raise HTTPException(status_code=400, detail={"code": "acm_rate_version_update_invalid", "message": str(exc)}) from exc
    payload_data = get_acm_rate_version_payload(db, version.id)
    if not payload_data:
        raise HTTPException(status_code=500, detail={"code": "acm_rate_version_missing", "message": "Version ACM introuvable apres modification."})
    return payload_data


@router.get(
    "/acm-rate-versions/{version_id}",
    response_model=AcmRateVersionRead,
    dependencies=[Depends(require_roles([UserRole.SUPER_ADMIN,UserRole.SUPPORT]))],
)
def get_acm_rate_version(version_id: int, db: Session = Depends(get_db)):
    payload = get_acm_rate_version_payload(db, version_id)
    if not payload:
        raise HTTPException(status_code=404, detail={"code": "acm_rate_version_not_found", "message": "Version ACM introuvable."})
    return payload


@router.get(
    "/sectors",
    response_model=list[ActivitySectorRead],
    dependencies=[Depends(require_roles([UserRole.SUPER_ADMIN,UserRole.SUPPORT]))],
)
def list_sectors(db: Session = Depends(get_db)):
    return db.scalars(select(ActivitySector).order_by(ActivitySector.name.asc())).all()


@router.post(
    "/sectors",
    response_model=ActivitySectorRead,
    dependencies=[Depends(require_roles([UserRole.SUPER_ADMIN,UserRole.SUPPORT]))],
)
def create_sector(
    payload: ActivitySectorCreate,
    db: Session = Depends(get_db),
    actor: User = Depends(get_current_user),
):
    existing = db.scalar(select(ActivitySector).where(func.lower(ActivitySector.name) == payload.name.strip().lower()))
    if existing:
        raise HTTPException(status_code=409, detail={"code": "duplicate_sector", "message": "Ce secteur existe deja."})
    sector = ActivitySector(name=payload.name.strip(), acm_rate=payload.acm_rate)
    db.add(sector)
    db.commit()
    db.refresh(sector)
    return sector


@router.put(
    "/sectors/{sector_id}",
    response_model=ActivitySectorRead,
    dependencies=[Depends(require_roles([UserRole.SUPER_ADMIN,UserRole.SUPPORT]))],
)
def update_sector(
    sector_id: int,
    payload: ActivitySectorUpdate,
    db: Session = Depends(get_db),
    actor: User = Depends(get_current_user),
):
    sector = db.get(ActivitySector, sector_id)
    if not sector:
        raise HTTPException(status_code=404, detail={"code": "sector_not_found", "message": "Secteur introuvable."})
    duplicate = db.scalar(
        select(ActivitySector).where(
            func.lower(ActivitySector.name) == payload.name.strip().lower(),
            ActivitySector.id != sector_id,
        )
    )
    if duplicate:
        raise HTTPException(status_code=409, detail={"code": "duplicate_sector", "message": "Ce secteur existe deja."})
    sector.name = payload.name.strip()
    sector.acm_rate = payload.acm_rate
    db.commit()
    db.refresh(sector)
    return sector


@router.delete(
    "/sectors/{sector_id}",
    response_model=Message,
    dependencies=[Depends(require_roles([UserRole.SUPER_ADMIN,UserRole.SUPPORT]))],
)
def delete_sector(
    sector_id: int,
    db: Session = Depends(get_db),
    actor: User = Depends(get_current_user),
):
    sector = db.get(ActivitySector, sector_id)
    if not sector:
        raise HTTPException(status_code=404, detail={"code": "sector_not_found", "message": "Secteur introuvable."})
    db.execute(delete(CategorySectorMapping).where(CategorySectorMapping.activity_sector_id == sector_id))
    db.delete(sector)
    db.commit()
    return Message(message="Secteur supprime.")


@router.get(
    "/category-mappings",
    response_model=list[CategorySectorMappingRead],
    dependencies=[Depends(require_roles([UserRole.SUPER_ADMIN,UserRole.SUPPORT]))],
)
def get_category_mappings(db: Session = Depends(get_db)):
    return list_category_sector_mappings(db)


@router.put(
    "/category-mappings",
    response_model=CategorySectorMappingRead,
    dependencies=[Depends(require_roles([UserRole.SUPER_ADMIN,UserRole.SUPPORT]))],
)
def upsert_category_mapping(payload: CategorySectorMappingUpdate, db: Session = Depends(get_db)):
    sector = db.get(ActivitySector, payload.activity_sector_id)
    if not sector:
        raise HTTPException(status_code=404, detail={"code": "sector_not_found", "message": "Secteur introuvable."})
    normalized_category = payload.category_desc.strip()
    if not normalized_category:
        raise HTTPException(status_code=400, detail={"code": "invalid_category", "message": "CATEGORY_DESC invalide."})
    mapping = db.scalar(select(CategorySectorMapping).where(CategorySectorMapping.category_desc == normalized_category))
    if mapping:
        mapping.activity_sector_id = payload.activity_sector_id
    else:
        mapping = CategorySectorMapping(
            category_desc=normalized_category,
            activity_sector_id=payload.activity_sector_id,
        )
        db.add(mapping)
    db.commit()
    db.refresh(mapping)
    return CategorySectorMappingRead(
        id=mapping.id,
        category_desc=mapping.category_desc,
        occurrence_count=0,
        activity_sector_id=mapping.activity_sector_id,
        activity_sector_name=sector.name,
        updated_at=mapping.updated_at,
    )


@router.delete(
    "/category-mappings/{mapping_id}",
    response_model=Message,
    dependencies=[Depends(require_roles([UserRole.SUPER_ADMIN]))],
)
def delete_category_mapping(mapping_id: int, db: Session = Depends(get_db)):
    mapping = db.get(CategorySectorMapping, mapping_id)
    if not mapping:
        raise HTTPException(status_code=404, detail={"code": "mapping_not_found", "message": "Mapping introuvable."})
    db.delete(mapping)
    db.commit()
    return Message(message="Mapping supprime.")
