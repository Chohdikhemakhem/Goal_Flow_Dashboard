from __future__ import annotations

import re
from datetime import datetime
from decimal import Decimal
from io import BytesIO
from typing import Any

import openpyxl
from fastapi import APIRouter, Depends, File, HTTPException, UploadFile, status
from sqlalchemy import and_, select
from sqlalchemy.orm import Session

from app.api.deps import get_current_user, require_roles
from app.db.session import get_db
from app.models.entities import Agency, Target, User
from app.models.enums import TargetType, UserRole

router = APIRouter(
    prefix="/targets/import",
    tags=["targets-import"],
    dependencies=[Depends(require_roles([UserRole.SUPER_ADMIN, UserRole.AGENCY_MANAGER, UserRole.PORTFOLIO_MANAGER]))],
)

MONTH_MAP = {
    "janvier": 1, "fevrier": 2, "mars": 3, "avril": 4,
    "mai": 5, "juin": 6, "juillet": 7, "aout": 8,
    "septembre": 9, "octobre": 10, "novembre": 11, "decembre": 12,
}

def normalize_agency_name(name: str) -> str:
    if not name:
        return ""
    return re.sub(r"\s+", " ", name.strip().upper())

def detect_month_year_from_title(title: str) -> tuple:
    if not title:
        return None, None
    pattern = r"Objectif(?:s)?\s+([a-zA-Z]+)\s+(\d{4})"
    match = re.search(pattern, title, re.IGNORECASE)
    if not match:
        return None, None
    month_name = match.group(1).lower()
    year = int(match.group(2))
    month_number = MONTH_MAP.get(month_name)
    return month_number, year


def parse_excel_file(file_content: bytes) -> dict:
    wb = openpyxl.load_workbook(BytesIO(file_content), data_only=True)
    ws = wb.active
    title = ws.cell(row=1, column=1).value
    if not title:
        raise HTTPException(status_code=400, detail={"code": "missing_title", "message": "Le fichier Excel doit contenir un titre a la cellule A1"})
    title = str(title).strip()
    month, year = detect_month_year_from_title(title)
    headers = {}
    for col in range(1, ws.max_column + 1):
        header_value = ws.cell(row=2, column=col).value
        if header_value:
            headers[col] = str(header_value).strip()
    agency_col = 1
    rows_raw = []
    for row_idx in range(3, ws.max_row + 1):
        agency_value = ws.cell(row=row_idx, column=agency_col).value
        if agency_value is None:
            continue
        agency_name = str(agency_value).strip()
        if agency_name.upper() == "TOTAL":
            continue
        row_data = {"agence": agency_name}
        for col_idx, header in headers.items():
            if col_idx == agency_col:
                continue
            value = ws.cell(row=row_idx, column=col_idx).value
            if value is not None:
                try:
                    row_data[header] = float(value)
                except (ValueError, TypeError):
                    row_data[header] = 0
            else:
                row_data[header] = 0
        rows_raw.append(row_data)
    return {"title": title, "month": month, "year": year, "rows": rows_raw}





def validate_and_map_excel_data(parsed_data: dict, db: Session) -> dict:
    month = parsed_data["month"]
    year = parsed_data["year"]
    rows_raw = parsed_data["rows"]

    if month is None or year is None:
        return {
            "success": False,
            "error": "Mois/annee non detectes. Veuillez les selectionner.",
            "month": None,
            "year": None,
            "rows": [],
        }

    db_agencies = db.scalars(select(Agency)).all()
    agency_name_map = {normalize_agency_name(a.name): a for a in db_agencies}

    errors = []
    valid_rows = []

    for row_raw in rows_raw:
        agency_name_raw = row_raw.get("agence", "")
        normalized_name = normalize_agency_name(agency_name_raw)
        agency = agency_name_map.get(normalized_name)

        if agency is None:
            errors.append(f'Agence inconnue : "{agency_name_raw}"')
            continue

        par_1_30 = row_raw.get("[1-30]", 0)
        par_31_60 = row_raw.get("[31-60]", 0)
        total_general = row_raw.get("Total general", 0)
        encours_sain = row_raw.get("Encours sain", 0)
        par_0 = row_raw.get("PAR 0", 0)
        par_30 = row_raw.get("PAR 30", 0)

        par_30_montant = Decimal("0")
        for col in [
            "[31-60]", "[61-90]", "[91-120]", "[121-150]",
            "[151-180]", "[181-210]", "[211-240]", "[241-270]",
            "[271-300]", "[301-330]", "[331-360]",
        ]:
            val = row_raw.get(col, 0)
            if val:
                par_30_montant += Decimal(str(val))

        objectif_encours = Decimal(str(total_general)) + Decimal(str(encours_sain))

        valid_rows.append({
            "agency_id": agency.id,
            "agency_name": agency.name,
            "target_par_1_30": Decimal(str(par_1_30)),
            "target_par_31_60": Decimal(str(par_31_60)),
            "target_par_30": par_30_montant,
            "target_healthy_outstanding": Decimal(str(encours_sain)),
            "target_outstanding": objectif_encours,
            "target_par_0": Decimal(str(par_0)),
            "target_par": Decimal(str(par_30)),
        })

    return {
        "success": len(errors) == 0,
        "month": month,
        "year": year,
        "rows": valid_rows,
        "errors": errors,
        "detected_agencies": len(rows_raw),
        "valid_agencies": len(valid_rows),
    }


@router.post("/preview")
async def preview_objectives_import(
    file: UploadFile = File(...),
    month: int | None = None,
    year: int | None = None,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
) -> dict:
    content = await file.read()
    try:
        parsed = parse_excel_file(content)
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(
            status_code=400,
            detail={"code": "parse_error", "message": f"Erreur lors de l'analyse du fichier : {str(e)}"},
        )

    if month is not None:
        parsed["month"] = month
    if year is not None:
        parsed["year"] = year

    result = validate_and_map_excel_data(parsed, db)

    if not result["success"]:
        return {
            "success": False,
            "error": result["error"],
            "month": result["month"],
            "year": result["year"],
            "rows": [],
            "detected_agencies": len(parsed.get("rows", [])),
            "valid_agencies": 0,
            "errors": result.get("errors", []),
        }

    existing_count = 0
    existing_agencies = []
    for row in result["rows"]:
        existing = db.scalar(
            select(Target).where(
                and_(
                    Target.target_type == TargetType.AGENCY,
                    Target.agency_id == row["agency_id"],
                    Target.month == result["month"],
                    Target.year == result["year"],
                )
            )
        )
        if existing:
            existing_count += 1
            existing_agencies.append(row["agency_name"])

    return {
        "success": True,
        "month": result["month"],
        "year": result["year"],
        "title": parsed["title"],


@router.post("/confirm")
async def confirm_objectives_import(
    payload: dict,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
) -> dict:
    action = payload.get("action")
    if action not in ("update", "ignore", "cancel"):
        raise HTTPException(
            status_code=400,
            detail={
                "code": "invalid_action",
                "message": "Action invalide. Doit etre 'update', 'ignore' ou 'cancel'",
            },
        )

    if action == "cancel":
        return {"message": "Import annule", "imported": 0, "updated": 0, "skipped": 0}

    month = payload.get("month")
    year = payload.get("year")
    rows = payload.get("rows", [])

    if not rows:
        raise HTTPException(
            status_code=400,
            detail={"code": "no_data", "message": "Aucune donnee a importer"},
        )

    imported = 0
    updated = 0
    skipped = 0

    try:
        for row in rows:
            agency_id = row["agency_id"]
            existing = db.scalar(
                select(Target).where(
                    and_(
                        Target.target_type == TargetType.AGENCY,
                        Target.agency_id == agency_id,
                        Target.month == month,
                        Target.year == year,
                    )
                )
            )

            if existing:
                if action == "ignore":
                    skipped += 1
                    continue

                if action == "update":
                    existing.target_par_1_30 = row["target_par_1_30"]
                    existing.target_par_31_60 = row["target_par_31_60"]
                    existing.target_par_30 = row["target_par_30"]
                    existing.target_healthy_outstanding = row["target_healthy_outstanding"]
                    existing.target_outstanding = row["target_outstanding"]
                    existing.target_par_0 = row["target_par_0"]
                    existing.target_par = row["target_par"]
                    updated += 1
                continue

            new_target = Target(
                target_type=TargetType.AGENCY,
                agency_id=agency_id,
                month=month,
                year=year,
                target_par_1_30=row["target_par_1_30"],
                target_par_31_60=row["target_par_31_60"],
                target_par_30=row["target_par_30"],
                target_healthy_outstanding=row["target_healthy_outstanding"],
                target_outstanding=row["target_outstanding"],
                target_par_0=row["target_par_0"],
                target_par=row["target_par"],
                target_disbursement_count=0,
                target_nb_clients=0,
                target_disbursement=Decimal("0"),
                created_by=user.id,
                created_at=datetime.now(),
            )
            db.add(new_target)
            imported += 1

        db.commit()

    except Exception as e:
        db.rollback()
        raise HTTPException(
            status_code=500,
            detail={"code": "import_error", "message": f"Erreur lors de l'import : {str(e)}"},
        )

    return {
        "message": f"Import termine : {imported} objectifs crees, {updated} mis a jour, {skipped} ignores",
        "imported": imported,
