"""Parsing and transactional persistence for monthly agency-objective workbooks.

The workbook is deliberately parsed again during confirmation.  The browser never
gets to submit calculated values as the source of truth: it only displays the
preview produced by this module.
"""

from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
import re
import unicodedata

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models.entities import Agency, Target, User
from app.models.enums import TargetType
from app.services.objective_active_period import calculate_active_period


MAX_OBJECTIVES_EXCEL_BYTES = 10 * 1024 * 1024
COHORT_HEADERS = (
    "[1-30]", "[31-60]", "[61-90]", "[91-120]", "[121-150]", "[151-180]",
    "[181-210]", "[211-240]", "[241-270]", "[271-300]", "[301-330]", "[331-360]",
)
REQUIRED_HEADERS = ("Agence/Cohorte", *COHORT_HEADERS, "Total général", "Encours sain", "PAR 0", "PAR 30")
DATA_HEADERS = REQUIRED_HEADERS[1:]
MONTHS = {
    "janvier": 1, "fevrier": 2, "mars": 3, "avril": 4, "mai": 5, "juin": 6,
    "juillet": 7, "aout": 8, "septembre": 9, "octobre": 10, "novembre": 11,
    "decembre": 12,
}
MONTH_LABELS = (
    "Janvier", "Février", "Mars", "Avril", "Mai", "Juin", "Juillet", "Août",
    "Septembre", "Octobre", "Novembre", "Décembre",
)


def normalise(value: object) -> str:
    """Case/space/accent-insensitive comparison key, without fuzzy matching."""
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(char for char in text if not unicodedata.combining(char))
    return re.sub(r"\s+", " ", text.replace("\xa0", " ").strip()).casefold()


HEADER_KEYS = {normalise(header): header for header in DATA_HEADERS}


def _period_from_text(value: object) -> tuple[int, int] | None:
    text = normalise(value)
    match = re.search(r"objectif\s+([a-z]+)\s+(20\d{2})", text)
    if not match:
        return None
    month = MONTHS.get(match.group(1))
    return (month, int(match.group(2))) if month else None


def _read_decimal(value: object, label: str, *, percentage: bool = False) -> Decimal:
    if isinstance(value, bool) or value is None or (isinstance(value, str) and not value.strip()):
        raise ValueError(f"Valeur numérique invalide pour {label}")
    is_percent_string = isinstance(value, str) and "%" in value
    text = str(value).strip().replace("\xa0", " ").replace(" ", "").replace("%", "")
    # French spreadsheets commonly use a comma decimal separator.  Keep a dot
    # decimal separator when both separators are present.
    if "," in text and "." in text:
        text = text.replace(",", "") if text.rfind(".") > text.rfind(",") else text.replace(".", "").replace(",", ".")
    elif "," in text:
        text = text.replace(",", ".")
    try:
        result = Decimal(text)
    except (InvalidOperation, ValueError) as exc:
        raise ValueError(f"Valeur numérique invalide pour {label}: {value!r}") from exc
    if not result.is_finite() or result < 0:
        raise ValueError(f"Valeur numérique invalide pour {label}: {value!r}")
    if percentage and (is_percent_string or result > 1):
        result /= Decimal("100")
    if percentage and result > 1:
        raise ValueError(f"Pourcentage invalide pour {label}: {value!r}")
    return result


def _columns_from_second_row(sheet) -> dict[str, int]:
    """The workbook has a title row followed by the actual data headers.

    A1 labels the agency column and B1 contains the period title.  Row 2 starts
    with an intentionally blank cell, then carries the cohort and PAR headers.
    """
    columns: dict[str, int] = {}
    header_cells = next(sheet.iter_rows(min_row=2, max_row=2))
    for index, cell in enumerate(header_cells, start=1):
        canonical = HEADER_KEYS.get(normalise(cell.value))
        if canonical:
            columns[canonical] = index
    missing = next((header for header in DATA_HEADERS if header not in columns), None)
    if missing:
        raise ValueError(f"Colonne obligatoire manquante : {missing}")
    if normalise(sheet.cell(row=1, column=1).value) != normalise("Agence/Cohorte"):
        raise ValueError("Structure Excel invalide : cellule A1 doit contenir Agence/Cohorte")
    return columns


def parse_objectives_excel(content: bytes, filename: str) -> tuple[tuple[int, int] | None, list[dict[str, object]]]:
    if not content:
        raise ValueError("Le fichier Excel est vide")
    if len(content) > MAX_OBJECTIVES_EXCEL_BYTES:
        raise ValueError("Le fichier Excel dépasse la taille maximale autorisée (10 Mo)")
    if not filename.lower().endswith(".xlsx"):
        raise ValueError("Format non autorisé : utilisez un fichier Excel .xlsx")
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True, keep_vba=False)
        sheet = workbook.active
        columns = _columns_from_second_row(sheet)
        # The period title is defined by the workbook contract: B1.
        period = _period_from_text(sheet.cell(row=1, column=2).value)
        rows: list[dict[str, object]] = []
        for row_number, cells in enumerate(sheet.iter_rows(min_row=3), start=3):
            agency_cell = cells[0]
            agency_name = str(agency_cell.value or "").strip()
            if not agency_name:
                continue
            if normalise(agency_name) == "total":
                continue
            values: dict[str, Decimal] = {}
            for header in DATA_HEADERS:
                if header == "PAR 0":
                    continue
                cell = cells[columns[header] - 1]
                values[header] = _read_decimal(
                    cell.value,
                    f"{header} (ligne {row_number}, agence {agency_name})",
                    percentage=header == "PAR 30",
                )
            par30_amount = sum((values[header] for header in COHORT_HEADERS[1:]), Decimal("0"))
            total_general = values["Total général"]
            healthy_outstanding = values["Encours sain"]
            par0_denominator = total_general + healthy_outstanding
            par0 = total_general / par0_denominator if par0_denominator else Decimal("0")
            rows.append({
                "row_number": row_number,
                "agency_name": agency_name,
                "agency_key": normalise(agency_name),
                "target_par_1_30": values["[1-30]"],
                "target_par_31_60": values["[31-60]"],
                "target_par_30": par30_amount,
                "target_healthy_outstanding": healthy_outstanding,
                "target_outstanding": par0_denominator,
                # PAR 0 is an Excel formula (=Total général / (Total général +
                # Encours sain)); calculate it here rather than relying on a
                # cached workbook formula result.
                "target_par_0": par0,
                "target_par": values["PAR 30"],
            })
        if not rows:
            raise ValueError("Aucune agence n'a été trouvée dans le fichier Excel")
        return period, rows
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("Le fichier Excel est illisible ou corrompu") from exc


def _period_label(month: int, year: int) -> str:
    return f"{MONTH_LABELS[month - 1]} {year}"


def build_objective_import_preview(
    db: Session,
    content: bytes,
    filename: str,
    month: int | None = None,
    year: int | None = None,
) -> dict[str, object]:
    detected_period, rows = parse_objectives_excel(content, filename)
    if (month is None) != (year is None):
        raise ValueError("Le mois et l'année doivent être renseignés ensemble")
    if month is not None and not 1 <= month <= 12:
        raise ValueError("Mois invalide")
    if year is not None and year < 2000:
        raise ValueError("Année invalide")
    period = detected_period or ((month, year) if month and year else None)

    agencies = db.scalars(select(Agency)).all()
    agency_by_key = {normalise(agency.name): agency for agency in agencies}
    errors: list[str] = []
    seen: set[str] = set()
    preview_rows: list[dict[str, object]] = []
    agency_ids: list[int] = []
    for row in rows:
        key = str(row["agency_key"])
        if key in seen:
            errors.append(f"Agence dupliquée dans le fichier : {row['agency_name']}")
            continue
        seen.add(key)
        agency = agency_by_key.get(key)
        if agency is None:
            errors.append(f"Agence inconnue : {row['agency_name']} (ligne {row['row_number']})")
            continue
        entry = {key: value for key, value in row.items() if key not in {"agency_key", "row_number"}}
        entry["agency_id"] = agency.id
        entry["agency_name"] = agency.name
        preview_rows.append(entry)
        agency_ids.append(agency.id)

    existing_ids: set[int] = set()
    if period and agency_ids:
        existing_ids = set(db.scalars(select(Target.agency_id).where(
            Target.target_type == TargetType.AGENCY,
            Target.agent_id.is_(None),
            Target.month == period[0], Target.year == period[1],
            Target.agency_id.in_(agency_ids),
        )).all())
    for entry in preview_rows:
        entry["existing"] = entry["agency_id"] in existing_ids
    return {
        "file_name": filename,
        "period_detected": detected_period is not None,
        "period_available": period is not None,
        "month": period[0] if period else None,
        "year": period[1] if period else None,
        "period_label": _period_label(*period) if period else None,
        "detected_agencies": len(rows),
        "valid_agencies": len(preview_rows),
        "unknown_agencies": sum(error.startswith("Agence inconnue") for error in errors),
        "existing_count": len(existing_ids),
        "errors": errors,
        "can_import": bool(preview_rows),
        "rows": preview_rows,
    }


def import_objectives(
    db: Session,
    content: bytes,
    filename: str,
    actor: User,
    mode: str,
    month: int | None,
    year: int | None,
    active_period_reference_date: date | None = None,
) -> dict[str, object]:
    if mode not in {"update", "skip"}:
        raise ValueError("Mode d'import invalide")
    preview = build_objective_import_preview(db, content, filename, month, year)
    if not preview["can_import"]:
        raise ValueError("Aucune agence valide ne peut être importée")
    if not preview["period_available"]:
        raise ValueError("Le mois et l'année doivent être sélectionnés avant l'import")
    period_month, period_year = int(preview["month"]), int(preview["year"])
    rows = preview["rows"]
    active_from, active_until = calculate_active_period(active_period_reference_date)
    agency_ids = [int(row["agency_id"]) for row in rows]
    targets = db.scalars(select(Target).where(
        Target.target_type == TargetType.AGENCY, Target.agent_id.is_(None),
        Target.month == period_month, Target.year == period_year,
        Target.agency_id.in_(agency_ids),
    )).all()
    existing_by_agency = {target.agency_id: target for target in targets}
    imported = updated = skipped = 0
    excel_fields = (
        "target_par_1_30", "target_par_31_60", "target_par_30", "target_healthy_outstanding",
        "target_outstanding", "target_par_0", "target_par",
    )
    try:
        for row in rows:
            target = existing_by_agency.get(int(row["agency_id"]))
            if target:
                if mode == "skip":
                    skipped += 1
                    continue
                for field in excel_fields:
                    setattr(target, field, row[field])
                target.updated_by = actor.id
                target.updated_at = datetime.now()
                updated += 1
            else:
                # Values absent from the workbook intentionally use the model's
                # current defaults; no fabricated objective is supplied here.
                db.add(Target(
                    target_type=TargetType.AGENCY, agency_id=int(row["agency_id"]), agent_id=None,
                    month=period_month, year=period_year, active_from=active_from, active_until=active_until,
                    created_by=actor.id, **{field: row[field] for field in excel_fields},
                ))
                imported += 1
        db.commit()
    except Exception:
        db.rollback()
        raise
    return {"imported": imported, "updated": updated, "skipped": skipped, "period_label": preview["period_label"]}
